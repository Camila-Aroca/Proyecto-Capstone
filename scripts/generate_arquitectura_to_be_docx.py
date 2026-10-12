"""Genera el documento de arquitectura to-be del SAAD: despliegue objetivo en AWS.

Complementa al Documento de Arquitectura de Software (DAS-SAAD-2026), que describe el
as-is. Este documento no lo modifica: describe solo la arquitectura objetivo, su costo
estimado y su justificación (D2-16), con el diagrama UML de despliegue (D2-21).

Ejecución: python scripts/generate_arquitectura_to_be_docx.py
Salida:    docs/08-arquitectura/Arquitectura_ToBe_AWS.docx
Figuras:   docs/07-diagramas/despliegue_to_be_aws.png y contenedores_docker.png, exportadas
           desde sus .drawio (ver sección 11.1 del documento)
"""

from __future__ import annotations

from pathlib import Path

from docx.shared import Cm

from docx_builder import (
    BRAND_SKY_SOFT,
    BRAND_THEME,
    add_brand_cover,
    add_figure,
    add_table,
    add_toc,
    apply_theme,
    bullets,
    h1,
    h2,
    h3,
    new_document,
    note,
    p,
    save,
)

OUTPUT_PATH = Path("docs/08-arquitectura/Arquitectura_ToBe_AWS.docx")
DEPLOYMENT_UML = Path("docs/07-diagramas/despliegue_to_be_aws.png")
CONTAINERS_UML = Path("docs/07-diagramas/contenedores_docker.png")


def build_cover(doc) -> None:
    add_brand_cover(
        doc,
        title="Arquitectura to-be: despliegue en AWS",
        subtitle=(
            "SAAD — Sistema de Análisis y Anticipación de Demanda de Urgencia "
            "en Salud Mental para la Región Metropolitana"
        ),
        project="Capstone APT · Ingeniería en Informática · Duoc UC Antonio Varas",
        kicker="Arquitectura objetivo",
        meta_rows=[
            ("Código del documento", "ATB-SAAD-2026"),
            ("Versión", "1.0"),
            ("Fecha de emisión", "12 de octubre de 2026"),
            ("Estado", "Propuesta para revisión de la docente y del equipo"),
            ("Clasificación", "Uso académico e interno del equipo de proyecto"),
            ("Tareas", "D2-16 Modificar o justificar la arquitectura (#83) · "
                       "D2-21 Diagrama UML de despliegue (#88)"),
            ("Elaborado por", "Equipo SAAD — Camila A., Felipe R., Cristopher R., Catalina"),
            ("Documento relacionado", "Documento de Arquitectura de Software (DAS-SAAD-2026), as-is"),
            ("Repositorio", "docs/08-arquitectura/"),
        ],
    )


def build_control(doc) -> None:
    h1(doc, "Control del documento")
    add_table(
        doc,
        ["Versión", "Fecha", "Autor(es)", "Descripción del cambio"],
        [
            ("1.0", "12-10-2026", "Equipo SAAD",
             "Primera versión: arquitectura objetivo en AWS, diagrama UML de despliegue, estimación "
             "de costos, justificación del proveedor, de no usar varias nubes y de las tecnologías, "
             "decisiones, relación con la demo y riesgos."),
        ],
        [Cm(2.0), Cm(2.6), Cm(3.6), Cm(9.5)],
        font_size=9.5,
    )


def build_purpose(doc) -> None:
    h1(doc, "1. Propósito y alcance")
    p(
        doc,
        "Este documento describe **cómo se desplegaría el SAAD en la nube** (arquitectura to-be). "
        "La arquitectura vigente (as-is) está en el Documento de Arquitectura de Software "
        "(DAS-SAAD-2026) y no se repite aquí. El despliegue descrito está **diseñado y "
        "presupuestado, pero no implementado**: no se han creado recursos en AWS.",
    )
    p(doc, "Responde a la revisión de la arquitectura (D2-16), que pidió:")
    bullets(
        doc,
        [
            "una estimación de precios (sección 4);",
            "justificar por qué AWS y no otra nube (sección 5);",
            "justificar por qué no se usan varias nubes (sección 6);",
            "justificar por qué se eligieron las tecnologías (sección 7 y DAS, sección 10.1).",
        ],
    )
    p(doc, "Incluye además el diagrama UML de despliegue que pide D2-21 (sección 3).")


def build_as_is_to_be(doc) -> None:
    h1(doc, "2. Del as-is al to-be")
    p(
        doc,
        "**El estilo arquitectónico se mantiene**: pipeline por lotes, API REST de solo lectura que "
        "responde desde memoria y aplicación web separada (DAS, secciones 3.1 y 3.2). Lo que cambia "
        "es dónde se ejecuta cada parte. Esa separación es justamente la que permite llevarla a la "
        "nube sin rehacer el código: la API y la aplicación ya son procesos independientes que solo "
        "se comunican por HTTP.",
    )
    add_table(
        doc,
        ["Aspecto", "As-is (repositorio)", "To-be (AWS)"],
        [
            ("Ejecución", "Docker Compose en el equipo de un integrante.", "Servicios gestionados de AWS en us-east-1."),
            ("Aplicación web", "Contenedor nginx que sirve el build y reenvía `/api`.",
             "Build en S3, servido por CloudFront, que también reenvía `/api/*`."),
            ("API", "Contenedor `saad-api` en el equipo.", "La misma imagen en ECS con Fargate, detrás de un balanceador."),
            ("Datos", "Volumen local `data/` y `reports/`.", "Un bucket S3 con RAW, manifest, outputs y reportes."),
            ("Pipeline", "Contenedor bajo demanda en el equipo.",
             "Tarea programada semanal en Fargate; la reconstrucción completa y el benchmark siguen en el equipo."),
            ("Base relacional", "Diseñada, no usada (ADR-02).", "Aurora Serverless v2, opcional, cuando la API la use."),
            ("HTTPS", "No (HTTP local).", "CloudFront y balanceador con certificados de ACM."),
            ("Acceso", "Solo desde el equipo.", "URL pública para usuarios fuera del equipo."),
        ],
        [Cm(3.2), Cm(6.6), Cm(7.9)],
        font_size=9.0,
        first_col_bold=True,
    )


def build_deployment(doc) -> None:
    h1(doc, "3. Diagrama de despliegue")
    add_figure(
        doc,
        DEPLOYMENT_UML,
        "**Figura 1.** Diagrama de despliegue to-be en AWS (D2-21). Fuente editable en draw.io: "
        "`docs/07-diagramas/despliegue_to_be_aws.drawio`.",
        width_cm=17.0,
    )
    h2(doc, "3.1 Nodos y componentes desplegados")
    add_table(
        doc,
        ["Nodo", "Qué despliega o guarda", "Por qué así"],
        [
            ("Amazon CloudFront", "Punto de entrada HTTPS: `/*` al sitio y `/api/*` al balanceador.",
             "Conserva el origen único del as-is (sin CORS) y entrega certificados con ACM."),
            ("Amazon S3 · sitio", "Build de Vite (`frontend/dist`).",
             "Archivos estáticos: no necesitan un contenedor encendido. Privado, solo lo lee CloudFront."),
            ("Application Load Balancer", "Reenvía `/api/*` a la tarea de la API.",
             "Da una dirección estable a la tarea de Fargate, cuya IP cambia en cada reinicio."),
            ("ECS con AWS Fargate", "Una tarea con el contenedor `saad-api` (0,5 vCPU, 2 GB).",
             "Contenedor siempre encendido sin servidor que administrar. 2 GB dejan margen sobre los "
             "~460 MB observados en ejecución."),
            ("Amazon S3 · datos", "RAW y manifest (~12 GB), outputs (`data/processed/`) y `reports/`.",
             "Un solo bucket con prefijos: la tarea pipeline lo lee y escribe; la API lee los outputs al "
             "arrancar."),
            ("ECS con Fargate · tarea pipeline", "Tarea programada con la imagen `saad-pipeline`, independiente de la API (4 vCPU, 16 GB, 50 GiB de disco).",
             "Cómputo solo durante la ejecución semanal (~1 h); sin máquina encendida el resto del tiempo."),
            ("EventBridge Scheduler", "Disparo semanal de la tarea pipeline.",
             "Programación gestionada, sin un servidor con cron."),
            ("Amazon ECR", "Imágenes versionadas de `saad-api`.", "Registro privado del que Fargate descarga."),
            ("CloudWatch Logs", "Registros de la API.", "Diagnóstico sin acceder a un servidor."),
            ("Aurora Serverless v2 (opcional)", "PostgreSQL con PostGIS (RT-03).",
             "Se agrega cuando la API consulte una base (ADR-02). Se pausa sola cuando no se usa."),
            ("Equipo del proyecto", "Reconstrucción completa del pipeline y benchmark del modelo.",
             "Son ejecuciones largas y ocasionales (solo cuando cambia el modelo); no justifican cómputo "
             "en la nube."),
        ],
        [Cm(3.9), Cm(6.5), Cm(7.3)],
        font_size=8.8,
        first_col_bold=True,
    )
    h2(doc, "3.2 Red y publicación")
    bullets(
        doc,
        [
            "La tarea corre en subredes públicas con IP propia para descargar su imagen y sus datos. "
            "En subredes privadas exigiría un NAT Gateway, que por sí solo costaría más que la tarea. "
            "Su grupo de seguridad solo admite tráfico desde el balanceador.",
            "El equipo publica la imagen (`docker push` a ECR) y el build de la aplicación (`aws s3 sync`). "
            "Los datos los publica la tarea pipeline (sección 3.3).",
        ],
    )
    h2(doc, "3.3 Actualización de datos")
    p(
        doc,
        "Las fuentes no cambian al mismo ritmo, y el pipeline solo necesita seguir a la que cambia "
        "seguido:",
    )
    add_table(
        doc,
        ["Fuente", "Frecuencia observada", "Evidencia"],
        [
            ("Urgencias DEIS, año en curso", "Varias veces por semana, ~1 día de desfase",
             "Dos snapshots de 2026 con hash distinto (3 y 7 de octubre); la descarga del 7 traía datos "
             "hasta el 6. Inferido, no es un calendario oficial."),
            ("Urgencias DEIS, años cerrados", "No cambian", "Hash estable en el manifest."),
            ("Egresos DEIS", "Anual, con meses de rezago", "Un archivo por año (2020-2025)."),
            ("Catálogo de establecimientos", "Irregular", "El archivo vigente es del 25-08-2026."),
            ("Censo, proyecciones INE, Casen, Estadísticas de Género", "Fijas", "Años o ediciones cerrados."),
        ],
        [Cm(4.4), Cm(4.4), Cm(8.9)],
        font_size=8.8,
        first_col_bold=True,
    )
    bullets(
        doc,
        [
            "**Cuándo:** EventBridge Scheduler dispara la tarea pipeline una vez por semana, el lunes. La "
            "semana DEIS va de domingo a sábado, así que el lunes ya cerró y pasó el día de desfase. Coincide "
            "con el modelo, que pronostica por semana.",
            "**Qué corre:** baja de S3 el RAW y el manifest, ejecuta `run_pipeline.py "
            "--refresh-current-urgencias` (solo vuelve a consultar el año en curso y regenera lo afectado), "
            "regenera el pronóstico y sube los resultados a S3. Así el pronóstico publicado siempre está "
            "vigente (R-02).",
            "**Cómo lo toma la API:** la API lee los datos solo al arrancar (ADR-06 del DAS), por lo que la "
            "tarea termina forzando un nuevo despliegue del servicio `saad-api` (R-10).",
            "**Qué queda en el equipo:** la reconstrucción completa y el benchmark del modelo, que se lanzan "
            "a mano cuando cambia el modelo.",
            "**Alternativas descartadas:** Lambda (límite de 15 minutos y 10 GB de disco), GitHub Actions "
            "(runners gratuitos con ~14 GB de disco, justo para el RAW) y seguir solo en el equipo (depende "
            "de que alguien lo ejecute cada semana).",
        ],
    )
    h2(doc, "3.4 Contenedores")
    p(
        doc,
        "El to-be no crea imágenes nuevas: despliega las que ya construye el repositorio. La figura 2 "
        "muestra esos contenedores tal como corren hoy con Docker Compose. En AWS, `saad-api` pasa a "
        "ser la tarea de Fargate, el contenido de `saad-web` (el build) se publica en S3 en lugar de "
        "servirse con nginx, y el contenedor del pipeline pasa a ser la tarea programada de Fargate "
        "(sigue disponible en el equipo para la reconstrucción completa).",
    )
    add_figure(
        doc,
        CONTAINERS_UML,
        "**Figura 2.** Diagrama UML de despliegue de los contenedores Docker (`docker-compose.yml`). "
        "Fuente editable en draw.io: `docs/07-diagramas/contenedores_docker.drawio`.",
        width_cm=17.0,
    )


def build_costs(doc) -> None:
    h1(doc, "4. Estimación de costos")
    p(
        doc,
        "Estimación mensual en us-east-1 con precios de lista a la fecha de redacción, 730 horas "
        "al mes y tráfico bajo (uso académico y de validación). Son valores **estimados, no "
        "medidos**: deben confirmarse ingresando estas mismas configuraciones en AWS Pricing "
        "Calculator y adjuntando el enlace como evidencia.",
    )
    add_table(
        doc,
        ["Servicio", "Configuración para la calculadora", "Precio unitario usado (US$)", "US$/mes"],
        [
            ("Fargate · tarea api", "Linux x86, 1 tarea, 0,5 vCPU, 2 GB, 730 h, 20 GB efímeros",
             "0,04048 por vCPU·h; 0,004445 por GB·h", "21,27"),
            ("Fargate · tarea pipeline", "Linux x86, 4 vCPU, 16 GB, 1 h × 4,35 semanas, 50 GB efímeros",
             "Ídem; 0,000111 por GiB·h sobre los 30 GiB que exceden los 20 gratuitos", "1,03"),
            ("Application Load Balancer", "1 balanceador, 730 h, 1 LCU", "0,0225 por h; 0,008 por LCU·h", "22,27"),
            ("IPv4 públicas", "3 direcciones 730 h (2 del balanceador, 1 de la API) y 1 de la tarea "
             "pipeline ~4 h", "0,005 por IP·h", "10,97"),
            ("Amazon CloudFront", "Menos de 1 TB de salida y 10 millones de solicitudes",
             "Capa gratuita permanente", "0,00"),
            ("Amazon S3 Standard", "~12,5 GB (RAW, outputs, reportes y sitio)", "0,023 por GB-mes", "0,29"),
            ("Amazon ECR", "~3,5 GB (imágenes saad-api y saad-pipeline, ~865 MB cada una, dos versiones de cada una)", "0,10 por GB-mes", "0,35"),
            ("CloudWatch Logs", "~1 GB ingerido al mes", "0,50 por GB", "0,53"),
            ("EventBridge Scheduler", "~4 invocaciones al mes", "Capa gratuita", "0,00"),
            ("AWS Certificate Manager", "Certificados públicos", "Sin costo", "0,00"),
            ("**Total base**", "", "", "**≈ 56,71**"),
            ("Aurora Serverless v2 (opcional)", "0,5 ACU durante 60 h al mes, pausada el resto; 1 GB",
             "0,12 por ACU·h; 0,10 por GB-mes", "≈ 3,70"),
            ("Aurora Serverless v2 sin pausa", "0,5 ACU durante 730 h; 1 GB", "Ídem", "≈ 43,90"),
        ],
        [Cm(3.9), Cm(6.4), Cm(5.0), Cm(2.4)],
        font_size=8.6,
        first_col_bold=True,
    )
    bullets(
        doc,
        [
            "**Base: ≈ US$57 al mes (≈ US$680 al año).** Casi el 60 % es costo fijo del balanceador y "
            "de las IPv4, que se paga aunque no haya visitas. La actualización semanal de datos agrega "
            "poco más de US$1.",
            "**Con base de datos:** ≈ US$60 si Aurora se pausa fuera de uso; ≈ US$100 si queda encendida.",
            "**Para reducirlo:** Fargate sobre ARM (Graviton) baja el cómputo cerca de un 20 %; dejar la "
            "tarea en 0 fuera de las demos ahorra los ~US$21 de Fargate, no el balanceador. Fargate Spot "
            "no sirve: puede interrumpir la única tarea.",
            "**Fuera de la estimación:** la reconstrucción completa y el benchmark corren en el equipo "
            "(sin costo de nube); un dominio propio en Route 53 sumaría US$0,50 al mes más su registro.",
        ],
    )
    h2(doc, "4.1 Alternativa más barata considerada")
    p(
        doc,
        "Una sola máquina EC2 (t3.small) ejecutando el `docker-compose.yml` actual costaría cerca de "
        "US$15 a US$20 al mes. Se descartó porque el equipo tendría que mantener el servidor "
        "(parches, disco, certificados, reinicios), lo que choca con un equipo de cuatro personas con "
        "dedicación parcial (RO-01). Es la alternativa de respaldo si el costo resulta inviable.",
    )


def build_why_aws(doc) -> None:
    h1(doc, "5. Por qué AWS y no otra nube")
    p(
        doc,
        "La comparación se hace contra Google Cloud y Microsoft Azure, con criterios que salen de la "
        "propia arquitectura del SAAD y no de preferencias del equipo.",
    )
    add_table(
        doc,
        ["Criterio", "AWS", "Google Cloud", "Azure"],
        [
            ("Contenedor de API con estado en memoria, arranque lento y cálculo en segundo plano "
             "(ADR-06 del DAS)",
             "Fargate: tarea siempre encendida con CPU dedicada; es su modo normal.",
             "Cloud Run: para esta carga necesita una instancia mínima y CPU asignada siempre; pierde "
             "su ventaja de escalar a cero.",
             "Container Apps: igual que Cloud Run, requiere una réplica mínima encendida."),
            ("PostgreSQL con PostGIS gestionado que se pausa solo (RT-03)",
             "Aurora Serverless v2 con pausa automática.",
             "Cloud SQL tiene PostGIS, pero la instancia queda encendida; solo se detiene a mano.",
             "Flexible Server tiene PostGIS; solo se detiene a mano y por un tiempo limitado."),
            ("Sitio estático y API bajo un mismo dominio",
             "CloudFront y S3.", "Firebase Hosting o balanceador con Cloud CDN.", "Static Web Apps o Front Door."),
            ("Región más cercana a los usuarios",
             "São Paulo (este diseño usa us-east-1 por costo).", "Santiago.", "São Paulo."),
        ],
        [Cm(4.6), Cm(4.3), Cm(4.4), Cm(4.4)],
        font_size=8.6,
        first_col_bold=True,
    )
    p(
        doc,
        "**Conclusión.** Las tres nubes pueden ejecutar el SAAD. AWS se elige por dos razones "
        "técnicas: es la única que ofrece la base espacial planificada (PostgreSQL con PostGIS) "
        "gestionada y con pausa automática, lo que permite sumar RT-03 sin un costo fijo adicional; "
        "y la principal ventaja de las alternativas, escalar a cero, no aplica a una API que debe "
        "quedar encendida con sus datos en memoria. A cambio se acepta el costo fijo del balanceador "
        "y una latencia mayor que la que daría la región de Google Cloud en Santiago. Si la base "
        "relacional se descartara, Google Cloud sería una alternativa equivalente; por eso el código "
        "no depende de servicios propietarios (ADR-11).",
    )
    note(
        doc,
        "Las características de los servicios de cada proveedor cambian con frecuencia. Deben "
        "verificarse en su documentación oficial a la fecha de cada entrega.",
        fill=BRAND_SKY_SOFT,
    )


def build_single_cloud(doc) -> None:
    h1(doc, "6. Por qué una sola nube")
    bullets(
        doc,
        [
            "**No hay un requisito que lo pida.** El SAAD es un servicio de solo lectura para "
            "planificación, sin compromiso de disponibilidad que justifique redundancia entre "
            "proveedores ni datos que deban residir en dos lugares.",
            "**Duplicaría el trabajo y el costo.** Dos balanceadores, dos registros de imágenes, dos "
            "esquemas de permisos y dos monitoreos para un equipo de cuatro personas con dedicación "
            "parcial (RO-01) y sin presupuesto (RO-02). El tráfico de datos entre nubes también se cobra.",
            "**La portabilidad ya se obtiene sin varias nubes.** Imágenes Docker estándar, datos en "
            "Parquet y PostgreSQL estándar: mover el SAAD a otro proveedor es cambiar el destino del "
            "despliegue, no el código (ADR-11).",
        ],
    )


def build_technologies(doc) -> None:
    h1(doc, "7. Tecnologías del to-be y su justificación")
    p(
        doc,
        "Las tecnologías del as-is (Python, FastAPI, pandas y PyArrow, Parquet, LightGBM, React, "
        "Vite, Tailwind, Leaflet, entre otras) se justifican en la sección 10.1 del DAS. Esta tabla "
        "cubre solo lo que agrega el despliegue.",
    )
    add_table(
        doc,
        ["Tecnología", "Uso", "Por qué", "Alternativa descartada"],
        [
            ("Docker", "Empaquetado de la API y el pipeline.",
             "Dos imágenes independientes (`saad-api` y `saad-pipeline`) que corren igual en el equipo y en Fargate.",
             "Instalar dependencias en cada servidor."),
            ("ECS con AWS Fargate", "Ejecución de la API y de la tarea pipeline.",
             "La API como contenedor siempre encendido y el pipeline como tarea que solo cobra mientras "
             "corre, ambos sin administrar servidores (sección 5).",
             "EC2 con Docker Compose: más barato, pero el equipo mantiene el servidor (sección 4.1)."),
            ("EventBridge Scheduler", "Ejecución semanal del pipeline.",
             "Programación gestionada y gratuita a esta escala (sección 3.3).",
             "Cron en un servidor propio; Lambda, por sus límites de tiempo y disco."),
            ("Application Load Balancer", "Entrada estable a la API.",
             "La IP de una tarea Fargate cambia en cada reinicio.", "IP pública directa: inestable y sin HTTPS."),
            ("S3 y CloudFront", "Sitio estático, datos y HTTPS.",
             "Los estáticos no necesitan un contenedor; CloudFront une sitio y API en un dominio.",
             "Servir el sitio desde un contenedor encendido."),
            ("Amazon ECR", "Registro de imágenes.", "Privado y en la misma red que Fargate.",
             "Registro público: expone la imagen."),
            ("CloudWatch Logs", "Registros de la API.", "Integrado con Fargate, sin agentes.", "—"),
            ("Aurora Serverless v2", "PostgreSQL con PostGIS (opcional).",
             "Gestionada y con pausa automática: sin costo de cómputo mientras no se usa.",
             "RDS siempre encendida (~US$15 al mes)."),
        ],
        [Cm(3.5), Cm(3.6), Cm(6.6), Cm(4.0)],
        font_size=8.6,
        first_col_bold=True,
    )


def build_decisions(doc) -> None:
    h1(doc, "8. Decisiones arquitectónicas")
    p(doc, "Continúan la numeración del registro de decisiones del DAS (ADR-01 a ADR-09).")
    decisiones = [
        (
            "ADR-10 · Despliegue objetivo en AWS con contenedores gestionados y sitio estático",
            "La plataforma debe poder publicarse con HTTPS para usuarios fuera del equipo, sin que "
            "nadie administre servidores y con un costo acotado.",
            "La API corre como una tarea de ECS en Fargate detrás de un balanceador; la aplicación se "
            "publica en S3 con CloudFront, que también enruta `/api/*`; los datos viven en S3. Una tarea "
            "de Fargate programada con EventBridge actualiza los datos y el pronóstico cada semana. La "
            "base PostgreSQL con PostGIS se agrega en Aurora Serverless v2 cuando la API la use. La "
            "reconstrucción completa y el benchmark siguen en el equipo.",
            "A favor: mismas imágenes que en el as-is, sin servidores, con HTTPS y con datos que se "
            "actualizan solos. En contra: ≈ US$57 al mes, la mayor parte fija por el balanceador; exige "
            "resolver R-01 antes de desplegar.",
        ),
        (
            "ADR-11 · Una sola nube y sin servicios propietarios en el código",
            "Repartir el sistema entre proveedores no responde a ningún requisito, y quedar atado a "
            "uno haría costoso cambiarlo.",
            "Se usa un solo proveedor. El código solo depende de estándares: imágenes Docker, archivos "
            "Parquet en almacenamiento de objetos y PostgreSQL. Lo propio de AWS queda en la "
            "configuración del despliegue.",
            "A favor: operación y costo simples, migrable cambiando el destino. En contra: sin "
            "redundancia ante una caída del proveedor, aceptable para un servicio de planificación.",
        ),
    ]
    for titulo, contexto, decision, consecuencias in decisiones:
        h3(doc, titulo)
        add_table(
            doc,
            ["Aspecto", "Detalle"],
            [("Contexto", contexto), ("Decisión", decision), ("Consecuencias", consecuencias)],
            [Cm(3.0), Cm(14.7)],
            font_size=9.2,
            first_col_bold=True,
            keep_together=True,
        )


def build_demo(doc) -> None:
    h1(doc, "9. Relación con la solución y la demo")
    bullets(
        doc,
        [
            "**La demo usa el as-is.** `docker compose up` levanta la aplicación y la API con las mismas "
            "imágenes que irían a AWS; la plataforma completa se muestra sin depender de la nube.",
            "**El to-be se justifica para el acceso público**, por ejemplo la validación con un usuario "
            "piloto fuera del equipo: da una URL HTTPS estable sin exponer el equipo de un integrante.",
            "**Antes de desplegar** debe resolverse R-01 (sección 10); R-09 y R-10 ya están resueltos en la imagen del pipeline.",
        ],
    )


def build_risks(doc) -> None:
    h1(doc, "10. Riesgos y limitaciones")
    add_table(
        doc,
        ["ID", "Situación", "Riesgo", "Tratamiento propuesto"],
        [
            ("R-01", "Al arrancar, la API lee también RAW de Urgencias y de población (~12 GB), no solo outputs procesados.",
             "Desplegarla exige llevar el RAW completo a la nube: más almacenamiento y arranques lentos.",
             "Que la API consuma solo outputs procesados (~200 MB) y los baje desde S3 al arrancar."),
            ("R-02", "Si el pronóstico persistido no está vigente, la API lo recalcula al arrancar.",
             "Hoy toma ~15 min con 12 núcleos; con 0,5 vCPU serían horas sin pronóstico.",
             "La tarea pipeline semanal regenera el pronóstico antes de publicar (sección 3.3)."),
            ("R-03", "No hay infraestructura como código ni publicación automática.",
             "Crear los recursos a mano no es reproducible ni fácil de destruir.",
             "Describir la infraestructura como código antes del primer despliegue."),
            ("R-04", "Costos estimados con precios de lista, no medidos.",
             "El costo real puede diferir (tráfico, precios vigentes).",
             "Confirmar en AWS Pricing Calculator y revisar la facturación el primer mes."),
            ("R-05", "Costo fijo del balanceador y de las IPv4 (≈ US$33 al mes) aunque no haya tráfico.",
             "Pagar por un servicio sin uso entre demos.",
             "Destruir y recrear el entorno entre demos cuando exista infraestructura como código."),
            ("R-06", "us-east-1 está lejos de los usuarios en Chile.",
             "Latencia adicional en cada consulta a la API (no medida).",
             "Aceptable para un servicio de planificación; si se mide un problema, mover a São Paulo."),
            ("R-07", "Una sola tarea de la API, con estado en memoria.",
             "Un reinicio deja la API sin servicio mientras carga.",
             "Aceptado para el MVP (DT-07 del DAS); el frontend ya reintenta ante `503`."),
            ("R-08", "Sin presupuesto asignado (RO-02).",
             "El to-be no puede desplegarse sin financiamiento.",
             "El as-is cubre la demo; el to-be se activa solo si se aprueba su costo."),
            ("R-09", "El pipeline asume `data/` en disco local.",
             "La tarea programada no puede leer ni escribir S3 directamente.",
             "Implementado: la imagen `saad-pipeline` (`src/cloud/pipeline_task.py`) sincroniza `data/` "
             "y `reports/` con S3 antes y después, sin cambiar el pipeline; si un paso falla no sube nada."),
            ("R-10", "La API toma datos nuevos solo al reiniciarse.",
             "Tras la actualización semanal seguiría sirviendo los datos anteriores.",
             "Implementado: al terminar con cambios, la tarea fuerza un nuevo despliegue del servicio "
             "`saad-api` (variables `SAAD_ECS_CLUSTER` y `SAAD_ECS_SERVICE`)."),
            ("R-11", "La frecuencia del DEIS se infirió de dos snapshots, no de un calendario oficial.",
             "La ejecución del lunes podría no tomar la semana más reciente.",
             "Registrar `data_lag_days` durante unas semanas y ajustar el día de ejecución."),
        ],
        [Cm(1.5), Cm(5.4), Cm(5.2), Cm(5.6)],
        font_size=8.8,
        first_col_bold=True,
    )


def build_traceability(doc) -> None:
    h1(doc, "11. Trazabilidad con las tareas")
    add_table(
        doc,
        ["Criterio de aceptación", "Dónde se atiende"],
        [
            ("D2-16 · Diagrama actualizado o validado", "Figura 1; sección 2."),
            ("D2-16 · Justificación técnica redactada", "Secciones 2, 5, 6 y 7; ADR-10 y ADR-11."),
            ("D2-16 · Estimación de costos (revisión)", "Sección 4."),
            ("D2-16 · Relación con solución y demo explicada", "Sección 9."),
            ("D2-16 · Riesgos o limitaciones mencionados", "Sección 10."),
            ("D2-21 · Nodos principales definidos", "Figura 1 y sección 3.1."),
            ("D2-21 · Componentes desplegados representados",
             "Figura 1 (servicios por nodo) y figura 2 (contenedores e imágenes)."),
            ("D2-21 · Relación con infraestructura clara", "Figura 1 y secciones 3.1 y 3.2."),
            ("D2-21 · Coherencia con arquitectura propuesta", "Sección 2; ADR-10 y ADR-11."),
        ],
        [Cm(8.2), Cm(9.5)],
        font_size=9.5,
        first_col_bold=True,
    )
    h2(doc, "11.1 Mantenimiento de este documento")
    p(
        doc,
        "Se regenera con `python scripts/generate_arquitectura_to_be_docx.py`. Las figuras 1 y 2 se "
        "editan en draw.io (`despliegue_to_be_aws.drawio` y `contenedores_docker.drawio`, en "
        "`docs/07-diagramas/`) y se exportan a PNG con el mismo nombre desde draw.io (Archivo → "
        "Exportar como → PNG, escala 150 %) o con `drawio -x -f png -s 1.5 -o <archivo>.png "
        "<archivo>.drawio`. Los precios de la sección 4 deben "
        "revisarse en AWS Pricing Calculator antes de cada entrega.",
    )


def main() -> None:
    apply_theme(BRAND_THEME)
    doc = new_document("ATB-SAAD-2026 · Arquitectura to-be: despliegue en AWS · v1.0")
    build_cover(doc)
    build_control(doc)
    add_toc(doc)
    build_purpose(doc)
    build_as_is_to_be(doc)
    build_deployment(doc)
    build_costs(doc)
    build_why_aws(doc)
    build_single_cloud(doc)
    build_technologies(doc)
    build_decisions(doc)
    build_demo(doc)
    build_risks(doc)
    build_traceability(doc)
    ruta = save(doc, OUTPUT_PATH)
    print(f"Documento generado: {ruta.as_posix()}")


if __name__ == "__main__":
    main()
