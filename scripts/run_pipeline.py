"""
Orquestador Principal del Pipeline - Proyecto Capstone

Permite la ejecución parcial o total del pipeline de procesamiento de datos,
asegurando la idempotencia (salto de etapas ya procesadas).
"""

import argparse
import logging
import subprocess
import sys
from pathlib import Path

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
)
logger = logging.getLogger(__name__)

# Definición de las etapas del pipeline y sus outputs esperados para idempotencia
STAGES = {
    "download_censo": {
        "module": "src.data.download_censo",
        "outputs": [
            "data/raw/censo/Cartografia_censo2024_Pais_Comunal.parquet",
            "data/raw/censo/Cartografia_censo2024_Pais_Distrital.parquet",
            "data/raw/censo/Cartografia_censo2024_Pais_Zonal.parquet",
            "data/raw/censo/Cartografia_censo2024_Pais_Entidades.parquet",
            "data/raw/censo/Cartografia_censo2024_Pais_Manzanas.parquet",
            "data/raw/censo/Diccionario_variables_geograficas_CPV24.xlsx"
        ],
        "depends_on": []
    },
    "download_censo_poblacion": {
        "module": "src.data.download_censo_poblacion",
        "outputs": [
            "data/raw/censo/D1_Poblacion-censada-por-sexo-y-edad-en-grupos-quinquenales.xlsx"
        ],
        "depends_on": []
    },
    "download_poblacion_proyecciones": {
        "module": "src.data.download_poblacion_proyecciones",
        "outputs": [
            "data/raw/poblacion/ine_estimaciones_proyecciones_2002_2035_comunas.xlsx"
        ],
        "depends_on": []
    },
    "download_pobreza_comunal": {
        "module": "src.data.download_pobreza_comunal",
        "outputs": [
            "data/raw/pobreza/mds_casen_2022_tasa_pobreza_ingresos_comunal.xlsx"
        ],
        "depends_on": []
    },
    "download_deis": {
        "module": "src.data.download_deis_sources",
        "outputs": [
            "data/raw/urgencias/AtencionesUrgencia2020.csv",
            "data/raw/urgencias/AtencionesUrgencia2021.csv",
            "data/raw/urgencias/AtencionesUrgencia2022.csv",
            "data/raw/urgencias/AtencionesUrgencia2023.csv",
            "data/raw/urgencias/AtencionesUrgencia2024.csv",
            "data/raw/urgencias/AtencionesUrgencia2025.csv",
            "data/raw/urgencias/AtencionesUrgencia2026.csv",
            "data/raw/egresos/egresos_2020.csv",
            "data/raw/egresos/egresos_2021.csv",
            "data/raw/egresos/egresos_2022.csv",
            "data/raw/egresos/egresos_2023.csv",
            "data/raw/egresos/egresos_2024.csv",
            "data/raw/egresos/egresos_2025.csv"
        ],
        "depends_on": []
    },
    "download_establishments": {
        "module": "src.data.download_establishments",
        "outputs": ["data/raw/deis/establecimientos_salud_actualizado.csv"],
        "depends_on": []
    },
    "download_contexto_genero": {
        "module": "src.data.download_gender_statistics",
        "outputs": [
            "data/raw/contexto_genero/egresos_intento_suicida_sexo_anio.xlsx",
            "data/raw/contexto_genero/suicidio_ratio_hm_tasas_nacional_regional.xlsx",
            "data/raw/contexto_genero/ansiedad_depresion_sintomas_18_mas_sexo.xlsx",
            "data/raw/contexto_genero/prevalencia_sintomas_depresivos_sexo.xlsx",
        ],
        "depends_on": []
    },
    "normalize_contexto_genero": {
        "module": "src.data.normalize_gender_statistics",
        "outputs": [
            "data/processed/contexto_genero/egresos_intento_suicida_sexo_anio.parquet",
            "data/processed/contexto_genero/suicidio_ratio_hm_tasas_nacional_regional.parquet",
            "data/processed/contexto_genero/ansiedad_depresion_sintomas_18_mas_sexo.parquet",
            "data/processed/contexto_genero/prevalencia_sintomas_depresivos_sexo.parquet",
        ],
        "depends_on": ["download_contexto_genero"]
    },
    "build_mart_contexto_genero_indicador_sexo": {
        "module": "src.data.build_contexto_genero_mart",
        "outputs": [
            "data/processed/marts/mart_contexto_genero_indicador_sexo.parquet",
        ],
        "depends_on": ["normalize_contexto_genero"]
    },
    "clean_establishments": {
        "module": "src.data.clean_establishments",
        "outputs": [
            "data/processed/establecimientos_rm_clean.csv",
            "data/processed/establecimientos_rm_clean.parquet",
            "data/processed/establecimientos_salud_clean.parquet"
        ],
        "depends_on": ["download_establishments"]
    },
    "clean_censo": {
        "module": "src.data.clean_censo_comunas",
        "outputs": [
            "data/processed/censo/Cartografia_censo2024_RM_Comunal.parquet",
            "data/processed/censo/Cartografia_censo2024_RM_Distrital.parquet",
            "data/processed/censo/Cartografia_censo2024_RM_Zonal.parquet",
            "data/processed/censo/Cartografia_censo2024_RM_Entidades.parquet",
            "data/processed/censo/Cartografia_censo2024_RM_Manzanas.parquet"
        ],
        "depends_on": ["download_censo"]
    },
    "clean_censo_poblacion": {
        "module": "src.data.clean_censo_poblacion",
        "outputs": ["data/processed/censo/dim_poblacion_comuna_censo2024.parquet"],
        "depends_on": ["download_censo_poblacion"]
    },
    "clean_poblacion_proyecciones": {
        "module": "src.data.clean_poblacion_proyecciones",
        "outputs": ["data/processed/censo/dim_poblacion_comuna_anual.parquet"],
        "depends_on": ["download_poblacion_proyecciones", "clean_censo_poblacion"]
    },
    "clean_pobreza_comunal": {
        "module": "src.data.clean_pobreza_comunal",
        "outputs": ["data/processed/pobreza/dim_vulnerabilidad_comuna.parquet"],
        "depends_on": ["download_pobreza_comunal", "clean_censo_poblacion"]
    },
    "build_catalogs": {
        "module": "src.data.build_catalogs",
        "outputs": ["data/processed/urgencias/catalogo_f00_f99.csv"],
        "depends_on": []
    },
    "clean_urgencias": {
        "module": "src.data.clean_urgencias",
        "outputs": [
            "data/processed/urgencias/urgencias_rm_2020.parquet",
            "data/processed/urgencias/urgencias_rm_2021.parquet",
            "data/processed/urgencias/urgencias_rm_2022.parquet",
            "data/processed/urgencias/urgencias_rm_2023.parquet",
            "data/processed/urgencias/urgencias_rm_2024.parquet",
            "data/processed/urgencias/urgencias_rm_2025.parquet",
            "data/processed/urgencias/urgencias_rm_2026.parquet"
        ],
        "depends_on": ["download_deis", "clean_establishments", "clean_censo", "build_catalogs"]
    },
    "clean_egresos": {
        "module": "src.data.clean_egresos",
        "outputs": [
            "data/processed/egresos/egresos_2020.parquet",
            "data/processed/egresos/egresos_2021.parquet",
            "data/processed/egresos/egresos_2022.parquet",
            "data/processed/egresos/egresos_2023.parquet",
            "data/processed/egresos/egresos_2024.parquet",
            "data/processed/egresos/egresos_2025.parquet"
        ],
        "depends_on": ["download_deis"]
    },
    "build_egresos_f00_f99": {
        "module": "src.data.build_egresos_f00_f99",
        "outputs": [
            "data/processed/egresos/egresos_f00_f99_nacional_2020_2025.parquet",
            "data/processed/egresos/catalogo_cie10_f00_f99.csv",
        ],
        "depends_on": ["clean_egresos"]
    },
    "eda_establishments": {
        "module": "scripts.eda_establecimientos_rm",
        "outputs": ["reports/eda/registros_sin_coordenadas.csv"],
        "depends_on": ["clean_establishments"]
    },
    "eda_urgencias": {
        "module": "scripts.eda_demanda_urgencias_rm",
        "outputs": ["data/processed/urgencias/tabla1_demanda_anual_rm.csv"],
        "depends_on": ["clean_urgencias"]
    },
    "profile_urgencias_sm_coverage": {
        "module": "src.data.profile_urgencias_sm_coverage",
        "outputs": [
            "data/processed/urgencias/"
            "perfil_cobertura_sm_comuna_semanal_2021_2025.parquet"
        ],
        "depends_on": ["clean_urgencias"]
    },
    "build_urgencias_comuna_marts": {
        "module": "src.data.build_urgencias_comuna_marts",
        "outputs": [
            "data/processed/marts/mart_urgencias_comuna_weekly.parquet",
            "data/processed/marts/mart_urgencias_comuna_monthly.parquet",
        ],
        "depends_on": ["clean_urgencias", "clean_poblacion_proyecciones"]
    },
    "build_urgencias_establecimiento_mart": {
        "module": "src.data.build_urgencias_establecimiento_mart",
        "outputs": [
            "data/processed/marts/mart_urgencias_establecimiento_monthly.parquet",
        ],
        "depends_on": ["clean_urgencias", "build_urgencias_comuna_marts"]
    },
    "build_urgencias_comuna_etario_mart": {
        "module": "src.data.build_urgencias_comuna_etario_mart",
        "outputs": [
            "data/processed/marts/mart_urgencias_comuna_etario_monthly.parquet",
        ],
        "depends_on": ["clean_urgencias", "build_urgencias_comuna_marts"]
    },
    "build_dim_oferta_urgencia_rm": {
        "module": "src.data.build_dim_oferta_urgencia_rm",
        "outputs": ["data/processed/geo/dim_oferta_urgencia_rm.parquet"],
        "depends_on": ["clean_establishments", "clean_urgencias"]
    },
    "build_mart_mvp_territorial_comuna": {
        "module": "src.data.build_mart_mvp_territorial_comuna",
        "outputs": ["data/processed/marts/mart_mvp_territorial_comuna.parquet"],
        "depends_on": [
            "clean_censo_poblacion", "clean_poblacion_proyecciones", "clean_pobreza_comunal",
            "clean_urgencias", "build_urgencias_comuna_marts", "build_dim_oferta_urgencia_rm",
        ]
    },
    "eda_contexto_genero": {
        "module": "scripts.eda_contexto_genero",
        "outputs": ["reports/eda/eda_contexto_genero_estadisticas_genero.md"],
        "depends_on": ["normalize_contexto_genero"]
    },
    "eda_mart_contexto_genero_indicador_sexo": {
        "module": "scripts.eda_mart_contexto_genero_indicador_sexo",
        "outputs": ["reports/eda/eda_mart_contexto_genero_indicador_sexo.md"],
        "depends_on": ["build_mart_contexto_genero_indicador_sexo"]
    }
}

# Orden estricto de ejecución
PIPELINE_ORDER = [
    "download_censo",
    "download_censo_poblacion",
    "download_poblacion_proyecciones",
    "download_pobreza_comunal",
    "download_deis",
    "download_establishments",
    "download_contexto_genero",
    "normalize_contexto_genero",
    "build_mart_contexto_genero_indicador_sexo",
    "clean_establishments",
    "clean_censo",
    "clean_censo_poblacion",
    "clean_poblacion_proyecciones",
    "clean_pobreza_comunal",
    "build_catalogs",
    "clean_urgencias",
    "clean_egresos",
    "build_egresos_f00_f99",
    "eda_establishments",
    "eda_urgencias",
    "profile_urgencias_sm_coverage",
    "build_urgencias_comuna_marts",
    "build_urgencias_establecimiento_mart",
    "build_urgencias_comuna_etario_mart",
    "build_dim_oferta_urgencia_rm",
    "build_mart_mvp_territorial_comuna",
    "eda_contexto_genero",
    "eda_mart_contexto_genero_indicador_sexo"
]

SUPPORTED_URGENCIAS_YEARS = tuple(range(2020, 2027))


def check_outputs_exist(outputs: list[str]) -> bool:
    """Verifica que todos los archivos de salida existan, tengan tamaño > 0 y sean legibles."""
    if not outputs:
        return False
    
    for out in outputs:
        p = Path(out)
        if not p.exists() or p.stat().st_size == 0:
            return False
            
        # Validación mínima según tipo
        try:
            ext = p.suffix.lower()
            if ext == '.csv':
                import pandas as pd
                df = pd.read_csv(p, nrows=0, encoding="latin-1")
                if len(df.columns) == 0:
                    return False
            elif ext == '.parquet':
                import pyarrow.parquet as pq
                schema = pq.read_schema(p)
                if len(schema.names) == 0:
                    return False
            elif ext == '.xlsx':
                from openpyxl import load_workbook
                workbook = load_workbook(p, read_only=True, data_only=False)
                try:
                    if not workbook.sheetnames:
                        return False
                finally:
                    workbook.close()
            elif ext == '.md':
                if not p.read_text(encoding="utf-8").strip():
                    return False
        except Exception as e:
            logger.warning(f"Archivo corrupto o ilegible {p}: {e}")
            return False

    return True


def stage_outputs(stage_name: str, year: int | None = None) -> list[str]:
    """Devuelve outputs de la etapa, acotados al año si corresponde."""
    if year is None:
        return STAGES[stage_name]["outputs"]
    if stage_name != "clean_urgencias":
        raise ValueError("--year solo está disponible para clean_urgencias.")
    return [f"data/processed/urgencias/urgencias_rm_{year}.parquet"]


def run_stage(
    stage_name: str,
    force: bool = False,
    upstream_changed: bool = False,
    year: int | None = None,
) -> bool:
    """Ejecuta una etapa específica. Devuelve True si se ejecutó, False si hizo SKIP."""
    if stage_name not in STAGES:
        logger.error(f"Stage desconocido: {stage_name}")
        sys.exit(1)
        
    config = STAGES[stage_name]
    module = config["module"]
    outputs = stage_outputs(stage_name, year)
    
    logger.info(f"--- Evaluando etapa: {stage_name} ---")
    
    if upstream_changed:
        logger.info(f"[FORZADO] Dependencias upstream fueron modificadas. Se ejecutará '{stage_name}'.")
    
    if not force and not upstream_changed and check_outputs_exist(outputs):
        logger.info(f"[SKIP] Etapa '{stage_name}' omitida. Los outputs ya existen y son válidos.")
        return False
        
    logger.info(f"[EJECUTANDO] Etapa '{stage_name}' -> Módulo: {module}")
    try:
        cmd = [sys.executable, "-m", module]
        if year is not None:
            cmd.extend(["--year", str(year)])
        if force:
            cmd.append("--force")
        subprocess.run(cmd, check=True)
        logger.info(f"[EXITO] Etapa '{stage_name}' finalizada correctamente.\n")
        return True
    except subprocess.CalledProcessError as e:
        logger.error(f"[ERROR] La etapa '{stage_name}' falló con código {e.returncode}.")
        sys.exit(1)


def main():
    parser = argparse.ArgumentParser(description="Orquestador del Pipeline de Datos")
    parser.add_argument(
        "--stage", 
        type=str, 
        choices=PIPELINE_ORDER + ["all"],
        default="all",
        help="Especifica la etapa a ejecutar (o 'all' para todo el pipeline)."
    )
    parser.add_argument(
        "--force", 
        action="store_true", 
        help="Fuerza la ejecución ignorando si los archivos de salida ya existen."
    )
    parser.add_argument(
        "--year", type=int,
        help="Año individual, disponible únicamente con --stage clean_urgencias.",
    )
    args = parser.parse_args()

    if args.year is not None:
        if args.stage != "clean_urgencias":
            parser.error("--year solo puede usarse con --stage clean_urgencias.")
        if args.year not in SUPPORTED_URGENCIAS_YEARS:
            parser.error(f"Año no soportado para clean_urgencias: {args.year}.")
        raw_path = Path(f"data/raw/urgencias/AtencionesUrgencia{args.year}.csv")
        if not raw_path.is_file() or raw_path.stat().st_size == 0:
            parser.error(f"RAW no disponible o vacío: {raw_path.as_posix()}")

    logger.info("Iniciando ejecución del pipeline...")
    
    # Determinar qué etapas evaluar según el --stage indicado
    target_stages = set(PIPELINE_ORDER) if args.stage == "all" else {args.stage}
    if args.stage != "all" and args.year is None:
        # Propagar recursivamente dependencias downstream (transitive closure)
        added = True
        while added:
            added = False
            for stage in PIPELINE_ORDER:
                if stage not in target_stages:
                    deps = STAGES[stage].get("depends_on", [])
                    if any(d in target_stages for d in deps):
                        target_stages.add(stage)
                        added = True
                        
    executed_stages = set()
    for stage in PIPELINE_ORDER:
        if stage in target_stages:
            deps = STAGES[stage].get("depends_on", [])
            upstream_changed = any(d in executed_stages for d in deps)
            
            # Solo aplicar args.force a la etapa explícitamente solicitada o a todas si es "all"
            is_explicit_target = (args.stage == "all" or stage == args.stage)
            stage_force = args.force if is_explicit_target else False
            
            executed = run_stage(
                stage,
                force=stage_force,
                upstream_changed=upstream_changed,
                year=args.year if stage == "clean_urgencias" else None,
            )
            if executed:
                executed_stages.add(stage)
                
    logger.info("Pipeline completado satisfactoriamente.")


if __name__ == "__main__":
    main()
