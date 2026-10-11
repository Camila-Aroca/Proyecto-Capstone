"""Genera el Documento de Arquitectura de Software (DAS) del proyecto SAAD en .docx.

Marco de referencia: ISO/IEC/IEEE 42010:2011 (descripción de arquitectura) y el modelo de
vistas 4+1 de Kruchten. Usa la identidad visual del kit de marca SAAD.

El documento describe la arquitectura **vigente**, verificada contra el código, y la separa
de la arquitectura **objetivo**; toda brecha entre ambas se declara de forma explícita.

Ejecución: python scripts/generate_das_docx.py
Salida:    docs/08-arquitectura/Documento_Arquitectura_Software.docx
           docs/07-diagramas/arquitectura_{capas,ejecucion,despliegue}.png
"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

from docx.shared import Cm
from PIL import Image

from arquitectura_figuras import generar_todas
from docx_builder import (
    BRAND_MINT_SOFT,
    BRAND_SKY_SOFT,
    BRAND_THEME,
    add_brand_cover,
    add_figure,
    add_picture,
    add_table,
    add_toc,
    apply_theme,
    bullets,
    h1,
    h2,
    h3,
    new_document,
    note,
    numbered,
    p,
    save,
)

OUTPUT_PATH = Path("docs/08-arquitectura/Documento_Arquitectura_Software.docx")
FIGURES_DIR = Path("docs/07-diagramas")
CLASS_DIAGRAM = FIGURES_DIR / "clases_saad.png"
LOGICAL_MODEL = FIGURES_DIR / "15-modelo-base-datos" / "modelo_datos_logico.png"

ANCHO_TOTAL = 17.7
# El modelo lógico es cuatro veces más ancho que alto: entero en una página queda
# ilegible. Se muestra en mitades, con una franja repetida para no cortar entidades.
SOLAPE_MODELO = 0.04


def mitades(ruta: Path, solape: float = SOLAPE_MODELO) -> list[BytesIO]:
    """Divide una imagen apaisada en mitad izquierda y derecha, como flujos PNG."""
    with Image.open(ruta) as imagen:
        ancho, alto = imagen.size
        corte, extra = ancho // 2, int(ancho * solape)
        partes = []
        for caja in ((0, 0, corte + extra, alto), (corte - extra, 0, ancho, alto)):
            flujo = BytesIO()
            imagen.crop(caja).save(flujo, format="PNG")
            flujo.seek(0)
            partes.append(flujo)
    return partes


def build_cover(doc) -> None:
    add_brand_cover(
        doc,
        title="Documento de Arquitectura de Software",
        subtitle=(
            "SAAD — Sistema de Análisis y Anticipación de Demanda de Urgencia "
            "en Salud Mental para la Región Metropolitana"
        ),
        project="Capstone APT · Ingeniería en Informática · Duoc UC Antonio Varas",
        kicker="Arquitectura de software",
        meta_rows=[
            ("Código del documento", "DAS-SAAD-2026"),
            ("Versión", "1.0"),
            ("Fecha de emisión", "11 de octubre de 2026"),
            ("Estado", "Emitido para revisión de la docente y del equipo"),
            ("Clasificación", "Uso académico e interno del equipo de proyecto"),
            (
                "Marco de referencia",
                "ISO/IEC/IEEE 42010:2011 (descripción de arquitectura); modelo de vistas 4+1; "
                "ISO/IEC 25010 para los atributos de calidad",
            ),
            ("Elaborado por", "Equipo SAAD — Camila A., Felipe R., Cristopher R., Catalina"),
            (
                "Documentos relacionados",
                "Especificación de Requisitos Funcionales (ERF-SAAD-2026) y No Funcionales "
                "(ERNF-SAAD-2026)",
            ),
            ("Repositorio", "docs/08-arquitectura/"),
        ],
    )


def build_control(doc) -> None:
    h1(doc, "Control del documento")

    h2(doc, "Historial de versiones")
    add_table(
        doc,
        ["Versión", "Fecha", "Autor(es)", "Descripción del cambio"],
        [
            ("1.0", "11-10-2026", "Equipo SAAD",
             "Primera versión formal: arquitectura seleccionada, vistas lógica, de datos, de "
             "ejecución, de desarrollo y de despliegue, integraciones y API, tecnologías "
             "justificadas, decisiones arquitectónicas, respuesta a los atributos de calidad, "
             "restricciones y deuda técnica. Verificada contra el código del repositorio."),
        ],
        [Cm(2.0), Cm(2.6), Cm(3.6), Cm(9.5)],
        font_size=9.5,
    )

    h2(doc, "Revisión y aprobación")
    add_table(
        doc,
        ["Rol", "Nombre", "Estado", "Fecha"],
        [
            ("Data Scientist (autor)", "Cristopher R.", "Emitido para revisión", "11-10-2026"),
            ("Jefa de Proyecto", "Camila A.", "Pendiente de revisión", "—"),
            ("Data Engineer", "Felipe R.", "Pendiente de revisión", "—"),
            ("Full Stack / Geoespacial", "Catalina", "Pendiente de revisión", "—"),
            ("Docente guía", "Por confirmar", "Pendiente de validación", "—"),
        ],
        [Cm(4.5), Cm(5.0), Cm(5.2), Cm(3.0)],
        font_size=9.5,
    )

    h2(doc, "Cómo leer este documento")
    p(
        doc,
        "El documento distingue en todo momento tres estados, que también usan las figuras: "
        "**implementado** (existe en el repositorio y fue verificado contra el código), "
        "**planificado** (especificado, aún no construido) y **externo** (sistemas o fuentes "
        "que el proyecto consume pero no controla). Cuando la arquitectura vigente difiere de "
        "la arquitectura objetivo declarada en los requisitos, la diferencia se expone como "
        "brecha en la sección 12 en lugar de presentarse como resuelta.",
    )
    note(
        doc,
        "**Criterio de verdad.** Ante una contradicción entre este documento y el código, "
        "prevalece el código: la arquitectura descrita aquí se obtuvo leyendo `src/`, "
        "`scripts/`, `frontend/`, `db/` y `PIPELINE.md` en la fecha de emisión. Las cifras "
        "(cantidad de etapas, módulos y pruebas) corresponden a esa fecha.",
        fill=BRAND_MINT_SOFT,
    )


def build_introduction(doc) -> None:
    h1(doc, "1. Introducción")

    h2(doc, "1.1 Propósito")
    p(
        doc,
        "Este documento describe la arquitectura de software del SAAD: qué estilo se eligió y "
        "por qué, qué componentes lo forman, cómo se comunican, sobre qué tecnologías se "
        "construye, qué decisiones la condicionan y qué restricciones debe respetar. Sirve "
        "como referencia común para el equipo y como evidencia académica de que la solución "
        "responde a los requisitos funcionales y no funcionales del proyecto.",
    )

    h2(doc, "1.2 Alcance")
    p(
        doc,
        "Cubre el sistema completo del MVP: la ingesta y preparación reproducible de datos "
        "oficiales, el modelo de pronóstico de demanda, la interfaz de programación que sirve "
        "los resultados y la aplicación web que los presenta. Los componentes del MVP que "
        "todavía no están construidos (accesibilidad territorial y análisis de hospitalización) "
        "se describen como planificados. Quedan fuera las funcionalidades excluidas del MVP en "
        "la sección 1.3 de la ERF, en particular la gestión de pacientes, la gestión de cupos o "
        "camas en tiempo real y el apoyo a decisiones clínicas (FA-01 a FA-08).",
    )

    h2(doc, "1.3 Audiencia")
    add_table(
        doc,
        ["Lector", "Qué necesita de este documento"],
        [
            ("Equipo de desarrollo",
             "Dónde vive cada responsabilidad, qué contratos respetar y qué decisiones ya están tomadas."),
            ("Docente guía y comisión evaluadora",
             "Evidencia de que la arquitectura es coherente con el sistema real y con los requisitos."),
            ("Cliente piloto o contraparte técnica",
             "Qué consume el sistema, qué expone y bajo qué limitaciones debe interpretarse."),
            ("Quien mantenga el sistema a futuro",
             "Las razones de cada decisión, para poder revisarlas con fundamento."),
        ],
        [Cm(5.2), Cm(12.5)],
        font_size=9.5,
        first_col_bold=True,
    )

    h2(doc, "1.4 Marco de descripción")
    p(
        doc,
        "La descripción sigue ISO/IEC/IEEE 42010: identifica a los interesados y sus "
        "preocupaciones, y responde a ellas mediante vistas. Las vistas se organizan según el "
        "modelo 4+1 de Kruchten, adaptado a un sistema centrado en datos:",
    )
    add_table(
        doc,
        ["Vista", "Pregunta que responde", "Sección"],
        [
            ("Lógica", "¿Qué componentes existen y qué responsabilidad tiene cada uno?", "4"),
            ("De datos", "¿Qué datos se almacenan, en qué zonas y con qué garantías?", "5"),
            ("De ejecución (procesos)", "¿Qué ocurre al correr el pipeline y al atender una consulta?", "6"),
            ("De desarrollo", "¿Cómo se organiza el código y cómo se verifica?", "7"),
            ("De despliegue (física)", "¿Dónde se ejecuta cada componente?", "8"),
            ("Escenarios (+1)", "¿Qué casos de uso recorren la arquitectura?", "ERF, sección 5"),
        ],
        [Cm(4.4), Cm(10.4), Cm(2.9)],
        font_size=9.5,
        first_col_bold=True,
    )

    h2(doc, "1.5 Definiciones y abreviaturas")
    add_table(
        doc,
        ["Término", "Definición"],
        [
            ("API", "Interfaz de programación de aplicaciones; aquí, el servicio HTTP que entrega los resultados."),
            ("DAG", "Grafo dirigido acíclico: el orden y las dependencias entre las etapas del pipeline."),
            ("Etapa (stage)", "Unidad del pipeline con entradas, salidas y validación declaradas."),
            ("Mart", "Tabla analítica ya agregada al grano que consume un análisis o la plataforma."),
            ("Pipeline", "Proceso por lotes que transforma las fuentes oficiales en datos y resultados publicables."),
            ("Procedencia", "Registro de dónde, cuándo y con qué huella (SHA256) se obtuvo un archivo fuente."),
            ("RAW", "Copia fiel de un archivo fuente, sin modificar (`data/raw/`)."),
            ("SPA", "Aplicación web de una sola página, que se ejecuta en el navegador."),
            ("Snapshot", "Estado de una fuente en el momento de su descarga."),
        ],
        [Cm(3.6), Cm(14.1)],
        font_size=9.5,
        first_col_bold=True,
    )

    h2(doc, "1.6 Referencias")
    add_table(
        doc,
        ["Código", "Documento"],
        [
            ("DEF-01", "Documento de Definición de Proyecto APT — Fase 1 (`docs/01-definicion-proyecto/`)."),
            ("ERF-SAAD-2026", "Especificación de Requisitos Funcionales, v1.1 (`docs/06-requisitos/`)."),
            ("ERNF-SAAD-2026", "Especificación de Requisitos No Funcionales, v1.0 (`docs/06-requisitos/`)."),
            ("PIPELINE.md", "Fuente de verdad del DAG reproducible: etapas, entradas, salidas y dependencias."),
            ("PROJECT_CONTEXT.md", "Estado vigente y decisiones estables del repositorio."),
            ("REG-01", "Reglas de rigor analítico, portabilidad y reproducibilidad (`.agents/rules.md`)."),
            ("Diagramas", "Casos de uso, clases y modelo de datos lógico y físico (`docs/07-diagramas/`)."),
            ("DDL", "Diseño físico PostgreSQL/PostGIS (`db/schema_postgresql.sql`)."),
            ("ISO/IEC/IEEE 42010:2011", "Systems and software engineering — Architecture description."),
            ("Kruchten (1995)", "The 4+1 View Model of Architecture. IEEE Software 12(6)."),
        ],
        [Cm(4.4), Cm(13.3)],
        font_size=9.5,
        first_col_bold=True,
    )


def build_drivers(doc) -> None:
    h1(doc, "2. Contexto y fuerzas que dan forma a la arquitectura")

    h2(doc, "2.1 Qué debe lograr el sistema")
    p(
        doc,
        "El SAAD apoya la planificación de la red de urgencia en salud mental de la Región "
        "Metropolitana. Para eso anticipa la demanda semanal a un horizonte de 4 a 8 semanas, "
        "la muestra por territorio junto a su contexto, y declara siempre su incertidumbre y "
        "sus límites de interpretación. No es un tablero descriptivo ni una herramienta "
        "clínica: su valor está en anticipar y en hacerlo de forma verificable.",
    )

    h2(doc, "2.2 Interesados y sus preocupaciones")
    add_table(
        doc,
        ["Interesado", "Preocupación principal", "Cómo responde la arquitectura"],
        [
            ("Gestor o planificador de la red",
             "Que las cifras sean correctas, comprensibles y lleguen a tiempo.",
             "Resultados precalculados, incertidumbre visible y advertencias de lectura servidas por la API."),
            ("Equipo de datos",
             "Poder reconstruir cualquier cifra y detectar cuándo una fuente cambió.",
             "Pipeline reproducible, RAW inmutable con procedencia y contratos de output."),
            ("Equipo de modelamiento",
             "Comparar modelos de forma honesta y servir exactamente lo evaluado.",
             "Interfaz única de modelo, registro de configuraciones y selección congelada."),
            ("Equipo de plataforma",
             "Construir la interfaz sin depender de los detalles del análisis.",
             "Contrato HTTP/JSON tipado como única vía entre presentación y servicios."),
            ("Docente y comisión",
             "Que lo entregado sea reproducible, probado y trazable a los requisitos.",
             "Documentos y figuras generados desde código; pruebas automatizadas por función analítica."),
        ],
        [Cm(3.9), Cm(5.9), Cm(7.9)],
        font_size=9.0,
        first_col_bold=True,
    )

    h2(doc, "2.3 Atributos de calidad que más pesan")
    p(
        doc,
        "De los requisitos no funcionales, estos son los que determinan decisiones "
        "estructurales. La sección 11 detalla cómo se atiende cada uno.",
    )
    add_table(
        doc,
        ["Atributo", "Requisito", "Efecto sobre la arquitectura"],
        [
            ("Reproducibilidad y trazabilidad", "RNF-DAT-007, RN-08",
             "Todo resultado sale de un DAG versionado; cada fuente conserva su huella."),
            ("Corrección de las cifras", "RNF-AFU-001, RNF-AFU-005",
             "Validaciones bloqueantes y reconciliación entre niveles antes de publicar."),
            ("Mantenibilidad", "RNF-MAN-001, RNF-MAN-004",
             "Capas con dependencia en un solo sentido y prueba por cada función analítica."),
            ("Reemplazabilidad del modelo", "RNF-FLE-005",
             "Los modelos comparten una interfaz mínima y se instancian desde un registro."),
            ("Desempeño de consulta", "RNF-EDE-001, RNF-EDE-002",
             "La API no calcula al atender: lee al arrancar y responde desde memoria."),
            ("Degradación controlada", "RNF-FIA-002",
             "Un dato que aún no está listo se informa como tal; nunca se reemplaza por cero."),
            ("Uso responsable", "RNF-USO-001, RNF-USO-004",
             "La incertidumbre y las advertencias viajan con los datos, no solo en la interfaz."),
        ],
        [Cm(4.4), Cm(3.6), Cm(9.7)],
        font_size=9.0,
        first_col_bold=True,
    )

    h2(doc, "2.4 Restricciones que la arquitectura no puede relajar")
    bullets(
        doc,
        [
            "**Datos agregados y disociados.** El sistema no procesa datos personales (RNF-SEG-001, FA-03).",
            "**Relación ecológica entre fuentes.** Urgencias y Egresos nunca se vinculan por paciente "
            "ni por establecimiento; mantienen granularidades separadas (RN-05, RD-02, RD-03).",
            "**Consumo por lotes.** Las fuentes se descargan periódicamente; no hay integración en "
            "tiempo real con sistemas hospitalarios (FA-01, FA-08).",
            "**Sin presupuesto de infraestructura** (RO-02) y un equipo de cuatro personas con "
            "dedicación parcial (RO-01): la solución debe poder ejecutarse en un computador personal.",
            "**Stack definido** en la definición del proyecto: Python con FastAPI, React, PostgreSQL "
            "con PostGIS y Docker Compose (RT-01 a RT-04).",
        ],
    )


def build_selected_architecture(doc, figuras: dict[str, Path]) -> None:
    h1(doc, "3. Arquitectura seleccionada")

    h2(doc, "3.1 Estilo")
    p(
        doc,
        "El SAAD adopta una **arquitectura por capas con un pipeline de datos por lotes y una "
        "capa de servicio de solo lectura**. La figura 1 muestra sus capas; a continuación se "
        "explican las tres partes que las agrupan.",
    )
    add_figure(
        doc,
        figuras["capas"],
        "**Figura 1.** Vista por capas del SAAD. Las flechas indican el sentido en que fluyen los "
        "datos; cada capa solo depende de las que tiene encima.",
        width_cm=15.4,
    )
    p(
        doc,
        "Son tres partes que se ejecutan en momentos distintos y se comunican únicamente a través "
        "de artefactos bien definidos:",
    )
    numbered(
        doc,
        [
            "Un **pipeline reproducible** transforma las fuentes oficiales en datos limpios, marts "
            "analíticos, la evaluación del modelo y el pronóstico vigente. Lo ejecuta el equipo, no "
            "el usuario final.",
            "Una **API REST de solo lectura** carga esos artefactos al arrancar y los entrega como "
            "JSON. No escribe datos ni entrena modelos al atender una consulta.",
            "Una **aplicación web** consume la API y presenta la serie, el mapa comunal y la tabla, "
            "con sus advertencias de interpretación.",
        ],
    )
    p(
        doc,
        "La dependencia va en un solo sentido: la presentación conoce a la API, la API conoce a "
        "los artefactos del pipeline, y el pipeline no conoce a ninguna de las dos. Esa asimetría "
        "es la que permite cambiar la interfaz sin tocar el análisis y reemplazar el modelo sin "
        "tocar la interfaz.",
    )

    h2(doc, "3.2 Por qué este estilo")
    add_table(
        doc,
        ["Fuerza", "Cómo la resuelve el estilo elegido"],
        [
            ("Las fuentes se publican por lotes y cambian sin aviso",
             "El pipeline descarga, valida y registra cada snapshot antes de usarlo; lo ya publicado no se altera."),
            ("Las cifras deben poder reconstruirse",
             "Cada etapa declara entradas, salidas y validación; el mismo código sobre el mismo snapshot da el mismo resultado."),
            ("Ajustar el modelo toma segundos a minutos",
             "El cálculo ocurre en el pipeline; la consulta solo lee resultados ya calculados."),
            ("El equipo trabaja en paralelo por especialidad",
             "Cada capa tiene un responsable y un contrato de salida; el trabajo de una no bloquea a la otra."),
            ("No hay infraestructura pagada",
             "Todo corre en un computador personal: archivos locales, un proceso de API y un servidor de desarrollo."),
        ],
        [Cm(6.2), Cm(11.5)],
        font_size=9.5,
        first_col_bold=True,
    )

    h2(doc, "3.3 Alternativas consideradas y descartadas")
    add_table(
        doc,
        ["Alternativa", "Por qué no se eligió"],
        [
            ("Microservicios",
             "Agregan despliegue, red y observabilidad distribuida para un sistema con un solo flujo de "
             "datos, cuatro integrantes y sin infraestructura. No hay cargas que escalar por separado."),
            ("Procesamiento en tiempo real (streaming)",
             "Las fuentes oficiales se publican como archivos periódicos; no existe un flujo continuo que "
             "consumir, y la integración en línea está fuera del alcance (FA-08)."),
            ("Base de datos como almacenamiento primario desde el inicio",
             "Obliga a mantener un servidor y una carga de datos antes de saber qué tablas necesita el "
             "producto. Los archivos Parquet dan lo mismo para leer y son reproducibles desde Git. "
             "PostgreSQL/PostGIS queda como capa de servicio planificada (sección 5.4)."),
            ("Herramienta de inteligencia de negocios como producto",
             "Reduce el producto a un tablero descriptivo, lo que las reglas del proyecto excluyen, y no "
             "permite controlar cómo se comunican la incertidumbre y las advertencias."),
            ("Entrenar el modelo dentro de la API",
             "Haría depender el tiempo de respuesta del ajuste del modelo y permitiría servir una "
             "configuración distinta de la evaluada."),
        ],
        [Cm(5.2), Cm(12.5)],
        font_size=9.5,
        first_col_bold=True,
    )


def build_logical_view(doc) -> None:
    h1(doc, "4. Vista lógica: componentes")
    p(
        doc,
        "La tabla define los componentes principales, dónde viven en el repositorio y en qué "
        "estado se encuentran a la fecha de emisión.",
    )

    h2(doc, "4.1 Capa de datos")
    add_table(
        doc,
        ["Componente", "Responsabilidad", "Ubicación", "Estado"],
        [
            ("Orquestador del DAG",
             "Ejecuta las 32 etapas en orden, decide SKIP o regeneración y propaga cambios solo aguas abajo.",
             "`scripts/run_pipeline.py`", "Implementado"),
            ("Ingesta y procedencia",
             "Descarga a una ubicación temporal, valida, registra URL, fecha y SHA256, y recién entonces publica el RAW.",
             "`src/data/download_*`, `raw_provenance.py`", "Implementado"),
            ("Limpieza y normalización",
             "Tipado, codificación, filtro territorial y unión con el maestro de establecimientos.",
             "`src/data/clean_*`, `normalize_*`", "Implementado"),
            ("Dimensiones y catálogos",
             "Población comunal anual, vulnerabilidad, oferta de urgencia y catálogo F00–F99.",
             "`src/data/build_dim_*`, `build_catalogs`", "Implementado"),
            ("Marts analíticos",
             "Urgencias por comuna (semanal y mensual), por establecimiento y por grupo etario; mart territorial del MVP.",
             "`src/data/build_*_mart*`", "Implementado"),
            ("Contratos de output",
             "Comprueba esquema, tipos y vigencia de cada output conocido antes de aceptarlo como válido.",
             "`src/data/output_contracts.py`", "Implementado"),
        ],
        [Cm(3.4), Cm(7.3), Cm(4.6), Cm(2.4)],
        font_size=8.8,
        first_col_bold=True,
    )

    h2(doc, "4.2 Capa analítica")
    add_table(
        doc,
        ["Componente", "Responsabilidad", "Ubicación", "Estado"],
        [
            ("EDA y perfilado",
             "Informes reproducibles de cada fuente y de cada mart, con sus tablas de apoyo.",
             "`scripts/eda_*`, `reports/eda/`", "Implementado"),
            ("Modelos de demanda",
             "Panel jerárquico región–comunas, variables sin fuga temporal, baseline estacional, GLM de "
             "Poisson, LightGBM global, híbrido por nivel, reconciliación e intervalos conformes.",
             "`src/models/`", "Implementado"),
            ("Benchmark y holdout",
             "Backtesting con origen deslizante, comparación pareada, selección congelada del modelo y "
             "evaluación sobre el año en curso.",
             "`scripts/run_demand_benchmark.py`, `evaluate_holdout_demanda_sm.py`", "Implementado"),
            ("Persistencia del pronóstico",
             "Guarda el pronóstico vigente con un manifiesto de modelo, configuración y huella de sus entradas.",
             "`src/models/persistence.py`", "Implementado"),
            ("Accesibilidad territorial",
             "Tiempos de viaje e isócronas sobre red vial y transporte público; cobertura poblacional.",
             "`src/geo/` (no existe aún)", "Planificado"),
            ("Análisis de hospitalización",
             "Duración de estadía F00–F99 según sexo, previsión y pertenencia al SNSS. El dataset "
             "canónico y su perfilado ya existen; falta el análisis.",
             "`reports/hospitalization/`", "Planificado"),
        ],
        [Cm(3.4), Cm(7.3), Cm(4.6), Cm(2.4)],
        font_size=8.8,
        first_col_bold=True,
    )

    h2(doc, "4.3 Capas de servicios y de presentación")
    add_table(
        doc,
        ["Componente", "Responsabilidad", "Ubicación", "Estado"],
        [
            ("Aplicación y ciclo de vida",
             "Arranque, carga en segundo plano, compresión y política de orígenes permitidos.",
             "`src/api/main.py`", "Implementado"),
            ("Rutas",
             "Endpoints de la versión 1, validación de parámetros y errores HTTP explícitos.",
             "`src/api/routes.py`", "Implementado"),
            ("Contratos de respuesta",
             "Modelos tipados de cada respuesta; fuente de la documentación OpenAPI.",
             "`src/api/schemas.py`", "Implementado"),
            ("Servicio de demanda",
             "Estado en memoria: datos cargados, pronóstico reutilizado o recalculado, y funciones puras de resumen.",
             "`src/api/service.py`", "Implementado"),
            ("Cliente tipado",
             "Único punto de acceso a la API desde la interfaz; reintenta cuando el servidor informa que aún calcula.",
             "`frontend/src/api.ts`", "Implementado"),
            ("Portada y tablero",
             "Presentación del producto y vista de trabajo: serie con pronóstico, mapa comunal, tabla y alcance.",
             "`frontend/src/components/`", "Implementado"),
            ("Identidad visual",
             "Tokens de color y tipografía del kit de marca, aplicados a la interfaz y a este documento.",
             "`frontend/src/index.css`, `frontend/public/marca/`", "Implementado"),
        ],
        [Cm(3.4), Cm(7.3), Cm(4.6), Cm(2.4)],
        font_size=8.8,
        first_col_bold=True,
    )

    h2(doc, "4.4 Clases conceptuales")
    p(
        doc,
        "El diagrama de clases resume las entidades conceptuales del dominio y sus asociaciones. "
        "Se mantiene sincronizado con el modelo de datos por una regla del repositorio: cada "
        "cambio material del modelo canónico obliga a regenerarlo desde su fuente.",
    )
    add_figure(
        doc,
        CLASS_DIAGRAM,
        "**Figura 2.** Diagrama de clases del SAAD. Fuente editable: "
        "`docs/07-diagramas/clases_saad.puml`; la versión vectorial permite ampliar el detalle.",
        width_cm=9.4,
    )


def build_data_view(doc) -> None:
    h1(doc, "5. Vista de datos")

    h2(doc, "5.1 Zonas de almacenamiento")
    add_table(
        doc,
        ["Zona", "Contenido", "Garantía", "¿Se versiona en Git?"],
        [
            ("`data/raw/`",
             "Copia fiel de cada archivo fuente, en su formato original.",
             "Solo lectura para todo lo que no sea la ingesta. Los años cerrados no se sobrescriben.",
             "No; sí su manifiesto de procedencia."),
            ("`data/processed/`",
             "Datos limpios, dimensiones y marts en Parquet y GeoParquet.",
             "Regenerable por completo desde el RAW mediante el pipeline.",
             "No."),
            ("`data/processed/modeling/`",
             "Pronóstico vigente y su manifiesto.",
             "Válido solo mientras no cambien las entradas registradas en el manifiesto.",
             "Sí, por excepción explícita."),
            ("`reports/`",
             "Informes de EDA, benchmark, holdout y diccionarios de datos.",
             "Generados por etapas del pipeline; no se editan a mano.",
             "Sí, los de modelamiento."),
        ],
        [Cm(3.9), Cm(5.0), Cm(5.6), Cm(3.2)],
        font_size=9.0,
        first_col_bold=True,
    )

    h2(doc, "5.2 Garantías sobre los datos")
    bullets(
        doc,
        [
            "**Procedencia.** `data/raw/provenance_manifest.json` conserva el historial de snapshots "
            "de cada fuente: URL, fecha y hora real de descarga, tamaño y SHA256.",
            "**Publicación atómica.** Los outputs se escriben en un archivo temporal, se validan y se "
            "reemplazan de una vez; una ejecución interrumpida no deja un archivo a medias (RNF-FIA-003).",
            "**Contratos.** Un output solo cuenta como válido si existe, se puede leer y cumple su "
            "esquema; en el caso del pronóstico, además si sigue vigente.",
            "**Sin imputación silenciosa.** Una comuna o semana sin dato se informa como ausente; no se "
            "reemplaza por cero ni por un valor estimado (RN-04, RN-10).",
            "**Fuente mutable identificada.** El año en curso de Urgencias cambia entre descargas; cada "
            "Parquet anual declara en sus metadatos el SHA256 del RAW del que proviene.",
        ],
    )

    h2(doc, "5.3 Modelo de datos")
    p(
        doc,
        "El modelo lógico organiza los datos en dimensiones territoriales y temporales, marts de "
        "Urgencias, el dataset de Egresos F00–F99 y el pronóstico. Dos invariantes lo recorren: "
        "Urgencias se relaciona con los establecimientos por `establecimiento_codigo`, mientras que "
        "Egresos no identifica el establecimiento y su territorio es el de residencia; y no existe "
        "ninguna unión entre ambas fuentes a nivel de paciente.",
    )
    izquierda, derecha = mitades(LOGICAL_MODEL)
    add_picture(doc, izquierda, width_cm=17.4)
    add_figure(
        doc,
        derecha,
        "**Figura 3.** Modelo de datos lógico, en dos mitades (izquierda arriba, derecha abajo). "
        "Fuente editable: `docs/07-diagramas/15-modelo-base-datos/modelo_datos_logico.mmd`. "
        "Para leer el detalle de cada entidad se recomienda la versión vectorial (`.svg`).",
        width_cm=17.4,
    )

    h2(doc, "5.4 Persistencia relacional: diseñada, aún no en uso")
    p(
        doc,
        "El diseño físico en PostgreSQL con PostGIS está escrito y versionado en "
        "`db/schema_postgresql.sql`, con ocho esquemas (`censo`, `pobreza`, `geo`, `urgencias`, "
        "`marts`, `egresos`, `contexto_genero` y `modeling`) y se mantiene coherente con el modelo "
        "lógico. Sin embargo, **el sistema vigente no carga ni consulta esa base**: la API lee los "
        "archivos Parquet directamente. La restricción RT-03 queda, por tanto, parcialmente "
        "atendida; la brecha y su justificación están en las secciones 10 (ADR-02) y 12.",
    )


def build_process_view(doc, figuras: dict[str, Path]) -> None:
    h1(doc, "6. Vista de ejecución")
    p(
        doc,
        "El sistema tiene dos procesos independientes, que se encuentran solo en los archivos que "
        "uno produce y el otro lee.",
    )
    add_figure(
        doc,
        figuras["ejecucion"],
        "**Figura 4.** Los dos procesos del SAAD. El pronóstico persistido y los marts son el único "
        "punto de contacto entre el pipeline y la consulta en línea.",
        width_cm=16.8,
    )

    h2(doc, "6.1 Proceso por lotes: el pipeline")
    p(
        doc,
        "Se ejecuta con `python scripts/run_pipeline.py`, completo o por etapa. Para cada etapa el "
        "orquestador decide:",
    )
    bullets(
        doc,
        [
            "**SKIP** si sus outputs existen, cumplen su contrato y nada cambió aguas arriba.",
            "**Regenerar** si falta un output, si su contrato no se cumple o si una dependencia cambió.",
            "**Forzar** con `--force`, que regenera la etapa pedida y propaga el cambio solo a las "
            "etapas que dependen de ella.",
        ],
    )
    p(
        doc,
        "Una etapa se considera exitosa únicamente si el proceso termina sin error **y** sus outputs "
        "declarados existen y son válidos. La actualización del año en curso de Urgencias tiene un "
        "comando propio (`--refresh-current-urgencias`) que vuelve a consultar la fuente, compara "
        "huellas y regenera solo lo afectado.",
    )

    h2(doc, "6.2 Consulta en línea: la API y la aplicación")
    numbered(
        doc,
        [
            "Al arrancar, la API carga en memoria el panel observado, el mart territorial, la "
            "cartografía simplificada y la selección congelada del modelo.",
            "En segundo plano obtiene el pronóstico: reutiliza el persistido si sigue vigente; si falta "
            "o quedó desactualizado, lo recalcula en memoria y lo declara en `/meta`.",
            "Cada consulta filtra y serializa tablas pequeñas ya cargadas; no lee disco ni ajusta modelos.",
            "Mientras la carga o el cálculo no terminan, los endpoints afectados responden `503` con "
            "`Retry-After`; el cliente reintenta y la interfaz muestra un estado de espera.",
        ],
    )

    h2(doc, "6.3 Concurrencia y estado")
    p(
        doc,
        "La API es un único proceso con estado en memoria. El cálculo del pronóstico está protegido "
        "por un candado, de modo que dos solicitudes simultáneas no lo disparan dos veces. Como todas "
        "las operaciones son de lectura sobre datos inmutables durante la vida del proceso, no hay "
        "condiciones de carrera sobre los datos. La consecuencia es que **un snapshot nuevo solo se "
        "refleja al reiniciar la API**; es una limitación aceptada para el MVP y coherente con el uso "
        "acotado previsto (hasta 20 sesiones concurrentes, ERNF sección 3).",
    )


def build_development_view(doc) -> None:
    h1(doc, "7. Vista de desarrollo")

    h2(doc, "7.1 Organización del repositorio")
    add_table(
        doc,
        ["Ruta", "Contenido", "Regla"],
        [
            ("`src/data/`", "27 módulos de ingesta, limpieza, dimensiones, marts y contratos.",
             "Solo se ejecutan a través del pipeline para generar outputs oficiales."),
            ("`src/models/`", "14 módulos: variables, modelos, backtesting, jerarquía, intervalos, registro y persistencia.",
             "Sin años ni comunas fijos en el código; semilla explícita en todo lo estocástico."),
            ("`src/api/`", "5 módulos: aplicación, rutas, contratos y servicio.",
             "Solo lectura; consume outputs del pipeline, no los genera."),
            ("`scripts/`", "Orquestador, etapas de análisis y generadores de documentos.",
             "Todo script permanente debe estar registrado en el DAG."),
            ("`frontend/`", "Aplicación React con TypeScript, su configuración y los recursos de marca.",
             "Accede a los datos únicamente por `frontend/src/api.ts`."),
            ("`tests/`", "32 archivos, 415 pruebas automatizadas.",
             "Cada función de transformación o cálculo analítico tiene su prueba."),
            ("`db/`", "Diseño físico PostgreSQL/PostGIS.", "Coherente con el modelo lógico y el UML."),
            ("`docs/`, `reports/`", "Documentación formal, diagramas e informes reproducibles.",
             "Los informes generados no se editan a mano."),
        ],
        [Cm(3.2), Cm(7.6), Cm(6.9)],
        font_size=9.0,
        first_col_bold=True,
    )

    h2(doc, "7.2 Reglas de dependencia")
    bullets(
        doc,
        [
            "`src/api/` puede importar `src/models/` y `src/data/`; lo contrario está prohibido.",
            "`src/models/` no conoce la API ni la interfaz; solo recibe tablas y devuelve tablas.",
            "`frontend/` no importa nada de Python: su único contrato es HTTP/JSON.",
            "Ningún módulo escribe en `data/raw/` salvo las etapas de ingesta.",
        ],
    )

    h2(doc, "7.3 Verificación")
    add_table(
        doc,
        ["Mecanismo", "Qué cubre", "Estado"],
        [
            ("Pruebas unitarias y de contrato (pytest)",
             "Transformaciones, modelos, persistencia, contratos de output, orquestación y API.",
             "415 pruebas, ejecutadas localmente."),
            ("Verificación de tipos del frontend",
             "Compilación estricta de TypeScript y construcción de producción con Vite.",
             "Manual, antes de integrar."),
            ("Contratos de output en el pipeline",
             "Esquema y vigencia de los outputs antes de aceptar un SKIP.",
             "Automático en cada ejecución."),
            ("Integración continua",
             "Ejecución automática de las pruebas en cada cambio propuesto (RNF-MAN-005).",
             "Pendiente: no existe flujo de CI."),
        ],
        [Cm(5.2), Cm(8.1), Cm(4.4)],
        font_size=9.0,
        first_col_bold=True,
    )

    h2(doc, "7.4 Gestión de configuración")
    bullets(
        doc,
        [
            "Las dependencias de Python están fijadas por versión en `requirements.txt`; las del "
            "frontend, en `frontend/package.json` y su archivo de bloqueo.",
            "Los valores sensibles se reciben por variables de entorno y nunca se versionan: "
            "`frontend/.env.local` para la clave del mapa base, con `frontend/.env.example` como plantilla.",
            "La configuración no sensible es versionada: orígenes permitidos de la API "
            "(`SAAD_CORS_ORIGINS`) y destino del proxy de desarrollo (`SAAD_API_URL`).",
            "El trabajo se integra en la rama principal de GitHub y se gestiona en GitHub Projects (RT-05).",
        ],
    )


def build_deployment_view(doc, figuras: dict[str, Path]) -> None:
    h1(doc, "8. Vista de despliegue")
    p(
        doc,
        "A la fecha de emisión el sistema se ejecuta en modo de desarrollo local. El despliegue "
        "empaquetado que exigen los requisitos está planificado y se muestra a la derecha de la "
        "figura para dejar clara la distancia entre ambos.",
    )
    add_figure(
        doc,
        figuras["despliegue"],
        "**Figura 5.** Despliegue vigente (izquierda) y objetivo (derecha). El diagrama UML de "
        "despliegue formal es una tarea abierta del equipo (D2-21).",
        width_cm=16.8,
    )

    h2(doc, "8.1 Despliegue vigente")
    add_table(
        doc,
        ["Nodo o proceso", "Qué ejecuta", "Cómo se inicia"],
        [
            ("Servidor de la API", "Uvicorn con la aplicación FastAPI, puerto 8000.",
             "`uvicorn src.api.main:app --port 8000`"),
            ("Servidor de desarrollo", "Vite: sirve la aplicación y reenvía `/api` al backend, puerto 5173.",
             "`npm run dev` en `frontend/`"),
            ("Pipeline", "Proceso por lotes bajo demanda del equipo.",
             "`python scripts/run_pipeline.py`"),
            ("Almacenamiento", "Sistema de archivos local: `data/` y `reports/`.", "No requiere servicio."),
        ],
        [Cm(3.9), Cm(7.6), Cm(6.2)],
        font_size=9.0,
        first_col_bold=True,
    )
    p(
        doc,
        "En desarrollo el navegador ve un solo origen, porque el servidor de Vite reenvía las "
        "llamadas a la API; por eso no se requiere configuración de CORS para trabajar localmente.",
    )

    h2(doc, "8.2 Despliegue objetivo")
    p(
        doc,
        "El objetivo declarado en RT-04 es un empaquetado con Docker Compose: un contenedor que "
        "sirva los archivos estáticos de la aplicación y actúe como proxy inverso, uno para la API "
        "y uno para PostgreSQL con PostGIS, más el pipeline como tarea por lotes sobre un volumen de "
        "datos. **Nada de esto existe todavía en el repositorio**: no hay archivos de contenedor ni "
        "de composición. La arquitectura vigente lo facilita, porque la API y la aplicación ya son "
        "procesos separados que se comunican por HTTP y no comparten estado.",
    )


def build_integrations(doc) -> None:
    h1(doc, "9. Integraciones, servicios y API")

    h2(doc, "9.1 Fuentes externas")
    p(
        doc,
        "Todas las integraciones con fuentes de datos son descargas por lotes sobre HTTPS, "
        "iniciadas por el pipeline. Ninguna es una llamada en línea durante una consulta.",
    )
    add_table(
        doc,
        ["Fuente", "Qué aporta", "Formato", "Frecuencia de cambio"],
        [
            ("DEIS / MINSAL — Atenciones de Urgencia",
             "Atenciones por causa, establecimiento, día y grupo etario, 2020 en adelante.",
             "ZIP con CSV anual",
             "El año en curso se sobrescribe periódicamente; los años cerrados no."),
            ("DEIS / MINSAL — Egresos Hospitalarios",
             "Egresos disociados 2020–2025, con diagnóstico, previsión y días de estadía.",
             "ZIP con CSV anual", "Anual."),
            ("DEIS / MINSAL — Maestro de establecimientos",
             "Catálogo de establecimientos, tipo, Servicio de Salud y coordenadas.",
             "Planilla", "Se publica actualizado, sin versión."),
            ("INE — Censo 2024",
             "Cartografía comunal y subcomunal, y población censada.",
             "GeoParquet y planilla", "Cerrado."),
            ("INE — Proyecciones de población",
             "Población comunal anual estimada y proyectada, base Censo 2017.",
             "Planilla", "Cerrado."),
            ("MDS — Casen 2022",
             "Tasa de pobreza por ingresos comunal, como indicador de vulnerabilidad.",
             "Planilla", "Cerrado."),
            ("Estadísticas de Género",
             "Cuatro indicadores de salud mental por sexo, usados solo como contexto interpretativo.",
             "Planillas", "Cerrado."),
            ("OpenStreetMap y GTFS",
             "Red vial y transporte público para accesibilidad.",
             "Por definir", "Planificado."),
        ],
        [Cm(4.4), Cm(6.4), Cm(3.0), Cm(3.9)],
        font_size=8.8,
        first_col_bold=True,
    )
    note(
        doc,
        "**Tolerancia a fallos de las fuentes.** Si una descarga o su validación falla, el RAW "
        "previo se conserva y el comando termina con error; el sistema nunca reemplaza un snapshot "
        "válido por uno defectuoso (RNF-FIA-005).",
    )

    h2(doc, "9.2 Servicio externo en la interfaz")
    p(
        doc,
        "El mapa puede mostrar un mapa base de **CARTO Basemaps**, que el navegador solicita "
        "directamente por HTTPS. Es opcional: sin clave (`VITE_CARTO_API_KEY`), el mapa se dibuja "
        "solo con los límites comunales. La atribución a OpenStreetMap y CARTO se muestra siempre "
        "que el mapa base está activo, como exige su licencia de uso.",
    )

    h2(doc, "9.3 API REST")
    p(
        doc,
        "La API expone la versión 1 bajo el prefijo `/api/v1`. Es de solo lectura, responde JSON "
        "comprimido cuando la respuesta lo amerita y publica su documentación interactiva "
        "(OpenAPI) en `/docs`.",
    )
    add_table(
        doc,
        ["Método y ruta", "Qué entrega", "Parámetros"],
        [
            ("GET `/salud`", "Estado de la carga de datos y del pronóstico.", "—"),
            ("GET `/meta`",
             "Modelo servido, horizontes, nivel del intervalo, última semana observada, semana en "
             "curso, origen del pronóstico, cobertura real fuera de muestra y advertencias de lectura.",
             "—"),
            ("GET `/territorios`",
             "Región y las 52 comunas, indicando cuáles tienen pronóstico y, si no, por qué.", "—"),
            ("GET `/series/{series_id}`",
             "Historia observada de una serie y su pronóstico con intervalo.",
             "`semanas` (1 a 520; por defecto 104)"),
            ("GET `/pronosticos`",
             "Pronóstico de todas las series.",
             "`nivel` (region o comuna), `horizonte`"),
            ("GET `/resumen-comunal`",
             "Una fila por comuna: nivel reciente, pronóstico, cambio esperado y contexto territorial.",
             "—"),
            ("GET `/geo/comunas`",
             "Límites comunales simplificados en GeoJSON.", "—"),
        ],
        [Cm(4.2), Cm(9.2), Cm(4.3)],
        font_size=9.0,
        first_col_bold=True,
    )

    h3(doc, "Contrato de errores")
    add_table(
        doc,
        ["Código", "Cuándo ocurre", "Qué debe hacer el cliente"],
        [
            ("200", "La respuesta está completa.", "Usarla."),
            ("404", "La serie no existe o no es modelable; el cuerpo trae el motivo.",
             "Mostrar el motivo; no reintentar."),
            ("503 con `Retry-After`", "Los datos cargan o el pronóstico se calcula.",
             "Reintentar tras el tiempo indicado."),
            ("503 sin `Retry-After`", "Falta un output del pipeline; el cuerpo nombra la etapa a ejecutar.",
             "Informar el problema; no reintentar."),
        ],
        [Cm(4.2), Cm(8.1), Cm(5.4)],
        font_size=9.0,
        first_col_bold=True,
    )

    h3(doc, "Convenciones del contrato")
    bullets(
        doc,
        [
            "Todo valor ausente se entrega como `null`, nunca como cero.",
            "Las cifras son atenciones (eventos), no personas únicas; el contrato lo declara en sus descripciones.",
            "Cada pronóstico viaja con sus límites inferior y superior; no existe un endpoint que entregue solo el valor puntual.",
            "Las advertencias de interpretación forman parte de `/meta`, de modo que cualquier cliente las recibe.",
        ],
    )


def build_technologies(doc) -> None:
    h1(doc, "10. Tecnologías y decisiones arquitectónicas")

    h2(doc, "10.1 Tecnologías y su justificación")
    add_table(
        doc,
        ["Tecnología", "Uso", "Por qué", "Alternativa descartada"],
        [
            ("Python 3.12", "Pipeline, modelos y API.",
             "Un solo lenguaje desde la ingesta hasta el servicio; ecosistema analítico maduro (RT-01).",
             "R: separa el análisis del servicio."),
            ("pandas 3 y PyArrow 25", "Transformación y lectura/escritura columnar.",
             "Lectura selectiva de columnas y filtros sobre Parquet sin cargar archivos completos.",
             "Leer CSV completos en memoria."),
            ("Parquet y GeoParquet", "Formato de la capa procesada.",
             "Tipado, compresión y metadatos propios; regenerable y portable.",
             "CSV: sin tipos y más pesado."),
            ("LightGBM 4.7", "Modelo global de las comunas.",
             "Aprende de todas las series a la vez y soporta conteos con objetivo de Poisson.",
             "Un modelo por comuna: pocas observaciones en las pequeñas."),
            ("scikit-learn 1.9 y statsmodels 0.15", "GLM de Poisson regional y pruebas estadísticas.",
             "Modelo interpretable para la serie agregada; herramientas estándar de diagnóstico.",
             "Redes neuronales: sin datos suficientes."),
            ("FastAPI 0.141 y Pydantic 2", "API y contratos de respuesta.",
             "Validación tipada y documentación OpenAPI generada desde el mismo contrato (RNF-CMP-003).",
             "Flask: contratos y documentación a mano."),
            ("Uvicorn 0.53", "Servidor de la API.", "Servidor estándar de FastAPI; arranque simple en local.", "—"),
            ("React 19 y TypeScript 7", "Aplicación web.",
             "Componentes reutilizables y tipos que reflejan el contrato de la API (RT-02).",
             "Plantillas del servidor: acoplan interfaz y servicio."),
            ("Vite 6", "Construcción y servidor de desarrollo.",
             "Recarga inmediata y proxy hacia la API sin configuración adicional.", "—"),
            ("Tailwind CSS 4", "Estilos.",
             "Los tokens de la marca se declaran una vez y se usan en toda la interfaz.",
             "Hojas de estilo sueltas."),
            ("Recharts 3", "Gráfico de la serie.", "Composición declarativa de línea, pronóstico e intervalo.", "—"),
            ("Leaflet 1.9 y React Leaflet 5", "Mapa comunal.",
             "Cartografía interactiva liviana, sin servicio propio de mapas.", "—"),
            ("Shapely 2", "Simplificación de geometrías.",
             "Reduce el peso del GeoJSON que viaja al navegador.", "—"),
            ("pytest 9", "Pruebas automatizadas.", "Regla técnica del proyecto (RT-06).", "—"),
            ("python-docx y Matplotlib", "Documentos y figuras.",
             "La documentación formal se genera desde código versionado.", "Edición manual."),
            ("PostgreSQL con PostGIS", "Capa de servicio relacional (planificada).",
             "Consultas espaciales sobre polígonos y rutas (RT-03).", "—"),
            ("Docker Compose", "Empaquetado (planificado).",
             "Mismo entorno entre integrantes y en el despliegue (RT-04).", "—"),
        ],
        [Cm(3.5), Cm(3.6), Cm(6.6), Cm(4.0)],
        font_size=8.6,
        first_col_bold=True,
    )

    h2(doc, "10.2 Registro de decisiones arquitectónicas")
    p(
        doc,
        "Cada decisión se registra con el problema que resuelve, lo que se decidió y lo que cuesta. "
        "Las consecuencias incluyen las negativas: una decisión sin costo declarado no fue analizada.",
    )
    decisiones = [
        (
            "ADR-01 · Pipeline por lotes orquestado como DAG",
            "Las cifras deben poder reconstruirse y las fuentes cambian sin aviso.",
            "Todo output canónico se genera por `scripts/run_pipeline.py`, con etapas que declaran "
            "entradas, salidas y validación, y que hacen SKIP cuando su resultado sigue siendo válido.",
            "A favor: reproducibilidad y regeneración mínima. En contra: cada script nuevo debe "
            "registrarse en el DAG, en `PIPELINE.md` y en las pruebas.",
        ),
        (
            "ADR-02 · Parquet como almacenamiento de la capa procesada; base relacional diferida",
            "RT-03 pide PostgreSQL/PostGIS, pero el producto aún estaba definiendo qué tablas necesitaba "
            "y no hay infraestructura.",
            "La capa procesada vive en archivos Parquet que la API lee directamente. El diseño físico "
            "relacional se mantiene escrito y sincronizado, y se cargará cuando lo exija el análisis "
            "espacial o el despliegue.",
            "A favor: cero servicios que mantener y datos regenerables desde Git. En contra: RT-03 queda "
            "parcialmente atendida y las consultas espaciales complejas no están disponibles.",
        ),
        (
            "ADR-03 · RAW inmutable con procedencia por SHA256",
            "El archivo del año en curso se sobrescribe en la fuente y no conserva versiones.",
            "Se descarga a una ubicación temporal, se valida, se registra su huella y recién entonces se "
            "publica. Las etapas posteriores solo leen el RAW.",
            "A favor: se sabe con qué datos se obtuvo cada resultado. En contra: un snapshot antiguo no "
            "puede reconstruirse solo desde su URL si la fuente ya lo reemplazó.",
        ),
        (
            "ADR-04 · Interfaz única de modelo, registro y selección congelada",
            "Comparar modelos solo es honesto si todos se evalúan igual, y servir un modelo solo es "
            "correcto si es exactamente el evaluado.",
            "Todos los modelos cumplen una interfaz mínima. El benchmark elige uno por nombre y lo deja "
            "en un archivo de selección; holdout, pronóstico y API lo instancian desde un registro único.",
            "A favor: reemplazar el modelo no toca la API (RNF-FLE-005). En contra: cambiar un "
            "hiperparámetro obliga a reejecutar benchmark y holdout.",
        ),
        (
            "ADR-05 · Se persiste el pronóstico, no el modelo entrenado",
            "Ajustar el modelo en cada arranque de la API tomaba decenas de segundos.",
            "Una etapa del pipeline calcula el pronóstico una vez y lo guarda con un manifiesto de "
            "modelo, configuración y huella de sus entradas. La API lo reutiliza si sigue vigente.",
            "A favor: arranque en milisegundos y procedencia explícita. En contra: la vigencia depende "
            "de qué entradas vigila el manifiesto (ver deuda técnica DT-03).",
        ),
        (
            "ADR-06 · API de solo lectura que responde desde memoria",
            "La consulta debe ser rápida y no puede alterar datos.",
            "La API carga los artefactos al arrancar y solo filtra y serializa. Lo que aún no está listo "
            "se informa con `503` y `Retry-After`.",
            "A favor: latencia baja y superficie de riesgo mínima. En contra: estado por proceso; un "
            "snapshot nuevo exige reiniciar y no se puede escalar horizontalmente sin replantearlo.",
        ),
        (
            "ADR-07 · Aplicación web separada de la API",
            "La interfaz y el análisis evolucionan a ritmos distintos y con responsables distintos.",
            "La aplicación React se comunica con la API solo por HTTP/JSON, a través de un cliente tipado.",
            "A favor: cada parte se construye y prueba por separado. En contra: el contrato debe "
            "mantenerse a mano en dos lugares (`schemas.py` y `api.ts`).",
        ),
        (
            "ADR-08 · Urgencias y Egresos con granularidades separadas",
            "Egresos no identifica el establecimiento y ninguna fuente identifica pacientes.",
            "El modelo de datos mantiene ambas fuentes como entidades separadas; toda relación entre "
            "ellas es territorial y ecológica.",
            "A favor: impide por construcción un cruce indebido (RN-05). En contra: no se puede seguir "
            "la trayectoria de un caso desde la urgencia hasta la hospitalización.",
        ),
        (
            "ADR-09 · Documentación formal y figuras generadas desde código",
            "Los documentos editados a mano se desalinean del sistema y no dejan historial revisable.",
            "Los documentos de requisitos, este documento y sus figuras se generan con scripts versionados.",
            "A favor: cambios revisables y regeneración inmediata. En contra: editar el documento exige "
            "editar su script.",
        ),
    ]
    for titulo, contexto, decision, consecuencias in decisiones:
        h3(doc, titulo)
        add_table(
            doc,
            ["Aspecto", "Detalle"],
            [("Contexto", contexto), ("Decisión", decision), ("Consecuencias", consecuencias)],
            [Cm(3.0), Cm(14.7)],
            font_size=9.2,
            first_col_bold=True,
            keep_together=True,
        )


def build_quality(doc) -> None:
    h1(doc, "11. Cómo responde la arquitectura a los atributos de calidad")
    p(
        doc,
        "La tabla relaciona los requisitos no funcionales con el mecanismo arquitectónico que los "
        "atiende. La columna de estado es deliberadamente estricta: «Atendido» significa que el "
        "mecanismo existe y se puede señalar en el código, no que el valor objetivo del requisito "
        "haya sido medido.",
    )
    add_table(
        doc,
        ["Requisito", "Mecanismo arquitectónico", "Estado"],
        [
            ("RNF-MAN-001 Modularidad",
             "Capas con dependencia en un solo sentido; contratos entre capas.", "Atendido"),
            ("RNF-MAN-004 Prueba por función analítica",
             "415 pruebas sobre transformaciones, modelos, contratos y API.", "Atendido"),
            ("RNF-MAN-005 Integración continua",
             "No existe un flujo automático que ejecute las pruebas.", "Brecha"),
            ("RNF-FLE-002 Extensibilidad territorial",
             "Sin años ni comunas fijos; el RAW de Urgencias es nacional y se filtra en la limpieza.",
             "Atendido en el diseño; no probado en otra región"),
            ("RNF-FLE-004 Instalación reproducible",
             "Dependencias fijadas por versión; carpetas creadas por el código.",
             "Parcial: falta el empaquetado en contenedores"),
            ("RNF-FLE-005 Reemplazabilidad del modelo",
             "Interfaz de modelo, registro único y selección congelada.", "Atendido"),
            ("RNF-FLE-007 Portabilidad de rutas",
             "Rutas relativas al proyecto en código, informes y documentos.", "Atendido"),
            ("RNF-FIA-002 Degradación controlada",
             "`503` con `Retry-After`; recálculo en memoria si el pronóstico persistido no sirve.",
             "Atendido"),
            ("RNF-FIA-003 Publicación atómica",
             "Escritura en temporal, validación y reemplazo de una sola vez.", "Atendido"),
            ("RNF-FIA-005 Fallos de las fuentes",
             "Un RAW válido nunca se reemplaza por una descarga fallida.", "Atendido"),
            ("RNF-DAT-007 Trazabilidad",
             "Manifiesto de procedencia, metadatos del Parquet y manifiesto del pronóstico.", "Atendido"),
            ("RNF-EDE-002 Latencia de consulta",
             "Respuesta desde memoria, compresión y geometrías simplificadas.",
             "Atendido en el diseño; pendiente de medición formal"),
            ("RNF-CMP-003 Interfaz estándar",
             "REST con JSON y especificación OpenAPI generada.", "Atendido"),
            ("RNF-SEG-001 Sin datos personales",
             "Solo fuentes agregadas o disociadas; ningún identificador de persona en el modelo.", "Atendido"),
            ("RNF-SEG-002 Cifrado del canal",
             "En desarrollo local se usa HTTP; HTTPS depende del despliegue.", "Brecha hasta desplegar"),
            ("RNF-SEG-004 Gestión de secretos",
             "Variables de entorno; plantilla sin valores; archivo local ignorado por Git.", "Atendido"),
            ("RNF-USO-001 Incertidumbre visible",
             "El contrato no permite entregar un pronóstico sin su intervalo.", "Atendido"),
            ("RNF-USO-004 Interpretación indebida",
             "Advertencias servidas por `/meta` y mostradas por la interfaz.", "Atendido"),
        ],
        [Cm(5.0), Cm(8.4), Cm(4.3)],
        font_size=8.8,
        first_col_bold=True,
    )


def build_constraints(doc) -> None:
    h1(doc, "12. Restricciones técnicas, brechas y deuda")

    h2(doc, "12.1 Restricciones técnicas y su cumplimiento")
    add_table(
        doc,
        ["ID", "Restricción", "Cumplimiento vigente"],
        [
            ("RT-01", "Backend en Python con FastAPI; datos con pandas y PyArrow.", "Cumplida."),
            ("RT-02", "Frontend en React con Tailwind CSS, Recharts y React Leaflet.", "Cumplida."),
            ("RT-03", "Persistencia en PostgreSQL con PostGIS.",
             "Parcial: el diseño físico existe y está sincronizado; el sistema no lo usa aún (ADR-02)."),
            ("RT-04", "Empaquetado y despliegue con Docker Compose.",
             "No cumplida: no hay archivos de contenedor ni de composición."),
            ("RT-05", "Control de versiones en Git/GitHub y gestión en GitHub Projects.", "Cumplida."),
            ("RT-06", "Pruebas automatizadas con pytest.", "Cumplida."),
            ("RT-07", "Dependencias explícitas y justificadas en `requirements.txt`.", "Cumplida."),
        ],
        [Cm(1.7), Cm(7.6), Cm(8.4)],
        font_size=9.0,
        first_col_bold=True,
    )

    h2(doc, "12.2 Restricciones derivadas de los datos")
    add_table(
        doc,
        ["ID", "Restricción", "Efecto sobre la arquitectura"],
        [
            ("RD-01", "La serie de salud mental es comparable desde 2021.",
             "El modelo entrena desde 2021; 2020 se excluye por regla, no por un filtro oculto."),
            ("RD-02", "Urgencias son recuentos agregados, no registros individuales.",
             "El target es un conteo de eventos; no hay entidad «paciente» en el modelo."),
            ("RD-03", "Egresos no incluye el código de establecimiento.",
             "No hay relación entre Egresos y establecimientos; su territorio es el de residencia."),
            ("RD-04", "Parte del maestro de establecimientos carece de coordenadas.",
             "Esos establecimientos se excluyen del cálculo espacial y se informan; no se les estima ubicación."),
            ("—", "Cada atención se registra en la comuna del establecimiento, no en la de residencia.",
             "El mapa comunal muestra dónde se atiende, no dónde vive la demanda; la interfaz lo advierte."),
        ],
        [Cm(1.7), Cm(7.0), Cm(9.0)],
        font_size=9.0,
        first_col_bold=True,
    )

    h2(doc, "12.3 Deuda técnica y riesgos de arquitectura")
    add_table(
        doc,
        ["ID", "Situación", "Riesgo", "Tratamiento propuesto"],
        [
            ("DT-01", "No hay empaquetado en contenedores (RT-04).",
             "El entorno puede diferir entre integrantes y no hay una forma única de desplegar.",
             "Agregar Docker Compose para API y aplicación; sumar la base cuando se use."),
            ("DT-02", "No hay integración continua (RNF-MAN-005).",
             "Un cambio puede romper pruebas sin que nadie lo note al integrar.",
             "Flujo que ejecute pytest y la construcción del frontend en cada cambio propuesto."),
            ("DT-03", "El manifiesto del pronóstico vigila el mart cerrado y la selección, no los datos del año en curso.",
             "Tras actualizar Urgencias, el pronóstico persistido puede seguir considerándose vigente.",
             "Incluir en el manifiesto la huella del snapshot del año en curso."),
            ("DT-04", "La base relacional está diseñada pero no cargada (RT-03).",
             "El diseño puede desalinearse de los Parquet si no se ejercita.",
             "Cargarla cuando se implemente la accesibilidad territorial, que es quien la necesita."),
            ("DT-05", "El contrato de la API se mantiene en dos archivos.",
             "Un campo puede cambiar en el servidor y no en la interfaz.",
             "Generar los tipos del frontend desde la especificación OpenAPI."),
            ("DT-06", "El diagrama de componentes (`componentes_saad`) es anterior a la construcción del modelo, la API y la interfaz.",
             "Marca como planificados componentes que ya existen.",
             "Actualizarlo; es parte de la tarea abierta D2-16."),
            ("DT-07", "La API es un proceso único con estado en memoria.",
             "No escala horizontalmente y exige reinicio para tomar datos nuevos.",
             "Aceptado para el MVP; revisar si el uso supera lo previsto."),
            ("DT-08", "Dos componentes del MVP siguen planificados: accesibilidad y análisis de hospitalización.",
             "El MVP declarado no está completo.",
             "Priorizarlos antes de agregar funciones nuevas."),
        ],
        [Cm(1.5), Cm(5.4), Cm(5.2), Cm(5.6)],
        font_size=8.8,
        first_col_bold=True,
    )
    note(
        doc,
        "**Sobre la arquitectura objetivo.** Las brechas DT-01 y DT-04 no cuestionan el estilo "
        "elegido: son partes de la arquitectura objetivo que aún no se construyen. La separación "
        "entre pipeline, API y aplicación es la que permite agregarlas sin rehacer lo existente.",
        fill=BRAND_SKY_SOFT,
    )


def build_traceability(doc) -> None:
    h1(doc, "13. Trazabilidad")

    h2(doc, "13.1 Criterios de aceptación del documento")
    add_table(
        doc,
        ["Criterio (tarea D2-15)", "Dónde se atiende"],
        [
            ("Arquitectura seleccionada descrita", "Sección 3."),
            ("Componentes principales definidos", "Sección 4."),
            ("Tecnologías justificadas", "Sección 10.1."),
            ("Integraciones y servicios/API documentados", "Sección 9."),
            ("Restricciones técnicas incluidas", "Secciones 2.4 y 12."),
        ],
        [Cm(8.2), Cm(9.5)],
        font_size=9.5,
        first_col_bold=True,
    )

    h2(doc, "13.2 Requisitos funcionales y componentes")
    add_table(
        doc,
        ["Grupo de requisitos (ERF)", "Componentes que lo realizan", "Estado"],
        [
            ("RF-ETL y RF-CAL · Ingesta y calidad de datos",
             "Ingesta y procedencia, limpieza, contratos de output.", "Implementado"),
            ("RF-DEM · Proyección de demanda",
             "Modelos de demanda, benchmark y holdout, persistencia del pronóstico.", "Implementado"),
            ("RF-ACC · Análisis territorial y accesibilidad",
             "Marts territoriales (implementado); accesibilidad por red (planificado).", "Parcial"),
            ("RF-HOS · Caracterización de hospitalizaciones",
             "Dataset canónico de Egresos F00–F99 (implementado); análisis (planificado).", "Parcial"),
            ("RF-VIS · Visualización de resultados",
             "Aplicación web: portada, serie, mapa comunal y tabla.", "Implementado"),
            ("RF-API · API de consulta", "API REST de solo lectura.", "Implementado"),
            ("RF-TRZ y RF-ADM · Trazabilidad y administración",
             "Orquestador, manifiestos de procedencia y del pronóstico.", "Implementado"),
        ],
        [Cm(6.0), Cm(8.6), Cm(3.1)],
        font_size=9.0,
        first_col_bold=True,
    )

    h2(doc, "13.3 Mantenimiento de este documento")
    p(
        doc,
        "El documento se regenera con `python scripts/generate_das_docx.py`, que también vuelve a "
        "dibujar las figuras 1, 4 y 5. Debe actualizarse cuando cambie un componente, un endpoint, "
        "una tecnología o una decisión registrada, y en particular cuando se cierre alguna de las "
        "brechas de la sección 12. Las figuras 2 y 3 provienen de los diagramas mantenidos en "
        "`docs/07-diagramas/` y se actualizan regenerándolos desde su fuente.",
    )


def main() -> None:
    apply_theme(BRAND_THEME)
    figuras = generar_todas(FIGURES_DIR)

    doc = new_document("DAS-SAAD-2026 · Documento de Arquitectura de Software · v1.0")
    build_cover(doc)
    build_control(doc)
    add_toc(doc)
    build_introduction(doc)
    build_drivers(doc)
    build_selected_architecture(doc, figuras)
    build_logical_view(doc)
    build_data_view(doc)
    build_process_view(doc, figuras)
    build_development_view(doc)
    build_deployment_view(doc, figuras)
    build_integrations(doc)
    build_technologies(doc)
    build_quality(doc)
    build_constraints(doc)
    build_traceability(doc)

    ruta = save(doc, OUTPUT_PATH)
    print(f"Documento generado: {ruta.as_posix()}")
    for nombre, figura in figuras.items():
        print(f"Figura {nombre}: {figura.as_posix()}")


if __name__ == "__main__":
    main()
