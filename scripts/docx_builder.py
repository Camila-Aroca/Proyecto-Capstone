"""Utilidades de construcción de documentos Word (.docx) para la documentación del proyecto.

Provee una capa delgada sobre `python-docx` con estilos, portadas, tablas e índices
consistentes, de modo que los documentos formales del proyecto sean reproducibles
mediante código versionado.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Optional, Sequence, Tuple

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# Paleta institucional del documento (tema clásico, usado por defecto)
COLOR_H1 = RGBColor(0x1F, 0x3B, 0x63)
COLOR_H2 = RGBColor(0x2E, 0x5E, 0x8C)
COLOR_H3 = RGBColor(0x44, 0x54, 0x6A)
COLOR_TEXT = RGBColor(0x1A, 0x1A, 0x1A)
COLOR_MUTED = RGBColor(0x59, 0x59, 0x59)
SHADE_HEADER = "1F3B63"
SHADE_SOFT = "EAF0F6"
SHADE_NOTE = "FFF4E5"

BODY_FONT = "Calibri"
HEADING_FONT = "Calibri"
MONO_FONT = "Consolas"
# Fuente que Word usa cuando la del tema no está instalada en el equipo del lector.
FALLBACK_FONT = "Calibri"
# Color de la regla bajo los títulos de primer nivel; None la omite.
RULE_COLOR: Optional[str] = None
ACCENT_COLOR = "C08B2F"


@dataclass(frozen=True)
class Theme:
    """Colores y tipografías de un documento. Los colores van en hexadecimal sin `#`."""

    h1: str
    h2: str
    h3: str
    text: str
    muted: str
    table_header: str
    zebra: str
    note: str
    accent: str
    heading_font: str
    body_font: str
    rule: Optional[str] = None


CLASSIC_THEME = Theme(
    h1="1F3B63", h2="2E5E8C", h3="44546A", text="1A1A1A", muted="595959",
    table_header="1F3B63", zebra="EAF0F6", note="FFF4E5", accent="C08B2F",
    heading_font="Calibri", body_font="Calibri",
)

# Identidad visual SAAD: valores de `saad-tokens.css` del kit de marca
# (primary-deep, primary, secondary, text-*, bg-alt, accent-soft, accent) y sus
# dos familias: Plus Jakarta Sans para titulares e Inter para texto.
BRAND_THEME = Theme(
    h1="0B4660", h2="0F5A78", h3="2F74A8", text="0F2A38", muted="5E7685",
    table_header="0B4660", zebra="EEF3F6", note="FCF1DC", accent="EBA845",
    heading_font="Plus Jakarta Sans", body_font="Inter", rule="86BEA9",
)
# Tonos suaves de la marca para cuadros informativos (`mint-soft`, `sky-soft`).
BRAND_MINT_SOFT = "E3F1EB"
BRAND_SKY_SOFT = "E3F0F8"
BRAND_LOGO_PATH = Path("frontend/public/marca/saad-logo.png")


def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color.upper())


def apply_theme(theme: Theme) -> None:
    """Fija el tema de los documentos que se creen a continuación.

    Los bloques de este módulo leen la paleta al momento de escribir, de modo que
    basta llamar a esta función antes de `new_document`. Sin llamarla rige el tema
    clásico, con lo que los documentos existentes no cambian.
    """
    global COLOR_H1, COLOR_H2, COLOR_H3, COLOR_TEXT, COLOR_MUTED
    global SHADE_HEADER, SHADE_SOFT, SHADE_NOTE, ACCENT_COLOR
    global BODY_FONT, HEADING_FONT, RULE_COLOR
    COLOR_H1, COLOR_H2, COLOR_H3 = _rgb(theme.h1), _rgb(theme.h2), _rgb(theme.h3)
    COLOR_TEXT, COLOR_MUTED = _rgb(theme.text), _rgb(theme.muted)
    SHADE_HEADER, SHADE_SOFT, SHADE_NOTE = theme.table_header, theme.zebra, theme.note
    ACCENT_COLOR = theme.accent
    BODY_FONT, HEADING_FONT, RULE_COLOR = theme.body_font, theme.heading_font, theme.rule


# --------------------------------------------------------------------------------------
# Infraestructura de bajo nivel (OOXML)
# --------------------------------------------------------------------------------------
def _set_cell_background(cell, hex_fill: str) -> None:
    """Aplica un color de fondo sólido a una celda de tabla."""
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    cell._tc.get_or_add_tcPr().append(shd)


def _keep_row_together(row) -> None:
    """Impide que una fila se parta entre dos páginas."""
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def _restart_numbering(document: Document, paragraph) -> None:
    """Hace que una lista numerada comience en 1 en vez de continuar la anterior.

    Todas las listas del estilo `List Number` comparten una misma definición; sin
    esto, la segunda lista del documento seguiría contando desde donde quedó la primera.
    """
    numbering = document.part.numbering_part.numbering_definitions._numbering
    style_num_pr = paragraph.style.element.pPr.numPr if paragraph.style.element.pPr is not None else None
    if style_num_pr is None or style_num_pr.numId is None:
        return
    base = numbering.num_having_numId(style_num_pr.numId.val)
    nueva = numbering.add_num(base.abstractNumId.val)
    nueva.add_lvlOverride(ilvl=0).add_startOverride(1)
    num_pr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
    num_pr.get_or_add_ilvl().val = 0
    num_pr.get_or_add_numId().val = nueva.numId


def _repeat_header_row(row) -> None:
    """Marca una fila como encabezado que se repite en cada página."""
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def _add_field(paragraph, instruction: str) -> None:
    """Inserta un campo dinámico de Word (por ejemplo, PAGE o TOC)."""
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    for element in (fld_begin, instr, fld_sep, fld_end):
        run._r.append(element)


def _set_bottom_border(paragraph_format_owner, hex_color: str, size_eighths: int = 8) -> None:
    """Dibuja una regla bajo un párrafo o estilo (`size_eighths` en octavos de punto)."""
    p_pr = paragraph_format_owner.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size_eighths))
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), hex_color)
    borders.append(bottom)
    p_pr.append(borders)


def _set_fonts(style_or_run, font_name: str) -> None:
    """Asigna la fuente a todos los rangos de escritura, no solo al latino básico."""
    style_or_run.font.name = font_name
    r_pr = style_or_run.element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    for attribute in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        r_fonts.set(qn(attribute), font_name)


def _register_font_fallbacks(document: Document, fonts: Sequence[str], fallback: str) -> None:
    """Declara en la tabla de fuentes qué usar si una fuente no está instalada.

    Las fuentes de marca son de Google Fonts y pueden faltar en el equipo de quien
    abre el archivo. Sin esta declaración Word elige el reemplazo por su cuenta;
    con `w:altName` el documento conserva una tipografía previsible.
    """
    from lxml import etree

    for relation in document.part.rels.values():
        if not relation.reltype.endswith("/fontTable"):
            continue
        part = relation.target_part
        root = etree.fromstring(part.blob)
        declared = {node.get(qn("w:name")) for node in root.findall(qn("w:font"))}
        for font in fonts:
            if font in declared or font == fallback:
                continue
            node = etree.SubElement(root, qn("w:font"))
            node.set(qn("w:name"), font)
            alt = etree.SubElement(node, qn("w:altName"))
            alt.set(qn("w:val"), fallback)
        part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _enable_update_fields(document: Document) -> None:
    """Solicita a Word que actualice los campos (índice, páginas) al abrir el archivo."""
    try:
        settings = document.settings.element
    except (AttributeError, KeyError):
        return
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)


# --------------------------------------------------------------------------------------
# Texto enriquecido
# --------------------------------------------------------------------------------------
_INLINE_PATTERN = re.compile(r"(\*\*.+?\*\*|`.+?`)")


def add_rich_text(paragraph, text: str, size: Optional[float] = None) -> None:
    """Escribe texto admitiendo marcas simples: **negrita** y `monoespaciado`."""
    for token in _INLINE_PATTERN.split(text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = MONO_FONT
            run.font.size = Pt((size or 10.5) - 0.5)
        else:
            run = paragraph.add_run(token)
        if size is not None and run.font.size is None:
            run.font.size = Pt(size)


# --------------------------------------------------------------------------------------
# Documento y estilos base
# --------------------------------------------------------------------------------------
def new_document(footer_text: str) -> Document:
    """Crea un documento con márgenes, tipografía, estilos y pie de página del proyecto."""
    document = Document()

    section = document.sections[0]
    section.page_width = Cm(21.59)
    section.page_height = Cm(27.94)
    section.left_margin = Cm(1.9)
    section.right_margin = Cm(1.9)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)

    normal = document.styles["Normal"]
    _set_fonts(normal, BODY_FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = COLOR_TEXT
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for style_name, size, color, space_before in (
        ("Heading 1", 16, COLOR_H1, 18),
        ("Heading 2", 13, COLOR_H2, 14),
        ("Heading 3", 11.5, COLOR_H3, 10),
    ):
        style = document.styles[style_name]
        _set_fonts(style, HEADING_FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(space_before)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
    if RULE_COLOR:
        _set_bottom_border(document.styles["Heading 1"].element, RULE_COLOR)

    _register_font_fallbacks(document, [BODY_FONT, HEADING_FONT], FALLBACK_FONT)
    _build_footer(section, footer_text)
    _enable_update_fields(document)
    return document


def _build_footer(section, footer_text: str) -> None:
    """Construye el pie de página con identificación del documento y numeración."""
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(f"{footer_text}  |  Página ")
    run.font.size = Pt(8)
    run.font.color.rgb = COLOR_MUTED
    _add_field(paragraph, " PAGE ")
    run = paragraph.add_run(" de ")
    run.font.size = Pt(8)
    run.font.color.rgb = COLOR_MUTED
    _add_field(paragraph, " NUMPAGES ")
    for run in paragraph.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = COLOR_MUTED


# --------------------------------------------------------------------------------------
# Bloques de contenido
# --------------------------------------------------------------------------------------
def add_cover(
    document: Document,
    title: str,
    subtitle: str,
    project: str,
    meta_rows: Sequence[Tuple[str, str]],
) -> None:
    """Genera la portada del documento."""
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(48)

    para = document.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run(project.upper())
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = COLOR_MUTED

    para = document.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(18)
    run = para.add_run(title)
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = COLOR_H1

    para = document.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_after = Pt(36)
    run = para.add_run(subtitle)
    run.font.size = Pt(12)
    run.font.color.rgb = COLOR_H2

    add_table(document, ["Campo", "Detalle"], meta_rows, [Cm(5.0), Cm(12.7)], font_size=10)
    document.add_page_break()


def add_brand_cover(
    document: Document,
    title: str,
    subtitle: str,
    project: str,
    meta_rows: Sequence[Tuple[str, str]],
    logo_path: Path = BRAND_LOGO_PATH,
    kicker: str = "",
) -> None:
    """Portada con la identidad del proyecto: logo, banda de título y ficha del documento.

    La banda usa el color primario profundo con el título en blanco y una regla en el
    color de acento, que es la misma jerarquía de la portada de la plataforma.
    """
    if logo_path.is_file():
        logo = document.add_paragraph()
        logo.paragraph_format.space_before = Pt(6)
        logo.paragraph_format.space_after = Pt(20)
        logo.add_run().add_picture(str(logo_path), width=Cm(6.2))

    band = document.add_table(rows=1, cols=1)
    band.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = band.cell(0, 0)
    _set_cell_background(cell, SHADE_HEADER)
    cell.width = Cm(17.8)

    first = cell.paragraphs[0]
    first.paragraph_format.space_before = Pt(26)
    first.paragraph_format.space_after = Pt(4)
    first.paragraph_format.left_indent = Cm(0.6)
    run = first.add_run((kicker or project).upper())
    _set_fonts(run, BODY_FONT)
    run.font.size = Pt(9)
    run.font.bold = True
    run.font.color.rgb = _rgb(ACCENT_COLOR)

    heading = cell.add_paragraph()
    heading.paragraph_format.left_indent = Cm(0.6)
    heading.paragraph_format.right_indent = Cm(0.6)
    heading.paragraph_format.space_after = Pt(8)
    heading.paragraph_format.line_spacing = 1.0
    run = heading.add_run(title)
    _set_fonts(run, HEADING_FONT)
    run.font.size = Pt(28)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    sub = cell.add_paragraph()
    sub.paragraph_format.left_indent = Cm(0.6)
    sub.paragraph_format.right_indent = Cm(0.6)
    sub.paragraph_format.space_after = Pt(26)
    run = sub.add_run(subtitle)
    _set_fonts(run, BODY_FONT)
    run.font.size = Pt(11.5)
    run.font.color.rgb = RGBColor(0xE3, 0xF0, 0xF8)

    rule = document.add_paragraph()
    rule.paragraph_format.space_after = Pt(18)
    _set_bottom_border(rule._p, ACCENT_COLOR, size_eighths=24)

    if kicker:
        line = document.add_paragraph()
        line.paragraph_format.space_after = Pt(14)
        run = line.add_run(project)
        run.font.size = Pt(10)
        run.font.color.rgb = COLOR_MUTED

    add_table(document, ["Campo", "Detalle"], meta_rows, [Cm(5.0), Cm(12.7)], font_size=10)
    document.add_page_break()


def add_picture(document: Document, image, width_cm: float = 17.4) -> None:
    """Inserta una imagen centrada, sin leyenda. `image` es una ruta o un flujo binario."""
    if isinstance(image, Path):
        if not image.is_file():
            raise FileNotFoundError(f"No existe la figura: {image.as_posix()}")
        image = str(image)
    picture = document.add_paragraph()
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.paragraph_format.space_before = Pt(6)
    picture.paragraph_format.space_after = Pt(4)
    picture.paragraph_format.keep_with_next = True
    picture.add_run().add_picture(image, width=Cm(width_cm))


def add_figure(document: Document, image, caption: str, width_cm: float = 17.4) -> None:
    """Inserta una imagen centrada con su leyenda numerada por quien llama."""
    add_picture(document, image, width_cm)

    legend = document.add_paragraph()
    legend.alignment = WD_ALIGN_PARAGRAPH.CENTER
    legend.paragraph_format.space_after = Pt(10)
    add_rich_text(legend, caption, size=9)
    for run in legend.runs:
        run.font.color.rgb = COLOR_MUTED


def add_toc(document: Document, title: str = "Tabla de contenidos") -> None:
    """Inserta un índice automático que Word actualiza al abrir el documento."""
    heading = document.add_paragraph(title, style="Heading 1")
    heading.paragraph_format.space_before = Pt(0)
    para = document.add_paragraph()
    _add_field(para, r' TOC \o "1-2" \h \z \u ')
    note = document.add_paragraph()
    run = note.add_run(
        "Si el índice no se muestra actualizado, seleccionarlo y presionar F9 en Microsoft Word."
    )
    run.font.size = Pt(8.5)
    run.font.italic = True
    run.font.color.rgb = COLOR_MUTED
    document.add_page_break()


def h1(document: Document, text: str) -> None:
    document.add_paragraph(text, style="Heading 1")


def h2(document: Document, text: str) -> None:
    document.add_paragraph(text, style="Heading 2")


def h3(document: Document, text: str) -> None:
    document.add_paragraph(text, style="Heading 3")


def p(document: Document, text: str, italic: bool = False, size: float = 10.5) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_rich_text(paragraph, text, size=size)
    if italic:
        for run in paragraph.runs:
            run.font.italic = True


def bullets(document: Document, items: Iterable[str], size: float = 10.5) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(3)
        add_rich_text(paragraph, item, size=size)


def numbered(document: Document, items: Iterable[str], size: float = 10.5) -> None:
    numero = None
    for item in items:
        paragraph = document.add_paragraph(style="List Number")
        paragraph.paragraph_format.space_after = Pt(3)
        add_rich_text(paragraph, item, size=size)
        if numero is None:
            # La primera entrada abre una numeración propia; las demás la continúan.
            _restart_numbering(document, paragraph)
            num_pr = paragraph._p.pPr.numPr
            numero = num_pr.numId.val if num_pr is not None and num_pr.numId is not None else 0
        elif numero:
            num_pr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
            num_pr.get_or_add_ilvl().val = 0
            num_pr.get_or_add_numId().val = numero


def note(document: Document, text: str, fill: Optional[str] = None) -> None:
    """Cuadro destacado para notas, advertencias o supuestos.

    Sin `fill` usa el color de nota del tema activo.
    """
    table = document.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    _set_cell_background(cell, fill or SHADE_NOTE)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(4)
    add_rich_text(paragraph, text, size=10)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def add_table(
    document: Document,
    headers: Sequence[str],
    rows: Sequence[Sequence[str]],
    widths: Optional[Sequence[Cm]] = None,
    font_size: float = 9.0,
    first_col_bold: bool = False,
    keep_together: bool = False,
) -> None:
    """Inserta una tabla con encabezado destacado, filas cebra y ancho fijo por columna.

    El encabezado nunca queda solo al pie de una página y ninguna fila se parte entre
    dos. Con `keep_together`, la tabla completa se mantiene en una misma página; es
    para tablas cortas que se leen como una unidad.
    """
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = widths is None

    header_cells = table.rows[0].cells
    for index, text in enumerate(headers):
        cell = header_cells[index]
        _set_cell_background(cell, SHADE_HEADER)
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_before = Pt(2)
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.keep_with_next = True
        run = paragraph.add_run(text)
        run.bold = True
        run.font.size = Pt(font_size)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    _repeat_header_row(table.rows[0])
    _keep_row_together(table.rows[0])

    ultima = len(rows) - 1
    for row_index, row_values in enumerate(rows):
        row = table.add_row()
        _keep_row_together(row)
        cells = row.cells
        for col_index, value in enumerate(row_values):
            cell = cells[col_index]
            if row_index % 2 == 1:
                _set_cell_background(cell, SHADE_SOFT)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(2)
            paragraph.paragraph_format.space_after = Pt(2)
            if keep_together and row_index < ultima:
                paragraph.paragraph_format.keep_with_next = True
            add_rich_text(paragraph, str(value), size=font_size)
            for run in paragraph.runs:
                run.font.size = Pt(font_size)
                if first_col_bold and col_index == 0:
                    run.bold = True

    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = width

    document.add_paragraph().paragraph_format.space_after = Pt(4)


def add_definition_block(document: Document, title: str, lines: Sequence[Tuple[str, str]]) -> None:
    """Bloque etiqueta-valor usado en casos de uso y fichas de requisito."""
    h3(document, title)
    add_table(
        document,
        ["Elemento", "Detalle"],
        lines,
        [Cm(4.2), Cm(13.5)],
        font_size=9.5,
        first_col_bold=True,
    )


def save(document: Document, path: Path) -> Path:
    """Guarda el documento creando el directorio destino si es necesario."""
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)
    return path
