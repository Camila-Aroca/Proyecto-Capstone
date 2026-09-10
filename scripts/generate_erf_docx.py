"""Genera la Especificación de Requisitos Funcionales (ERF) del proyecto SAAD en formato .docx.

Marco de referencia: ISO/IEC/IEEE 29148:2018 e IEEE Std 830-1998.
Ejecución: python scripts/generate_erf_docx.py
Salida:    docs/06-requisitos/Especificacion_Requisitos_Funcionales.docx
"""

from __future__ import annotations

from pathlib import Path

from docx.shared import Cm

from docx_builder import (
    add_cover,
    add_definition_block,
    add_table,
    add_toc,
    bullets,
    h1,
    h2,
    h3,
    new_document,
    note,
    numbered,
    p,
    save,
)

OUTPUT_PATH = Path("docs/06-requisitos/Especificacion_Requisitos_Funcionales.docx")

# Anchos de columna del catálogo de requisitos (total útil: 17,7 cm)
REQ_WIDTHS = [Cm(2.0), Cm(3.2), Cm(5.4), Cm(5.2), Cm(1.9)]
REQ_HEADERS = ["ID", "Requisito", "Descripción funcional", "Criterio de aceptación", "Prio. / Ver."]


def build_cover(doc) -> None:
    add_cover(
        doc,
        title="Especificación de Requisitos Funcionales",
        subtitle=(
            "SAAD — Sistema de Análisis y Anticipación de Demanda de Urgencia "
            "en Salud Mental para la Región Metropolitana"
        ),
        project="Capstone APT · Ingeniería en Informática · Duoc UC Antonio Varas",
        meta_rows=[
            ("Código del documento", "ERF-SAAD-2026"),
            ("Versión", "1.0"),
            ("Fecha de emisión", "10 de septiembre de 2026"),
            ("Estado", "Borrador para validación con la docente y el cliente piloto"),
            ("Clasificación", "Uso académico e interno del equipo de proyecto"),
            (
                "Marco de referencia",
                "ISO/IEC/IEEE 29148:2018 (Requirements engineering); IEEE Std 830-1998",
            ),
            ("Elaborado por", "Equipo SAAD — Camila A., Felipe R., Cristopher R., Catalina"),
            (
                "Documento complementario",
                "Especificación de Requisitos No Funcionales (ERNF-SAAD-2026)",
            ),
            ("Repositorio", "docs/06-requisitos/"),
        ],
    )


def build_control(doc) -> None:
    h1(doc, "Control del documento")

    h2(doc, "Historial de versiones")
    add_table(
        doc,
        ["Versión", "Fecha", "Autor(es)", "Descripción del cambio"],
        [
            (
                "0.1",
                "27-08-2026",
                "Equipo SAAD",
                "Levantamiento preliminar a partir de la entrevista de dominio y la bitácora "
                "de definición de alcance.",
            ),
            (
                "1.0",
                "10-09-2026",
                "Equipo SAAD",
                "Primera versión formal completa: catálogo de requisitos funcionales, historias "
                "de usuario, casos de uso, reglas de negocio, exclusiones y trazabilidad.",
            ),
        ],
        [Cm(2.0), Cm(2.6), Cm(3.6), Cm(9.5)],
        font_size=9.5,
    )

    h2(doc, "Revisión y aprobación")
    add_table(
        doc,
        ["Rol", "Nombre", "Estado", "Fecha"],
        [
            ("Jefa de Proyecto", "Camila A.", "Emitido para revisión", "10-09-2026"),
            ("Docente guía", "Por confirmar", "Pendiente de validación", "—"),
            ("Cliente piloto", "Por confirmar", "Pendiente de validación", "—"),
        ],
        [Cm(4.5), Cm(5.0), Cm(5.2), Cm(3.0)],
        font_size=9.5,
    )

    note(
        doc,
        "**Nota de estado.** El cliente piloto se encuentra en proceso de confirmación (bitácora "
        "de definición de alcance, 23-08-2026, sección 5). Los requisitos identificados como "
        "*por validar* dependen de esa confirmación y podrán ajustarse en la versión 1.1 sin "
        "alterar el alcance acordado con la docente.",
    )
    doc.add_page_break()


def build_introduction(doc) -> None:
    h1(doc, "1. Introducción")

    h2(doc, "1.1 Propósito")
    p(
        doc,
        "Este documento especifica **qué debe hacer** el sistema SAAD: el conjunto completo de "
        "capacidades funcionales que la plataforma debe proveer a sus usuarios y a los sistemas "
        "con los que interopera. Constituye la línea base técnica para el diseño, la construcción, "
        "las pruebas y la validación del producto durante las fases 2 y 3 del Capstone APT.",
    )
    p(
        doc,
        "Los atributos de calidad —desempeño, seguridad, usabilidad, fiabilidad, mantenibilidad y "
        "restricciones normativas— se especifican en el documento complementario "
        "**Especificación de Requisitos No Funcionales (ERNF-SAAD-2026)**.",
    )

    h2(doc, "1.2 Alcance del producto")
    p(
        doc,
        "SAAD es una **plataforma analítica web** para la Región Metropolitana de Chile que integra "
        "fuentes oficiales del DEIS/MINSAL con insumos territoriales auxiliares, con tres "
        "finalidades:",
    )
    numbered(
        doc,
        [
            "**Anticipar** la demanda de atenciones de urgencia por causas de salud mental "
            "(CIE-10 F00–F99) en un horizonte de 4 a 8 semanas, con intervalos de predicción y "
            "evaluación contra un baseline estacional ingenuo.",
            "**Territorializar** las brechas de accesibilidad a la red pública de urgencia, "
            "mediante isócronas de red vial y transporte público, cobertura poblacional y cruce "
            "con vulnerabilidad socioeconómica comunal.",
            "**Caracterizar** los factores asociados a la duración de la hospitalización "
            "psiquiátrica a partir de los egresos hospitalarios con diagnóstico F00–F99.",
        ],
    )
    p(
        doc,
        "El sistema es un **instrumento de apoyo a la planificación sanitaria territorial**. No es "
        "un sistema clínico, no gestiona pacientes ni derivaciones y no sustituye el juicio "
        "profesional. Las exclusiones explícitas se detallan en la sección 8.",
    )

    h2(doc, "1.3 Definiciones, acrónimos y abreviaturas")
    add_table(
        doc,
        ["Término", "Definición"],
        [
            ("APS", "Atención Primaria de Salud."),
            (
                "Backtesting",
                "Evaluación retrospectiva de un modelo de pronóstico sobre datos históricos no "
                "utilizados en su entrenamiento.",
            ),
            (
                "Baseline estacional ingenuo",
                "Pronóstico de referencia que repite el valor observado en el mismo período de la "
                "temporada anterior; opera como umbral mínimo de utilidad del modelo.",
            ),
            (
                "CIE-10 F00–F99",
                "Capítulo V de la Clasificación Internacional de Enfermedades: trastornos mentales "
                "y del comportamiento.",
            ),
            (
                "Cruce ecológico",
                "Comparación entre fuentes a nivel de agregado territorial o temporal, sin enlace "
                "individuo a individuo ni establecimiento a establecimiento.",
            ),
            ("CUT", "Código Único Territorial de comuna (Chile)."),
            ("DEIS", "Departamento de Estadísticas e Información de Salud del MINSAL."),
            (
                "Egreso hospitalario",
                "Registro de salida de un paciente hospitalizado; base individual disociada "
                "publicada por el DEIS.",
            ),
            ("GTFS", "General Transit Feed Specification: formato estándar de transporte público."),
            (
                "Isócrona",
                "Polígono que agrupa el territorio alcanzable desde un punto en un tiempo de viaje "
                "determinado.",
            ),
            ("MASE", "Mean Absolute Scaled Error: error absoluto medio escalado sobre un baseline."),
            ("MVP", "Producto mínimo viable: alcance comprometido para la entrega final."),
            ("OSM", "OpenStreetMap: base cartográfica y de red vial abierta."),
            (
                "SAPU / SAR / SUR / UEH",
                "Servicio de Atención Primaria de Urgencia / Servicio de Alta Resolutividad / "
                "Servicio de Urgencia Rural / Unidad de Emergencia Hospitalaria.",
            ),
            ("SNSS", "Sistema Nacional de Servicios de Salud."),
            (
                "Subregistro",
                "Ausencia o incompletitud sistemática de registros que puede distorsionar "
                "comparaciones entre territorios o períodos.",
            ),
            (
                "`id_causa`",
                "Identificador de causa de consulta en la base DEIS de atenciones de urgencia "
                "(`1` = total de atenciones; `36` = total F00–F99).",
            ),
        ],
        [Cm(4.4), Cm(13.3)],
        font_size=9.5,
        first_col_bold=True,
    )

    h2(doc, "1.4 Referencias")
    add_table(
        doc,
        ["Código", "Documento o fuente"],
        [
            ("DEF-01", "Documento de Definición de Proyecto APT — Fase 1 (docs/01-definicion-proyecto/)."),
            ("ENT-01", "Bitácora de entrevista a Belén Guzmán, psicóloga, 17-08-2026 (docs/02-entrevistas/)."),
            ("BIT-01", "Bitácora de reunión de definición y alcance, 23-08-2026 (docs/03-bitacoras/)."),
            ("CTX-01", "Contexto técnico y decisiones vigentes del proyecto (PROJECT_CONTEXT.md)."),
            ("REG-01", "Reglas de rigor analítico, portabilidad y alcance (.agent/rules.md)."),
            ("EDA-01", "Auditoría e integridad de establecimientos RM (reports/eda_establecimientos_rm.md)."),
            ("EDA-02", "Cartografía comunal Censo 2024 RM (reports/eda_censo_comunal.md)."),
            ("EDA-03", "Auditoría técnica de formato y encoding RAW (reports/eda_auditoria_formato_raw.md)."),
            ("EDA-04", "Normalización y filtrado de urgencias RM (reports/eda_urgencias_normalizacion.md)."),
            (
                "EDA-05",
                "Validación post-normalización de urgencias RM "
                "(reports/eda_urgencias_validacion_post_normalizacion.md).",
            ),
            ("EDA-06", "Homologación de causas y catálogo F00–F99 (reports/eda_urgencias_causas_2020_2026.md)."),
            ("EDA-07", "Maestro de establecimientos y red asistencial (reports/eda_maestro_establecimientos_y_red.md)."),
            ("EDA-08", "Perfilado descriptivo de la demanda de urgencias RM (reports/eda_demanda_urgencias_rm.md)."),
            ("NOR-01", "ISO/IEC/IEEE 29148:2018 — Requirements engineering."),
            ("NOR-02", "IEEE Std 830-1998 — Recommended Practice for Software Requirements Specifications."),
            ("NOR-03", "ISO/IEC 25010 — Modelo de calidad de producto (aplicado en el ERNF)."),
        ],
        [Cm(2.4), Cm(15.3)],
        font_size=9.5,
        first_col_bold=True,
    )

    h2(doc, "1.5 Estructura del documento")
    p(
        doc,
        "La sección 2 describe el contexto, los actores y las restricciones generales. La sección 3 "
        "fija las convenciones de identificación, prioridad y verificación. La sección 4 contiene el "
        "catálogo de requisitos funcionales, organizado en nueve módulos. Las secciones 5 y 6 "
        "expresan esos mismos requisitos desde la perspectiva del usuario, mediante historias de "
        "usuario y casos de uso. Las secciones 7 y 8 delimitan las reglas de negocio y las "
        "exclusiones. Las secciones 9 a 12 aseguran trazabilidad, planificación por incrementos, "
        "criterios de aceptación global y control de cambios.",
    )
    doc.add_page_break()


def build_overview(doc) -> None:
    h1(doc, "2. Descripción general del sistema")

    h2(doc, "2.1 Perspectiva del producto")
    p(
        doc,
        "SAAD es un producto nuevo y autocontenido, que no reemplaza sistemas existentes ni se "
        "integra en línea con sistemas clínicos. Consume datos públicos por lotes, los procesa en "
        "un pipeline reproducible y expone los resultados mediante una API y una interfaz web.",
    )
    add_table(
        doc,
        ["Entradas (fuentes externas)", "Procesamiento (SAAD)", "Salidas y destinatarios"],
        [
            (
                "DEIS/MINSAL — Atenciones de urgencia 2020–2026\n"
                "DEIS/MINSAL — Egresos hospitalarios 2020–2025\n"
                "DEIS/MINSAL — Maestro de establecimientos\n"
                "INE / Censo 2024 — Cartografía y población comunal\n"
                "OpenStreetMap — Red vial\n"
                "GTFS — Transporte público de la RM",
                "1. Módulo de datos: ingesta, normalización, filtrado territorial y agregación.\n"
                "2. Módulo de calidad: validación de esquema, integridad, consistencia y cobertura.\n"
                "3. Módulo de demanda: series semanales, pronóstico 4–8 semanas y backtesting.\n"
                "4. Módulo de accesibilidad: isócronas, cobertura poblacional y brechas.\n"
                "5. Módulo de hospitalización: duración de estadía y factores asociados.\n"
                "6. API de servicios y aplicación web.",
                "Panel de indicadores y series con proyección\n"
                "Mapas de accesibilidad y brechas comunales\n"
                "Caracterización de hospitalización psiquiátrica\n"
                "Descargas tabulares y capas geoespaciales\n"
                "API documentada para consumo externo\n\n"
                "Destinatarios: gestores de red, referentes comunales de salud mental, analistas "
                "y jefaturas de urgencia.",
            )
        ],
        [Cm(5.3), Cm(6.6), Cm(5.8)],
        font_size=9,
    )

    h2(doc, "2.2 Funciones principales")
    add_table(
        doc,
        ["Módulo", "Prefijo", "Función"],
        [
            ("Ingesta y preparación de datos", "RF-ETL", "Descargar, normalizar, tipar, filtrar territorialmente y persistir las fuentes oficiales."),
            ("Calidad y auditoría de datos", "RF-CAL", "Validar esquemas, integridad, consistencia y cobertura; documentar anomalías y subregistro."),
            ("Proyección de demanda", "RF-DEM", "Construir series semanales, entrenar, evaluar y publicar pronósticos con intervalos."),
            ("Accesibilidad territorial", "RF-ACC", "Georreferenciar, calcular isócronas, cobertura poblacional y brechas comunales."),
            ("Hospitalización psiquiátrica", "RF-HOS", "Caracterizar la duración de estadía F00–F99 y sus factores asociados."),
            ("Visualización y exploración", "RF-VIS", "Entregar paneles, mapas, series, comparadores y exportaciones a los usuarios."),
            ("Servicios de datos", "RF-API", "Exponer los resultados analíticos mediante una API REST documentada."),
            ("Administración y operación", "RF-ADM", "Ejecutar, monitorear, versionar y auditar el ciclo de actualización de datos."),
            ("Trazabilidad y documentación", "RF-TRZ", "Registrar linaje, metadatos, supuestos y limitaciones de cada resultado publicado."),
        ],
        [Cm(5.2), Cm(2.3), Cm(10.2)],
        font_size=9.5,
    )

    h2(doc, "2.3 Actores y usuarios")
    add_table(
        doc,
        ["ID", "Actor", "Descripción", "Necesidad principal", "Uso"],
        [
            (
                "AC-01",
                "Gestor de red / Planificador sanitario",
                "Profesional de un Servicio de Salud, municipio o corporación que planifica recursos "
                "de la red de urgencia.",
                "Saber dónde y cuándo aumentará la presión y qué territorios quedan peor cubiertos.",
                "Semanal",
            ),
            (
                "AC-02",
                "Referente comunal de salud mental",
                "Profesional responsable del programa de salud mental a nivel comunal o de CESFAM.",
                "Comparar su comuna con el resto de la RM y sustentar solicitudes de refuerzo.",
                "Mensual",
            ),
            (
                "AC-03",
                "Analista / Epidemiólogo de servicio",
                "Perfil técnico que requiere el dato desagregado y su metodología.",
                "Descargar series, revisar supuestos, métricas y limitaciones.",
                "Semanal",
            ),
            (
                "AC-04",
                "Jefatura de urgencia",
                "Responsable operativo de una unidad de urgencia.",
                "Ver la tendencia proyectada de su establecimiento y su entorno.",
                "Mensual",
            ),
            (
                "AC-05",
                "Administrador de datos",
                "Integrante del equipo SAAD a cargo del pipeline y de las publicaciones.",
                "Ejecutar, auditar y versionar las actualizaciones de datos y modelos.",
                "Por publicación",
            ),
            (
                "AC-06",
                "Docente / Comisión evaluadora",
                "Evalúa el cumplimiento del alcance comprometido.",
                "Verificar reproducibilidad, evidencias y trazabilidad.",
                "Por hito",
            ),
            (
                "AC-07",
                "Fuentes de datos oficiales",
                "Sistema externo: DEIS/MINSAL, INE/Censo, OSM y GTFS.",
                "Entregar archivos y servicios consumidos por lotes.",
                "Según publicación",
            ),
        ],
        [Cm(1.5), Cm(3.3), Cm(5.2), Cm(5.2), Cm(2.5)],
        font_size=9,
    )

    h2(doc, "2.4 Restricciones generales")
    add_table(
        doc,
        ["ID", "Restricción", "Origen"],
        [
            ("RG-01", "El territorio de análisis es exclusivamente la Región Metropolitana (52 comunas).", "BIT-01"),
            ("RG-02", "El foco es la red pública de urgencia (SAPU, SAR, SUR y UEH); el sector privado queda fuera.", "BIT-01"),
            ("RG-03", "Solo se procesan datos agregados o disociados; no se tratan datos personales identificables.", "BIT-01, REG-01"),
            ("RG-04", "El vínculo entre urgencias y egresos hospitalarios es exclusivamente ecológico.", "BIT-01, EDA-07"),
            ("RG-05", "La serie de urgencias de salud mental es comparable desde 2021; 2020 presenta discontinuidad de captura documentada.", "EDA-08"),
            ("RG-06", "Los tramos etarios de la base de urgencias no permiten aislar con precisión población infanto-adolescente.", "BIT-01"),
            ("RG-07", "El sistema debe ser reproducible de punta a punta mediante código versionado.", "REG-01"),
            ("RG-08", "El plazo de construcción está acotado al calendario académico (entrega final: 26-11-2026).", "DEF-01"),
        ],
        [Cm(1.7), Cm(12.5), Cm(3.5)],
        font_size=9.5,
        first_col_bold=True,
    )

    h2(doc, "2.5 Supuestos y dependencias")
    add_table(
        doc,
        ["ID", "Supuesto o dependencia", "Impacto si no se cumple"],
        [
            ("SD-01", "El DEIS mantiene la publicación anual de urgencias y egresos con esquema estable.", "Requiere re-mapeo de columnas y nueva auditoría de codificación."),
            ("SD-02", "Los egresos hospitalarios 2026 no se encuentran publicados al cierre del proyecto.", "El módulo de hospitalización cubre el período 2020–2025."),
            ("SD-03", "El feed GTFS del transporte público de la RM es accesible y vigente.", "Las isócronas de transporte público se degradan a alcance por red vial (RF-ACC-003 pasa a *Podría*)."),
            ("SD-04", "Se confirma un cliente piloto disponible para validación durante el semestre.", "La validación de utilidad se realiza con la experta de dominio y la docente."),
            ("SD-05", "El equipo dispone de capacidad de cómputo suficiente para volúmenes del orden de decenas de millones de registros.", "Se aplica procesamiento por particiones y muestreo documentado."),
        ],
        [Cm(1.7), Cm(8.5), Cm(7.5)],
        font_size=9.5,
        first_col_bold=True,
    )
    doc.add_page_break()


def build_conventions(doc) -> None:
    h1(doc, "3. Convenciones de especificación")

    h2(doc, "3.1 Identificación")
    p(
        doc,
        "Cada requisito se identifica como **RF-<MÓDULO>-<NNN>**, donde `<MÓDULO>` corresponde al "
        "prefijo definido en la sección 2.2 y `<NNN>` a un correlativo de tres dígitos. Los "
        "identificadores **no se reutilizan**: un requisito eliminado conserva su identificador "
        "con estado *Retirado* y la referencia a la decisión que lo retiró.",
    )

    h2(doc, "3.2 Prioridad (MoSCoW)")
    add_table(
        doc,
        ["Código", "Significado", "Compromiso"],
        [
            ("Debe", "Must have", "Obligatorio para el MVP. Su ausencia invalida la entrega."),
            ("Debería", "Should have", "Alto valor; se implementa salvo restricción de tiempo justificada."),
            ("Podría", "Could have", "Deseable; se implementa si el avance del proyecto lo permite."),
            ("No", "Won't have (this time)", "Explícitamente fuera del alcance actual (sección 8)."),
        ],
        [Cm(2.5), Cm(4.5), Cm(10.7)],
        font_size=9.5,
        first_col_bold=True,
    )

    h2(doc, "3.3 Método de verificación")
    p(
        doc,
        "Conforme a ISO/IEC/IEEE 29148:2018, cada requisito declara su método de verificación en la "
        "columna «Prio. / Ver.» del catálogo:",
    )
    add_table(
        doc,
        ["Código", "Método", "Descripción"],
        [
            ("I", "Inspección", "Revisión de código, configuración, documentación o salida generada."),
            ("A", "Análisis", "Verificación mediante cálculo, modelo o razonamiento sobre evidencia."),
            ("D", "Demostración", "Ejecución observada del sistema ante un evaluador."),
            ("P", "Prueba", "Ejecución de casos de prueba automatizados con resultado esperado."),
        ],
        [Cm(2.0), Cm(3.4), Cm(12.3)],
        font_size=9.5,
        first_col_bold=True,
    )

    h2(doc, "3.4 Reglas de redacción")
    bullets(
        doc,
        [
            "Los requisitos se redactan en forma «El sistema debe …», en voz activa y en presente.",
            "Cada requisito expresa **una sola** capacidad, verificable e inequívoca.",
            "Se evitan términos no medibles («rápido», «amigable», «óptimo»); esos atributos se "
            "cuantifican como requisitos no funcionales en el ERNF.",
            "Cada requisito declara un criterio de aceptación observable y su método de verificación.",
        ],
    )

    h2(doc, "3.5 Estados del requisito")
    p(
        doc,
        "`Propuesto` · `Aprobado` · `Implementado` · `Verificado` · `Retirado`. Al cierre de esta "
        "versión, todos los requisitos se encuentran en estado `Propuesto`, salvo los indicados "
        "como **Implementado** en la sección 9.2, correspondientes al avance efectivo de la Fase 1.",
    )
    doc.add_page_break()


# --------------------------------------------------------------------------------------
# Catálogo de requisitos funcionales
# --------------------------------------------------------------------------------------
RF_ETL = [
    ("RF-ETL-001", "Registro de fuentes oficiales",
     "El sistema debe permitir la descarga o incorporación de los archivos oficiales DEIS (atenciones de urgencia, egresos hospitalarios y maestro de establecimientos), registrando origen, fecha de obtención y tamaño de cada archivo.",
     "Existe un registro consultable con URL o ruta de origen, fecha de descarga y verificación de integridad para cada archivo incorporado.",
     "Debe / I"),
    ("RF-ETL-002", "Normalización de codificación y delimitador",
     "El sistema debe detectar la codificación de cada archivo de origen (Latin-1/CP1252 o UTF-8) y su delimitador, y convertir el contenido a UTF-8 sin pérdida ni sustitución silenciosa de caracteres.",
     "La conversión no introduce caracteres de reemplazo nuevos; los caracteres de reemplazo preexistentes de origen quedan reportados y no ocultos.",
     "Debe / P"),
    ("RF-ETL-003", "Homologación de esquema y tipos",
     "El sistema debe transformar los nombres de columna a snake_case y asignar tipos de dato explícitos y homogéneos entre todos los años de una misma fuente.",
     "El esquema resultante es idéntico para todos los años de la fuente; una prueba automatizada falla ante cualquier divergencia de nombre o tipo.",
     "Debe / P"),
    ("RF-ETL-004", "Filtrado territorial de la Región Metropolitana",
     "El sistema debe filtrar los registros correspondientes a la RM mediante el catálogo oficial DEIS de establecimientos, resolviendo el cruce de identificadores de forma unívoca.",
     "El 100 % de los registros retenidos posee correspondencia 1:1 con un establecimiento del maestro RM; los no resueltos se reportan con su causa.",
     "Debe / P"),
    ("RF-ETL-005", "Persistencia analítica particionada",
     "El sistema debe persistir los datos normalizados en formato columnar (Parquet), particionado por año.",
     "Se generan archivos por año legibles con el esquema declarado y sin pérdida de filas respecto del conteo posterior al filtrado.",
     "Debe / I"),
    ("RF-ETL-006", "Ingesta de cartografía comunal",
     "El sistema debe incorporar la cartografía comunal oficial del Censo 2024 para la RM, preservando el sistema de referencia de coordenadas de origen.",
     "La capa comunal contiene las 52 comunas de la RM, sin nulos ni duplicados en el código territorial y con geometrías válidas.",
     "Debe / P"),
    ("RF-ETL-007", "Ingesta de red vial y transporte público",
     "El sistema debe incorporar la red vial de OpenStreetMap y el feed GTFS del transporte público de la RM, registrando la fecha de extracción de cada insumo.",
     "Ambos insumos quedan almacenados, versionados por fecha y disponibles para el módulo de accesibilidad.",
     "Debe / I"),
    ("RF-ETL-008", "Carga en base de datos espacial",
     "El sistema debe cargar las tablas analíticas y las capas espaciales en PostgreSQL/PostGIS, conservando índices espaciales y claves de negocio.",
     "Las consultas espaciales y por comuna se resuelven desde la base de datos; existen índices sobre las columnas de cruce y de geometría.",
     "Debe / D"),
    ("RF-ETL-009", "Ejecución reproducible e idempotente",
     "El sistema debe permitir re-ejecutar el pipeline completo sobre las mismas fuentes obteniendo resultados idénticos, sin duplicar registros ni requerir pasos manuales.",
     "Dos ejecuciones consecutivas sobre el mismo insumo producen salidas equivalentes en conteo, esquema y agregados de control.",
     "Debe / P"),
    ("RF-ETL-010", "Agregación semanal analítica",
     "El sistema debe generar tablas agregadas por semana epidemiológica, establecimiento, comuna y causa, listas para el modelamiento y la visualización.",
     "La suma de los agregados coincide exactamente con el detalle de origen para cada combinación año-causa.",
     "Debe / P"),
    ("RF-ETL-011", "Catálogo dimensional de causas",
     "El sistema debe mantener un catálogo de causas de urgencia que identifique el macro-agregador de salud mental (F00–F99) y sus subcausas, homologado con el diccionario oficial DEIS.",
     "El catálogo declara para cada identificador su glosa, su equivalencia CIE-10 y su relación de anidamiento, y está versionado en el repositorio.",
     "Debe / I"),
    ("RF-ETL-012", "Normalización de egresos hospitalarios",
     "El sistema debe normalizar la base de egresos hospitalarios F00–F99 (2020–2025) aplicando las mismas reglas de codificación, tipado y homologación de esquema.",
     "El resultado supera los mismos controles de integridad definidos para urgencias y queda documentado en un informe reproducible.",
     "Debe / P"),
    ("RF-ETL-013", "Incorporación de indicadores complementarios",
     "El sistema debería incorporar la población comunal por grupo etario y un indicador de vulnerabilidad socioeconómica comunal, como denominadores del análisis.",
     "Cada comuna cuenta con población e indicador incorporados; se documenta fuente, año y unidad de cada variable.",
     "Debería / I"),
    ("RF-ETL-014", "Ejecución parcial por fuente y período",
     "El sistema debería permitir ejecutar el pipeline restringido a una fuente y a un rango de años, sin reprocesar el resto.",
     "Es posible actualizar un solo año y verificar que los años restantes permanecen sin modificación.",
     "Debería / D"),
]

RF_CAL = [
    ("RF-CAL-001", "Validación de esquema",
     "El sistema debe validar, antes de publicar cualquier salida, que columnas, tipos y dominios coincidan con el esquema declarado, deteniendo la publicación ante una divergencia.",
     "Ante una columna faltante, un tipo inesperado o un valor fuera de dominio, la ejecución se detiene e informa el elemento infractor.",
     "Debe / P"),
    ("RF-CAL-002", "Controles de integridad",
     "El sistema debe verificar ausencia de nulos en identificadores, ausencia de valores negativos en conteos, validez de fechas y semanas, y ausencia de duplicados exactos y de llave natural.",
     "El informe de calidad reporta cada control con su conteo asociado; cualquier incumplimiento bloquea la promoción del dato.",
     "Debe / P"),
    ("RF-CAL-003", "Control de doble conteo",
     "El sistema debe calcular el total de atenciones y el total de salud mental con los identificadores de causa correctos, y verificar la igualdad entre el macro-agregador de salud mental y la suma de sus subcausas.",
     "La verificación de anidamiento arroja diferencia cero; una discrepancia distinta de cero genera alerta y bloquea la publicación.",
     "Debe / P"),
    ("RF-CAL-004", "Consistencia de desagregación etaria",
     "El sistema debe verificar que el total de atenciones de cada registro coincida con la suma de sus tramos de edad.",
     "El control se ejecuta sobre el 100 % de los registros y reporta cero discrepancias o bien las enumera.",
     "Debe / P"),
    ("RF-CAL-005", "Detección de subregistro y discontinuidades",
     "El sistema debe identificar y documentar períodos o establecimientos con captura anómala respecto de su comportamiento histórico, antes de habilitar comparaciones territoriales.",
     "Cada caso marcado registra la regla de detección y el valor observado; la interfaz advierte al usuario al mostrarlo.",
     "Debe / A"),
    ("RF-CAL-006", "Informe de calidad por ejecución",
     "El sistema debe generar en cada ejecución un informe reproducible con los controles aplicados, sus resultados, las limitaciones vigentes y un veredicto explícito de aptitud del dato.",
     "El informe se genera automáticamente, es legible sin ejecutar código y declara qué controles se realizaron y qué limitaciones permanecen.",
     "Debe / I"),
    ("RF-CAL-007", "Cobertura geoespacial",
     "El sistema debe identificar los establecimientos sin coordenadas válidas, excluirlos de los cálculos espaciales e informar explícitamente esa exclusión en las vistas afectadas.",
     "Los establecimientos sin coordenadas aparecen listados en el informe de cobertura y las vistas espaciales muestran la advertencia con su conteo.",
     "Debe / D"),
    ("RF-CAL-008", "Diccionario de datos versionado",
     "El sistema debe mantener un diccionario de datos que declare, para cada variable publicada, su definición, tipo, dominio, granularidad, fuente y transformaciones aplicadas.",
     "El diccionario está versionado en el repositorio y cubre el 100 % de las variables expuestas por la API y la interfaz.",
     "Debe / I"),
    ("RF-CAL-009", "Registro de anomalías sin corrección silenciosa",
     "El sistema no debe modificar ni descartar valores considerados anómalos sin registrar el valor original, la evidencia de la anomalía, la regla de detección y la justificación de la transformación.",
     "Toda transformación no trivial cuenta con su registro; la revisión de código verifica la ausencia de correcciones no documentadas.",
     "Debe / I"),
    ("RF-CAL-010", "Comparación entre versiones de datos",
     "El sistema debería comparar cada nueva carga con la publicación anterior, reportando variaciones de volumen que superen un umbral configurable.",
     "Ante una variación superior al umbral, la ejecución emite alerta y requiere confirmación del administrador antes de publicar.",
     "Debería / P"),
]

RF_DEM = [
    ("RF-DEM-001", "Construcción de la serie de entrenamiento",
     "El sistema debe construir la serie semanal de atenciones de urgencia por salud mental (F00–F99) para el período comparable acordado, reservando el período más reciente como conjunto de evaluación fuera de muestra.",
     "La serie de entrenamiento excluye el período declarado no comparable y el conjunto de evaluación no participa del ajuste del modelo.",
     "Debe / I"),
    ("RF-DEM-002", "Generación de pronóstico",
     "El sistema debe generar proyecciones de demanda con un horizonte configurable entre 4 y 8 semanas.",
     "Para un horizonte dentro del rango, el sistema entrega una proyección por cada semana futura y por cada unidad de análisis habilitada.",
     "Debe / P"),
    ("RF-DEM-003", "Intervalos de predicción",
     "El sistema debe acompañar cada valor proyectado de un intervalo de predicción con nivel de confianza declarado.",
     "Toda proyección expuesta por la API o la interfaz incluye límite inferior, valor central, límite superior y nivel de confianza.",
     "Debe / P"),
    ("RF-DEM-004", "Baseline estacional obligatorio",
     "El sistema debe calcular, para el mismo horizonte y las mismas unidades, un pronóstico de referencia mediante un baseline estacional ingenuo.",
     "El baseline se calcula y almacena en cada ciclo de evaluación y es consultable junto al modelo propuesto.",
     "Debe / P"),
    ("RF-DEM-005", "Backtesting con origen deslizante",
     "El sistema debe evaluar el modelo mediante backtesting con origen deslizante sobre el histórico, reportando métricas de error por horizonte de pronóstico.",
     "Se reportan al menos MAE, RMSE y una métrica escalada respecto del baseline, desagregadas por horizonte.",
     "Debe / P"),
    ("RF-DEM-006", "Registro de experimentos",
     "El sistema debe registrar, para cada entrenamiento, la versión del modelo, los hiperparámetros, el rango de datos utilizado, las métricas obtenidas y la fecha de ejecución.",
     "El registro permite reproducir cualquier resultado publicado a partir de su identificador de versión.",
     "Debe / I"),
    ("RF-DEM-007", "Criterio explícito de promoción",
     "El sistema debe promover a producción únicamente modelos que superen al baseline según el criterio de selección documentado.",
     "Un modelo que no supera al baseline no se publica; la decisión y sus métricas quedan registradas.",
     "Debe / A"),
    ("RF-DEM-008", "Señalización de baja confianza",
     "El sistema debe marcar como de baja confianza, o abstenerse de proyectar, las unidades cuya serie no alcance el mínimo de observaciones o presente discontinuidades documentadas.",
     "Las unidades afectadas se muestran con la marca correspondiente y su explicación, y no se presentan como proyecciones equivalentes al resto.",
     "Debe / D"),
    ("RF-DEM-009", "Agregación territorial del pronóstico",
     "El sistema debe permitir consultar el pronóstico agregado por comuna, por servicio de salud y por tipo de establecimiento.",
     "La suma de los pronósticos de las unidades que componen un agregado es consistente con el valor mostrado para ese agregado.",
     "Debe / P"),
    ("RF-DEM-010", "Reentrenamiento programado",
     "El sistema debería permitir programar el reentrenamiento del modelo tras cada nueva publicación oficial de datos.",
     "Es posible configurar y ejecutar el reentrenamiento sin intervención manual sobre el código.",
     "Debería / D"),
    ("RF-DEM-011", "Desagregación por subcausa",
     "El sistema podría generar proyecciones desagregadas por subcausa de salud mental cuando el volumen de la serie lo permita.",
     "Las subcausas proyectadas cumplen el mínimo de observaciones definido y su suma es coherente con el agregado.",
     "Podría / A"),
    ("RF-DEM-012", "Comparación de escenarios",
     "El sistema podría permitir comparar el pronóstico vigente con el de la publicación anterior, para observar la revisión de la tendencia.",
     "La interfaz muestra ambas curvas identificadas por versión y fecha de generación.",
     "Podría / D"),
]

RF_ACC = [
    ("RF-ACC-001", "Georreferenciación validada",
     "El sistema debe georreferenciar los establecimientos de urgencia de la RM y validar que cada coordenada se ubique dentro del polígono regional.",
     "Las coordenadas fuera del polígono o ausentes quedan excluidas del cálculo espacial y listadas en el informe de cobertura.",
     "Debe / P"),
    ("RF-ACC-002", "Isócronas por red vial",
     "El sistema debe calcular isócronas de tiempo de viaje sobre la red vial para umbrales configurables (por defecto 15, 30 y 45 minutos) desde cada establecimiento de urgencia.",
     "Se genera un polígono por establecimiento y umbral, con geometría válida y sistema de referencia declarado.",
     "Debe / P"),
    ("RF-ACC-003", "Isócronas por transporte público",
     "El sistema debe calcular isócronas basadas en transporte público a partir del feed GTFS, para una franja horaria configurable.",
     "Se genera el conjunto de isócronas para al menos una franja de referencia, documentando hora, día tipo y versión del feed.",
     "Debe / P"),
    ("RF-ACC-004", "Cobertura poblacional",
     "El sistema debe estimar la población comunal cubierta y no cubierta por cada umbral de tiempo de viaje.",
     "Para cada comuna se reporta población total, cubierta, no cubierta y porcentaje, con la metodología de asignación documentada.",
     "Debe / A"),
    ("RF-ACC-005", "Cruce de brechas",
     "El sistema debe cruzar accesibilidad, demanda proyectada y vulnerabilidad socioeconómica para priorizar comunas con alta presión esperada y baja cobertura.",
     "Se produce un ranking comunal reproducible cuya fórmula, ponderadores y supuestos son documentados y consultables desde la interfaz.",
     "Debe / A"),
    ("RF-ACC-006", "Visualización cartográfica",
     "El sistema debe representar en un mapa interactivo las capas de establecimientos, isócronas, demanda comunal y brechas priorizadas, con activación independiente por capa.",
     "El usuario activa y desactiva cada capa y consulta el detalle de un elemento al seleccionarlo en el mapa.",
     "Debe / D"),
    ("RF-ACC-007", "Exportación de capas",
     "El sistema debería permitir exportar las capas resultantes en formatos geoespaciales estándar.",
     "La capa exportada se abre correctamente en software SIG de terceros, conservando atributos y sistema de referencia.",
     "Debería / D"),
    ("RF-ACC-008", "Documentación de supuestos del cálculo",
     "El sistema debe exponer, junto a cada resultado de accesibilidad, los supuestos aplicados: modo de transporte, franja horaria, fecha de la red vial y del feed GTFS.",
     "Toda vista o exportación de accesibilidad incluye su ficha de supuestos.",
     "Debe / I"),
    ("RF-ACC-009", "Escenario de cierre o refuerzo",
     "El sistema podría permitir simular el efecto sobre la cobertura de habilitar o suspender un establecimiento.",
     "Al excluir o incluir un establecimiento, el sistema recalcula y muestra la variación de cobertura comunal.",
     "Podría / D"),
]

RF_HOS = [
    ("RF-HOS-001", "Selección de la cohorte de análisis",
     "El sistema debe seleccionar los egresos hospitalarios con diagnóstico principal en el capítulo F00–F99 para el período disponible.",
     "El conteo de la cohorte es reproducible desde código y su criterio de selección está documentado.",
     "Debe / P"),
    ("RF-HOS-002", "Cálculo de duración de estadía",
     "El sistema debe calcular la duración de estadía de cada egreso y su distribución (tendencia central, dispersión y percentiles).",
     "Los estadísticos se calculan sobre la cohorte declarada, informando el tratamiento de valores extremos y faltantes.",
     "Debe / P"),
    ("RF-HOS-003", "Segmentación de resultados",
     "El sistema debe permitir segmentar la duración de estadía por grupo diagnóstico, grupo etario, sexo, previsión, pertenencia al SNSS y condición de egreso.",
     "Cada segmentación entrega conteo, medidas resumen y advertencia cuando el tamaño del subgrupo es insuficiente.",
     "Debe / D"),
    ("RF-HOS-004", "Análisis de factores asociados",
     "El sistema debe estimar la asociación entre las variables disponibles y la duración de estadía, reportando la magnitud del efecto y su incertidumbre.",
     "El resultado incluye estimadores, medida de incertidumbre y la declaración explícita de que la relación es asociativa y no causal.",
     "Debe / A"),
    ("RF-HOS-005", "Cruce ecológico controlado",
     "El sistema debe permitir comparar hospitalización y urgencias únicamente a nivel de agregado territorial y temporal, impidiendo cualquier enlace individual o por establecimiento.",
     "No existe función, endpoint ni consulta que devuelva un enlace entre un registro de egreso y uno de urgencia; la restricción se verifica por revisión de código.",
     "Debe / I"),
    ("RF-HOS-006", "Advertencias de interpretación",
     "El sistema debe mostrar, junto a los resultados de hospitalización, las advertencias sobre naturaleza ecológica, ausencia de causalidad y limitaciones de la fuente.",
     "Toda vista y exportación del módulo incluye el bloque de advertencias correspondiente.",
     "Debe / I"),
    ("RF-HOS-007", "Contexto con indicadores oficiales",
     "El sistema podría complementar la interpretación con indicadores oficiales de salud mental desagregados por sexo provenientes de fuentes estadísticas públicas.",
     "Los indicadores incorporados citan fuente, año y unidad, y se presentan claramente separados de los resultados propios.",
     "Podría / I"),
]

RF_VIS = [
    ("RF-VIS-001", "Panel general de indicadores",
     "El sistema debe presentar un panel inicial con los indicadores clave: volumen acumulado de atenciones, participación de salud mental, variación interanual, horizonte proyectado y fecha de corte.",
     "Cada indicador muestra su valor, su período de referencia y el acceso a su definición.",
     "Debe / D"),
    ("RF-VIS-002", "Filtros transversales",
     "El sistema debe ofrecer filtros por período, comuna, servicio de salud, tipo de establecimiento y causa, aplicables de forma consistente a todas las vistas.",
     "Al cambiar un filtro, todas las vistas activas se actualizan de forma coherente y el estado del filtro permanece visible.",
     "Debe / D"),
    ("RF-VIS-003", "Serie temporal con pronóstico",
     "El sistema debe graficar la serie observada junto a la proyección y su banda de intervalo de predicción, diferenciando visualmente el tramo observado del proyectado.",
     "El gráfico distingue ambos tramos, muestra la banda y permite consultar el valor puntual de cada semana.",
     "Debe / D"),
    ("RF-VIS-004", "Mapa interactivo",
     "El sistema debe ofrecer un mapa interactivo de la RM con navegación, selección de comuna o establecimiento y control de capas.",
     "Al seleccionar un elemento, el sistema muestra su ficha con indicadores y metadatos asociados.",
     "Debe / D"),
    ("RF-VIS-005", "Vista de brechas territoriales",
     "El sistema debe presentar una vista que combine demanda proyectada y accesibilidad para identificar comunas prioritarias.",
     "La vista ordena las comunas según el criterio de brecha y explica el cálculo aplicado.",
     "Debe / D"),
    ("RF-VIS-006", "Vista de hospitalización",
     "El sistema debe presentar la caracterización de la duración de estadía con sus segmentaciones y advertencias.",
     "La vista permite seleccionar la segmentación y muestra el tamaño de cada subgrupo.",
     "Debe / D"),
    ("RF-VIS-007", "Rankings y comparaciones",
     "El sistema debe entregar rankings de comunas y establecimientos por volumen, tasa y variación, con criterio de orden configurable.",
     "El usuario cambia el criterio de ordenamiento y el resultado se actualiza sin recargar la vista.",
     "Debe / D"),
    ("RF-VIS-008", "Estado y advertencias de datos",
     "El sistema debe mostrar de forma permanente la fecha de corte, la versión de los datos y las advertencias de calidad aplicables a la vista en uso.",
     "Ninguna vista analítica se presenta sin su indicador de vigencia y sus advertencias asociadas.",
     "Debe / I"),
    ("RF-VIS-009", "Comparador de comunas",
     "El sistema debería permitir comparar simultáneamente entre dos y cinco comunas sobre los mismos indicadores.",
     "La comparación muestra las series o valores seleccionados en una misma escala, identificados por color y etiqueta.",
     "Debería / D"),
    ("RF-VIS-010", "Ayuda contextual y glosario",
     "El sistema debe ofrecer la definición de cada indicador y su método de cálculo desde la propia vista.",
     "Cada indicador dispone de acceso directo a su definición sin abandonar la pantalla.",
     "Debe / D"),
    ("RF-VIS-011", "Descarga de datos de la vista",
     "El sistema debe permitir descargar en formato tabular los datos que sustentan la vista consultada, respetando los filtros aplicados.",
     "El archivo descargado reproduce exactamente los valores mostrados e incluye cabecera con filtros, fecha de corte y versión.",
     "Debe / P"),
    ("RF-VIS-012", "Enlace compartible del estado de vista",
     "El sistema debería permitir compartir un enlace que reproduzca la vista y los filtros aplicados.",
     "Al abrir el enlace en otra sesión se restituye la misma vista con los mismos filtros.",
     "Debería / D"),
    ("RF-VIS-013", "Ficha ejecutiva por comuna",
     "El sistema podría generar una ficha exportable por comuna con sus principales indicadores, proyección, cobertura y advertencias.",
     "La ficha se genera para cualquier comuna seleccionada y contiene fecha de emisión y versión de datos.",
     "Podría / D"),
]

RF_API = [
    ("RF-API-001", "API REST documentada",
     "El sistema debe exponer una API REST cuya especificación esté publicada y sea navegable.",
     "La especificación describe el 100 % de los endpoints disponibles, con parámetros, respuestas y ejemplos.",
     "Debe / I"),
    ("RF-API-002", "Servicios de demanda",
     "El sistema debe exponer endpoints para consultar la serie observada y la proyección, filtrables por período, unidad territorial y causa.",
     "Las respuestas incluyen valores, intervalos, versión del modelo y fecha de corte.",
     "Debe / P"),
    ("RF-API-003", "Servicios geoespaciales",
     "El sistema debe exponer las capas territoriales en un formato geoespacial estándar consumible por la interfaz y por clientes externos.",
     "Las capas se entregan con geometría válida, atributos y sistema de referencia declarado.",
     "Debe / P"),
    ("RF-API-004", "Servicios de hospitalización",
     "El sistema debe exponer endpoints con los indicadores agregados de duración de estadía y sus segmentaciones.",
     "Las respuestas entregan exclusivamente agregados, sin registros individuales.",
     "Debe / P"),
    ("RF-API-005", "Servicio de metadatos",
     "El sistema debe exponer un endpoint con la versión de los datos, la fecha de corte, la versión del modelo vigente y el estado de los controles de calidad.",
     "La interfaz consume este servicio para mostrar el estado de datos exigido por RF-VIS-008.",
     "Debe / P"),
    ("RF-API-006", "Manejo normalizado de errores",
     "El sistema debe responder los errores con un formato uniforme que incluya código, mensaje y detalle accionable, sin exponer trazas internas.",
     "Ante parámetros inválidos o recursos inexistentes, la respuesta usa el código HTTP correcto y el cuerpo normalizado.",
     "Debe / P"),
    ("RF-API-007", "Paginación y límites",
     "El sistema debe paginar las respuestas de colecciones extensas y declarar el límite aplicado.",
     "Ninguna respuesta excede el límite configurado; la paginación permite recorrer el conjunto completo sin omisiones ni repeticiones.",
     "Debe / P"),
    ("RF-API-008", "Versionado de la interfaz",
     "El sistema debe versionar la API en su ruta base, de modo que un cambio incompatible no afecte a los consumidores existentes.",
     "Las rutas incluyen el identificador de versión y los cambios incompatibles se publican bajo una versión nueva.",
     "Debe / I"),
]

RF_ADM = [
    ("RF-ADM-001", "Autenticación de funciones administrativas",
     "El sistema debe requerir autenticación para ejecutar el pipeline, promover modelos o publicar versiones de datos.",
     "Sin credenciales válidas, las funciones administrativas no son accesibles desde la interfaz ni desde la API.",
     "Debe / P"),
    ("RF-ADM-002", "Perfiles de acceso",
     "El sistema debe distinguir al menos los perfiles de consulta y de administración, con permisos diferenciados.",
     "Un usuario de consulta no puede ejecutar operaciones de administración; el intento queda registrado.",
     "Debe / P"),
    ("RF-ADM-003", "Ejecución manual y programada",
     "El sistema debe permitir ejecutar el ciclo de actualización de datos de forma manual y programada.",
     "Ambas modalidades registran inicio, término, resultado y volumen procesado.",
     "Debe / D"),
    ("RF-ADM-004", "Bitácora de ejecuciones",
     "El sistema debe mantener una bitácora consultable de ejecuciones, con resultado, duración, errores y controles de calidad aplicados.",
     "La bitácora permite identificar la causa de una ejecución fallida sin acceder al servidor.",
     "Debe / D"),
    ("RF-ADM-005", "Versionado de publicaciones",
     "El sistema debe versionar cada publicación de datos y modelos, permitiendo identificar qué versión sustenta cada resultado mostrado.",
     "Todo resultado expuesto puede asociarse a una versión de datos y de modelo identificable.",
     "Debe / I"),
    ("RF-ADM-006", "Reversión a versión anterior",
     "El sistema debería permitir revertir la publicación vigente a la versión inmediatamente anterior.",
     "La reversión restituye datos y modelo previos y queda registrada en la bitácora.",
     "Debería / D"),
    ("RF-ADM-007", "Verificación de disponibilidad",
     "El sistema debe exponer un punto de verificación de estado que informe la disponibilidad de la API y de la base de datos.",
     "La verificación responde con el estado de cada dependencia crítica.",
     "Debe / P"),
]

RF_TRZ = [
    ("RF-TRZ-001", "Linaje de datos",
     "El sistema debe registrar, para cada tabla publicada, la fuente de origen, las transformaciones aplicadas y la ejecución que la generó.",
     "Es posible reconstruir el camino completo desde un valor publicado hasta el archivo de origen que lo sustenta.",
     "Debe / I"),
    ("RF-TRZ-002", "Reproducibilidad de cifras",
     "Toda cifra publicada por el sistema debe ser reproducible mediante la ejecución del código versionado sobre las fuentes declaradas.",
     "Una re-ejecución independiente reproduce las cifras del informe o de la vista sin ajustes manuales.",
     "Debe / P"),
    ("RF-TRZ-003", "Declaración de supuestos y limitaciones",
     "El sistema debe publicar, junto a cada componente analítico, sus supuestos, limitaciones y el alcance de interpretación válido.",
     "Los tres módulos analíticos cuentan con su declaración vigente y fechada.",
     "Debe / I"),
    ("RF-TRZ-004", "Distinción entre hecho, inferencia e hipótesis",
     "El sistema debe distinguir explícitamente, en informes y vistas, entre valores observados, inferencias derivadas e hipótesis por validar.",
     "Ninguna inferencia se presenta con el mismo estatus visual o textual que un valor observado.",
     "Debe / I"),
    ("RF-TRZ-005", "Documentación de uso",
     "El sistema debe contar con documentación de instalación, despliegue, uso de la API y manual de usuario.",
     "La documentación permite a una persona ajena al equipo levantar el entorno y operar la plataforma.",
     "Debe / D"),
]

MODULES = [
    ("4.1 Módulo de ingesta y preparación de datos (RF-ETL)",
     "Cubre la incorporación de las fuentes oficiales, su normalización técnica y su preparación "
     "analítica. Es la base sobre la que operan los tres módulos analíticos.", RF_ETL),
    ("4.2 Módulo de calidad y auditoría de datos (RF-CAL)",
     "Materializa la exigencia de auditar calidad y subregistro antes de comparar territorios o "
     "entrenar modelos. Ningún dato se publica sin superar estos controles.", RF_CAL),
    ("4.3 Módulo de proyección de demanda (RF-DEM)",
     "Componente 1 del alcance acordado: anticipar la demanda de urgencia por salud mental en un "
     "horizonte de 4 a 8 semanas, con evaluación explícita frente a un baseline.", RF_DEM),
    ("4.4 Módulo de accesibilidad territorial (RF-ACC)",
     "Componente 2 del alcance acordado: territorializar la cobertura efectiva de la red pública "
     "de urgencia y las brechas resultantes.", RF_ACC),
    ("4.5 Módulo de hospitalización psiquiátrica (RF-HOS)",
     "Componente 3 del alcance acordado: caracterizar la duración de estadía en egresos F00–F99 y "
     "los factores asociados, siempre bajo restricción de cruce ecológico.", RF_HOS),
    ("4.6 Módulo de visualización y exploración (RF-VIS)",
     "Interfaz de la plataforma: traduce los resultados analíticos en vistas comprensibles, "
     "acompañadas de sus advertencias y de la posibilidad de exportar la evidencia.", RF_VIS),
    ("4.7 Módulo de servicios de datos (RF-API)",
     "Capa de servicios que desacopla la analítica de la interfaz y habilita el consumo por parte "
     "de terceros.", RF_API),
    ("4.8 Módulo de administración y operación (RF-ADM)",
     "Funciones internas de operación del ciclo de actualización, control de acceso y "
     "versionamiento de publicaciones.", RF_ADM),
    ("4.9 Módulo de trazabilidad y documentación (RF-TRZ)",
     "Garantiza que todo resultado sea atribuible, reproducible y correctamente interpretado, "
     "conforme a las reglas de rigor analítico del proyecto.", RF_TRZ),
]


def build_requirements(doc) -> None:
    h1(doc, "4. Requisitos funcionales")
    p(
        doc,
        "El catálogo se organiza en nueve módulos. La columna «Prio. / Ver.» indica la prioridad "
        "MoSCoW (sección 3.2) y el método de verificación (sección 3.3).",
    )
    for title, description, rows in MODULES:
        h2(doc, title)
        p(doc, description)
        add_table(doc, REQ_HEADERS, rows, REQ_WIDTHS, font_size=8.5, first_col_bold=True)
    doc.add_page_break()


# --------------------------------------------------------------------------------------
# Historias de usuario
# --------------------------------------------------------------------------------------
USER_STORIES = [
    ("HU-01 — Anticipar la presión sobre la red",
     "Como **gestor de red**, quiero ver la demanda proyectada de urgencias de salud mental para "
     "las próximas semanas por comuna, para anticipar refuerzos de dotación antes de que la "
     "presión se produzca.",
     ["**Dado** que existe una publicación de datos vigente, **cuando** selecciono un horizonte de "
      "6 semanas, **entonces** el sistema muestra el valor proyectado y su intervalo para cada semana.",
      "**Dado** que una comuna no cumple el mínimo de observaciones, **cuando** consulto su "
      "proyección, **entonces** el sistema la marca como de baja confianza y explica el motivo."],
     "RF-DEM-002, RF-DEM-003, RF-DEM-008, RF-DEM-009, RF-VIS-003"),
    ("HU-02 — Confiar en el pronóstico",
     "Como **analista de servicio**, quiero conocer el desempeño del modelo frente a un baseline, "
     "para decidir si utilizo su resultado en una propuesta formal.",
     ["**Dado** que consulto la ficha metodológica, **cuando** reviso el desempeño, **entonces** veo "
      "las métricas de backtesting por horizonte y su comparación con el baseline estacional.",
      "**Dado** que el modelo no supera al baseline, **cuando** se ejecuta la promoción, "
      "**entonces** el sistema no lo publica y registra la decisión."],
     "RF-DEM-004, RF-DEM-005, RF-DEM-006, RF-DEM-007"),
    ("HU-03 — Ver dónde la red no llega",
     "Como **planificador territorial**, quiero ver qué zonas quedan fuera de un tiempo de viaje "
     "razonable a un dispositivo de urgencia, para fundamentar dónde reforzar la oferta.",
     ["**Dado** que selecciono el umbral de 30 minutos en transporte público, **cuando** aplico el "
      "filtro, **entonces** el mapa muestra la cobertura resultante y el porcentaje de población "
      "cubierta por comuna.",
      "**Dado** que existen establecimientos sin coordenadas, **cuando** consulto la vista espacial, "
      "**entonces** el sistema informa cuántos fueron excluidos y cuáles."],
     "RF-ACC-002, RF-ACC-003, RF-ACC-004, RF-ACC-006, RF-CAL-007"),
    ("HU-04 — Priorizar comunas críticas",
     "Como **referente comunal de salud mental**, quiero identificar las comunas que combinan alta "
     "demanda proyectada, baja accesibilidad y mayor vulnerabilidad, para priorizar la gestión con "
     "evidencia.",
     ["**Dado** que abro la vista de brechas, **cuando** reviso el ranking, **entonces** veo las "
      "comunas ordenadas por criterio de brecha y puedo consultar cómo se calcula."],
     "RF-ACC-005, RF-VIS-005, RF-VIS-010"),
    ("HU-05 — Sustentar una solicitud de recursos",
     "Como **jefatura de urgencia**, quiero descargar los datos y una ficha de mi territorio, para "
     "adjuntarlos a una solicitud formal.",
     ["**Dado** que apliqué filtros a una vista, **cuando** descargo los datos, **entonces** el "
      "archivo contiene exactamente lo mostrado, más la fecha de corte y la versión de datos."],
     "RF-VIS-011, RF-VIS-013, RF-VIS-008"),
    ("HU-06 — Entender qué estoy mirando",
     "Como **usuario no técnico**, quiero ver la definición de cada indicador y sus limitaciones, "
     "para no interpretar mal el resultado.",
     ["**Dado** que consulto un indicador, **cuando** abro su ayuda, **entonces** obtengo su "
      "definición, su método de cálculo y sus limitaciones vigentes.",
      "**Dado** que un período presenta discontinuidad documentada, **cuando** lo visualizo, "
      "**entonces** el sistema advierte que no es comparable."],
     "RF-VIS-010, RF-VIS-008, RF-CAL-005, RF-TRZ-003"),
    ("HU-07 — Caracterizar la hospitalización",
     "Como **epidemiólogo de servicio**, quiero revisar la duración de estadía por diagnóstico, "
     "previsión y condición de egreso, para entender qué factores se asocian a estadías más "
     "prolongadas.",
     ["**Dado** que selecciono una segmentación, **cuando** el subgrupo es pequeño, **entonces** el "
      "sistema advierte sobre la limitación estadística del resultado.",
      "**Dado** que consulto cualquier resultado del módulo, **cuando** lo reviso, **entonces** "
      "aparece la advertencia sobre naturaleza ecológica y ausencia de causalidad."],
     "RF-HOS-003, RF-HOS-004, RF-HOS-006"),
    ("HU-08 — Mantener el dato actualizado",
     "Como **administrador de datos**, quiero ejecutar y auditar la actualización, para publicar "
     "solo datos que superen los controles de calidad.",
     ["**Dado** que ejecuto el pipeline, **cuando** un control de integridad falla, **entonces** el "
      "sistema detiene la publicación e informa el control infractor.",
      "**Dado** que la ejecución fue exitosa, **cuando** reviso la bitácora, **entonces** encuentro "
      "volumen procesado, controles aplicados y versión publicada."],
     "RF-CAL-001, RF-CAL-002, RF-CAL-006, RF-ADM-003, RF-ADM-004, RF-ADM-005"),
    ("HU-09 — Reproducir un resultado",
     "Como **docente evaluadora**, quiero reproducir una cifra publicada a partir del repositorio, "
     "para verificar el rigor del trabajo.",
     ["**Dado** que dispongo del repositorio y de las fuentes, **cuando** ejecuto el pipeline y el "
      "informe, **entonces** obtengo las mismas cifras publicadas."],
     "RF-TRZ-001, RF-TRZ-002, RF-ETL-009"),
    ("HU-10 — Integrar el dato en otros análisis",
     "Como **analista externo**, quiero consumir los indicadores mediante una API documentada, para "
     "integrarlos en mis propios tableros.",
     ["**Dado** que consulto la especificación, **cuando** invoco un endpoint con parámetros "
      "válidos, **entonces** recibo la respuesta documentada con su metadato de versión."],
     "RF-API-001, RF-API-002, RF-API-005"),
]


def build_user_stories(doc) -> None:
    h1(doc, "5. Historias de usuario y criterios de aceptación")
    p(
        doc,
        "Las historias expresan los requisitos desde la perspectiva del usuario. Sus criterios de "
        "aceptación se redactan en formato **Dado / Cuando / Entonces**, de modo que puedan "
        "traducirse directamente en casos de prueba.",
    )
    for title, story, criteria, traceability in USER_STORIES:
        h3(doc, title)
        p(doc, story, italic=False)
        bullets(doc, criteria, size=10)
        p(doc, f"**Trazabilidad:** {traceability}", size=9.5)
    doc.add_page_break()


# --------------------------------------------------------------------------------------
# Casos de uso
# --------------------------------------------------------------------------------------
USE_CASES = [
    ("CU-01 — Consultar la proyección de demanda de una comuna", [
        ("Actor principal", "AC-01 Gestor de red / AC-02 Referente comunal de salud mental"),
        ("Objetivo", "Obtener la demanda proyectada de urgencias de salud mental de una comuna para un horizonte de 4 a 8 semanas."),
        ("Precondiciones", "Existe una publicación vigente de datos y un modelo promovido a producción."),
        ("Postcondiciones", "El usuario dispone de la proyección, su intervalo y su ficha metodológica."),
        ("Flujo principal",
         "1. El usuario accede al panel general.\n"
         "2. Selecciona la comuna y el horizonte de proyección.\n"
         "3. El sistema solicita los datos a la API.\n"
         "4. El sistema grafica la serie observada y la proyección con su banda de intervalo.\n"
         "5. El sistema muestra fecha de corte, versión de datos y advertencias aplicables."),
        ("Flujo alternativo A",
         "3a. La serie de la comuna no alcanza el mínimo de observaciones: el sistema muestra la serie observada, omite la proyección y explica la causa."),
        ("Flujo alternativo B",
         "3b. La API no responde: el sistema informa la indisponibilidad y ofrece reintentar, sin mostrar datos parciales sin etiqueta."),
        ("Requisitos asociados", "RF-VIS-001, RF-VIS-002, RF-VIS-003, RF-VIS-008, RF-DEM-002, RF-DEM-003, RF-DEM-008, RF-API-002"),
    ]),
    ("CU-02 — Identificar brechas de accesibilidad territorial", [
        ("Actor principal", "AC-01 Gestor de red / Planificador sanitario"),
        ("Objetivo", "Identificar comunas con alta demanda proyectada y baja cobertura de la red de urgencia."),
        ("Precondiciones", "Las isócronas y la cobertura poblacional están calculadas para la publicación vigente."),
        ("Postcondiciones", "El usuario obtiene el ranking de brechas y puede exportar su sustento."),
        ("Flujo principal",
         "1. El usuario abre la vista de brechas territoriales.\n"
         "2. Selecciona modo de transporte y umbral de tiempo de viaje.\n"
         "3. El sistema cruza cobertura, demanda proyectada y vulnerabilidad.\n"
         "4. El sistema muestra el ranking comunal y el mapa asociado.\n"
         "5. El usuario consulta la explicación del cálculo y descarga los datos."),
        ("Flujo alternativo A",
         "2a. El feed GTFS no está disponible para la franja solicitada: el sistema ofrece el cálculo por red vial e informa la sustitución."),
        ("Requisitos asociados", "RF-ACC-004, RF-ACC-005, RF-ACC-006, RF-ACC-008, RF-VIS-005, RF-VIS-011, RF-CAL-007"),
    ]),
    ("CU-03 — Ejecutar y auditar una actualización de datos", [
        ("Actor principal", "AC-05 Administrador de datos"),
        ("Objetivo", "Incorporar una nueva publicación oficial y dejarla disponible solo si supera los controles de calidad."),
        ("Precondiciones", "El administrador está autenticado y existe una nueva publicación de la fuente."),
        ("Postcondiciones", "Se publica una versión nueva de datos, o bien se mantiene la anterior y se registra el motivo del rechazo."),
        ("Flujo principal",
         "1. El administrador registra la nueva fuente.\n"
         "2. Ejecuta el pipeline de actualización.\n"
         "3. El sistema normaliza, filtra territorialmente y agrega.\n"
         "4. El sistema ejecuta los controles de calidad y genera el informe.\n"
         "5. Con veredicto apto, publica la versión y actualiza los metadatos.\n"
         "6. La bitácora registra el resultado de la ejecución."),
        ("Flujo alternativo A",
         "4a. Un control falla: el sistema detiene la publicación, conserva la versión anterior y detalla el control infractor."),
        ("Flujo alternativo B",
         "4b. La variación de volumen supera el umbral configurado: el sistema solicita confirmación explícita antes de publicar."),
        ("Requisitos asociados", "RF-ETL-001 a RF-ETL-010, RF-CAL-001 a RF-CAL-006, RF-CAL-010, RF-ADM-001, RF-ADM-003, RF-ADM-004, RF-ADM-005"),
    ]),
    ("CU-04 — Reproducir un resultado publicado", [
        ("Actor principal", "AC-06 Docente / Comisión evaluadora"),
        ("Objetivo", "Verificar que una cifra mostrada por la plataforma es reproducible desde el código y las fuentes declaradas."),
        ("Precondiciones", "El repositorio y las fuentes oficiales están disponibles."),
        ("Postcondiciones", "La cifra se reproduce, o bien se documenta la discrepancia observada."),
        ("Flujo principal",
         "1. El evaluador identifica la cifra y su versión de datos.\n"
         "2. Consulta el linaje asociado.\n"
         "3. Ejecuta el pipeline y el informe correspondiente.\n"
         "4. Compara el resultado obtenido con lo publicado."),
        ("Requisitos asociados", "RF-TRZ-001, RF-TRZ-002, RF-ETL-009, RF-ADM-005"),
    ]),
]


def build_use_cases(doc) -> None:
    h1(doc, "6. Casos de uso principales")
    p(
        doc,
        "Se documentan los cuatro casos de uso de mayor valor y riesgo. El resto de la funcionalidad "
        "queda cubierta por el catálogo de la sección 4 y por las historias de la sección 5.",
    )
    for title, lines in USE_CASES:
        add_definition_block(doc, title, lines)
    doc.add_page_break()


def build_business_rules(doc) -> None:
    h1(doc, "7. Reglas de negocio")
    p(
        doc,
        "Las reglas de negocio condicionan de forma transversal el comportamiento del sistema. "
        "Su incumplimiento invalida los resultados, con independencia de que la funcionalidad "
        "opere correctamente.",
    )
    add_table(
        doc,
        ["ID", "Regla", "Fundamento"],
        [
            ("RN-01", "El total de atenciones de urgencia y el total de salud mental provienen de identificadores de causa distintos y no deben sumarse entre sí, para evitar doble conteo.", "EDA-06, EDA-08"),
            ("RN-02", "El macro-agregador de salud mental debe ser igual a la suma exacta de sus subcausas; una diferencia distinta de cero invalida la carga.", "EDA-06"),
            ("RN-03", "El año 2020 no se utiliza como base de comparación para la evolución de la demanda de salud mental, por discontinuidad de captura documentada.", "EDA-08"),
            ("RN-04", "Los establecimientos sin coordenadas válidas se excluyen del cálculo espacial y se informan; no se les asigna una ubicación estimada.", "REG-01, EDA-08"),
            ("RN-05", "Ningún resultado puede vincular un egreso hospitalario con una atención de urgencia a nivel individual o de establecimiento.", "BIT-01, EDA-07"),
            ("RN-06", "Un volumen alto de atenciones refleja utilización y oferta instalada, no prevalencia poblacional; toda vista debe evitar esa interpretación.", "EDA-08"),
            ("RN-07", "Ningún modelo se publica sin haber sido comparado contra el baseline estacional ingenuo.", "BIT-01, DEF-01"),
            ("RN-08", "Toda cifra publicada debe ser reproducible desde código; no se admiten valores calculados manualmente.", "REG-01"),
            ("RN-09", "Las anomalías se documentan antes de corregirse; ninguna corrección se aplica sin registrar valor original, evidencia, regla y justificación.", "REG-01"),
            ("RN-10", "Cuando una causa o relación no pueda determinarse con los datos disponibles, el sistema debe declararlo explícitamente en lugar de inferirlo.", "REG-01"),
            ("RN-11", "Los informes, rutas y enlaces del sistema deben usar exclusivamente rutas relativas al proyecto, sin exponer rutas absolutas del entorno local.", "REG-01"),
        ],
        [Cm(1.7), Cm(12.5), Cm(3.5)],
        font_size=9.5,
        first_col_bold=True,
    )


def build_out_of_scope(doc) -> None:
    h1(doc, "8. Requisitos fuera de alcance")
    p(
        doc,
        "Los siguientes elementos fueron evaluados y **excluidos explícitamente** del alcance del "
        "MVP. Se registran para dejar constancia de la decisión, de su fundamento y de su eventual "
        "reevaluación futura.",
    )
    add_table(
        doc,
        ["ID", "Elemento excluido", "Fundamento de la exclusión", "Reevaluable en"],
        [
            ("FA-01", "Gestión de cupos o camas de hospitalización en tiempo real.", "Requiere captura, estandarización y coordinación entre múltiples instituciones con sistemas heterogéneos, probablemente con intervención estatal; el equipo no tiene acceso ni autoridad sobre esos datos.", "Trabajo futuro"),
            ("FA-02", "Coordinación operativa de derivaciones y contacto entre profesionales.", "Depende de FA-01 y de acuerdos institucionales fuera del plazo del proyecto.", "Trabajo futuro"),
            ("FA-03", "Tratamiento de datos personales, fichas clínicas o información identificable de pacientes.", "Restricción ética y legal asumida por el proyecto; las fuentes utilizadas son agregadas o disociadas.", "No previsto"),
            ("FA-04", "Apoyo a decisiones clínicas individuales o diagnóstico.", "El producto es de planificación poblacional; una interpretación clínica constituiría un uso indebido.", "No previsto"),
            ("FA-05", "Cobertura de regiones distintas a la Metropolitana.", "Acotado por alcance y tiempo disponible; la arquitectura debe permitir la extensión (ver requisitos de flexibilidad del ERNF).", "Trabajo futuro"),
            ("FA-06", "Inclusión del sector privado de salud.", "La fuente utilizada corresponde a la red pública; el sector privado carece de una fuente pública equivalente.", "Trabajo futuro"),
            ("FA-07", "Predicción de demanda a nivel de individuo o de riesgo personal.", "Incompatible con la granularidad agregada de los datos y con las restricciones éticas del proyecto.", "No previsto"),
            ("FA-08", "Integración en tiempo real con sistemas hospitalarios.", "No existe acceso ni convenio institucional dentro del plazo del proyecto.", "Trabajo futuro"),
        ],
        [Cm(1.5), Cm(4.6), Cm(8.6), Cm(3.0)],
        font_size=9,
        first_col_bold=True,
    )
    doc.add_page_break()


def build_traceability(doc) -> None:
    h1(doc, "9. Matriz de trazabilidad")

    h2(doc, "9.1 Objetivos específicos → Requisitos funcionales")
    add_table(
        doc,
        ["Objetivo específico (DEF-01)", "Requisitos que lo satisfacen"],
        [
            ("OE1. Levantar y validar requerimientos, usuarios, criterios de éxito y restricciones.", "Este documento (ERF) y el ERNF; RF-TRZ-003, RF-TRZ-005"),
            ("OE2. Auditar, limpiar, documentar e integrar las fuentes DEIS/MINSAL y los insumos territoriales.", "RF-ETL-001 a RF-ETL-014; RF-CAL-001 a RF-CAL-010; RF-TRZ-001, RF-TRZ-002"),
            ("OE3. Construir y evaluar el modelo de proyección de demanda a 4–8 semanas.", "RF-DEM-001 a RF-DEM-012"),
            ("OE4. Implementar la capa geoespacial de accesibilidad.", "RF-ACC-001 a RF-ACC-009; RF-ETL-006, RF-ETL-007, RF-ETL-013"),
            ("OE5. Analizar los determinantes de la duración de estadía F00–F99.", "RF-HOS-001 a RF-HOS-007; RF-ETL-012"),
            ("OE6. Integrar los resultados en una plataforma web, verificarla y validarla con usuarios.", "RF-VIS-001 a RF-VIS-013; RF-API-001 a RF-API-008; RF-ADM-001 a RF-ADM-007; RF-TRZ-005"),
        ],
        [Cm(8.7), Cm(9.0)],
        font_size=9.5,
    )

    h2(doc, "9.2 Origen de los requisitos y estado de avance")
    add_table(
        doc,
        ["Origen", "Requisitos derivados", "Estado al 10-09-2026"],
        [
            ("ENT-01 — Entrevista de dominio", "Necesidad de visión integrada y territorializada: RF-VIS-004, RF-VIS-005, RF-ACC-005. Necesidad de cupos en tiempo real: FA-01 (excluido con fundamento).", "Propuesto"),
            ("BIT-01 — Bitácora de alcance", "RG-01 a RG-06; RN-03, RN-05, RN-07; FA-01, FA-05, FA-06.", "Aprobado por el equipo"),
            ("EDA-03 / EDA-04 / EDA-05", "RF-ETL-002, RF-ETL-003, RF-ETL-004, RF-ETL-005, RF-CAL-001, RF-CAL-002, RF-CAL-004.", "**Implementado** para urgencias 2020–2026"),
            ("EDA-01 / EDA-02 / EDA-07", "RF-ETL-006, RF-ACC-001, RF-CAL-007.", "**Implementado** (cartografía y maestro); accesibilidad pendiente"),
            ("EDA-06 / EDA-08", "RF-ETL-010, RF-ETL-011, RF-CAL-003, RF-CAL-005; RN-01, RN-02, RN-03, RN-06.", "**Implementado** (catálogo y perfilado descriptivo)"),
            ("REG-01 — Reglas de rigor", "RN-08 a RN-11; RF-CAL-009, RF-TRZ-002, RF-TRZ-004.", "Vigente para todo el desarrollo"),
            ("DEF-01 — Definición del proyecto", "RF-DEM-004, RF-DEM-005, RF-DEM-007, RF-HOS-003, módulos RF-VIS y RF-API.", "Propuesto"),
        ],
        [Cm(4.2), Cm(8.5), Cm(5.0)],
        font_size=9,
    )


def build_increments(doc) -> None:
    h1(doc, "10. Plan de incorporación por incrementos")
    add_table(
        doc,
        ["Incremento", "Hito asociado", "Requisitos comprometidos", "Resultado esperado"],
        [
            ("I1 — Base de datos confiable", "Fase 2.1 (15-10-2026)", "RF-ETL-001 a RF-ETL-012; RF-CAL-001 a RF-CAL-009; RF-TRZ-001, RF-TRZ-002", "Pipeline reproducible con urgencias y egresos normalizados y auditados."),
            ("I2 — Analítica central", "Fase 2.1 → Fase 2.3", "RF-DEM-001 a RF-DEM-009; RF-ACC-001 a RF-ACC-006; RF-HOS-001 a RF-HOS-006", "Pronóstico evaluado contra baseline, isócronas y caracterización de estadía."),
            ("I3 — Plataforma integrada", "Fase 2.3 (26-11-2026)", "RF-API-001 a RF-API-008; RF-VIS-001 a RF-VIS-011; RF-ADM-001 a RF-ADM-005, RF-ADM-007; RF-TRZ-003, RF-TRZ-005", "MVP desplegado con las tres capas analíticas integradas y documentadas."),
            ("I4 — Mejoras opcionales", "Si el avance lo permite", "RF-ETL-013, RF-ETL-014; RF-DEM-010 a RF-DEM-012; RF-ACC-007, RF-ACC-009; RF-HOS-007; RF-VIS-012, RF-VIS-013; RF-ADM-006", "Extensiones de valor no comprometidas en el MVP."),
        ],
        [Cm(3.6), Cm(3.2), Cm(6.4), Cm(4.5)],
        font_size=9,
        first_col_bold=True,
    )


def build_acceptance(doc) -> None:
    h1(doc, "11. Criterios de aceptación global del sistema")
    p(doc, "El sistema se considera aceptado cuando se verifican, en conjunto, las siguientes condiciones:")
    numbered(
        doc,
        [
            "**Cobertura funcional:** el 100 % de los requisitos de prioridad *Debe* está implementado y verificado por su método declarado.",
            "**Reproducibilidad:** una ejecución independiente del pipeline sobre las fuentes declaradas reproduce las cifras publicadas.",
            "**Calidad del dato:** la publicación vigente cuenta con informe de calidad con veredicto apto y sin controles críticos en falla.",
            "**Utilidad del modelo:** el modelo publicado supera al baseline estacional ingenuo según el criterio documentado; en caso contrario, el sistema lo declara explícitamente y se abstiene de publicarlo.",
            "**Trazabilidad:** cada resultado visible es atribuible a una versión de datos y de modelo identificable.",
            "**Validación con usuarios:** se ejecuta al menos una sesión de validación con el cliente piloto o la experta de dominio, con hallazgos registrados y plan de ajuste.",
            "**Documentación:** existe documentación de despliegue, API, manual de usuario y declaración de supuestos y limitaciones.",
            "**Conformidad no funcional:** se cumplen los requisitos no funcionales de prioridad *Debe* definidos en el ERNF.",
        ],
    )


def build_change_control(doc) -> None:
    h1(doc, "12. Control de cambios de requisitos")
    numbered(
        doc,
        [
            "Toda modificación de un requisito aprobado se solicita mediante una incidencia en GitHub etiquetada `requisito`, indicando identificador afectado, cambio propuesto, motivo e impacto en alcance, plazo y trazabilidad.",
            "La Jefa de Proyecto evalúa el impacto; los cambios que afecten el alcance comprometido requieren validación de la docente.",
            "Los cambios aprobados se incorporan en una nueva versión de este documento, con registro en el historial de versiones y actualización de la matriz de trazabilidad.",
            "Los requisitos retirados conservan su identificador con estado `Retirado` y la referencia a la decisión que los retiró.",
            "Los cambios de alcance relevantes se registran además en la bitácora del proyecto (docs/03-bitacoras/).",
        ],
    )
    note(
        doc,
        "**Documento complementario:** Especificación de Requisitos No Funcionales (ERNF-SAAD-2026), "
        "disponible en `docs/06-requisitos/Especificacion_Requisitos_No_Funcionales.docx`.",
    )


def main() -> None:
    doc = new_document("ERF-SAAD-2026 · Especificación de Requisitos Funcionales · v1.0")
    build_cover(doc)
    build_control(doc)
    add_toc(doc)
    build_introduction(doc)
    build_overview(doc)
    build_conventions(doc)
    build_requirements(doc)
    build_user_stories(doc)
    build_use_cases(doc)
    build_business_rules(doc)
    build_out_of_scope(doc)
    build_traceability(doc)
    build_increments(doc)
    build_acceptance(doc)
    build_change_control(doc)
    path = save(doc, OUTPUT_PATH)
    print(f"Documento generado: {path.as_posix()}")


if __name__ == "__main__":
    main()
