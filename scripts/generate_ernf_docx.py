"""Genera la Especificación de Requisitos No Funcionales (ERNF) del proyecto SAAD en .docx.

Marco de referencia: ISO/IEC 25010 (calidad de producto), ISO/IEC 25012 (calidad de datos)
e ISO/IEC/IEEE 29148:2018 (métodos de verificación).
Ejecución: python scripts/generate_ernf_docx.py
Salida:    docs/06-requisitos/Especificacion_Requisitos_No_Funcionales.docx
"""

from __future__ import annotations

from pathlib import Path

from docx.shared import Cm

from docx_builder import (
    add_cover,
    add_table,
    add_toc,
    bullets,
    h1,
    h2,
    new_document,
    note,
    numbered,
    p,
    save,
)

OUTPUT_PATH = Path("docs/06-requisitos/Especificacion_Requisitos_No_Funcionales.docx")

RNF_WIDTHS = [Cm(2.0), Cm(5.6), Cm(8.2), Cm(1.9)]
RNF_HEADERS = ["ID", "Requisito", "Métrica y valor objetivo", "Prio. / Ver."]


def build_cover(doc) -> None:
    add_cover(
        doc,
        title="Especificación de Requisitos No Funcionales",
        subtitle=(
            "SAAD — Sistema de Análisis y Anticipación de Demanda de Urgencia "
            "en Salud Mental para la Región Metropolitana"
        ),
        project="Capstone APT · Ingeniería en Informática · Duoc UC Antonio Varas",
        meta_rows=[
            ("Código del documento", "ERNF-SAAD-2026"),
            ("Versión", "1.0"),
            ("Fecha de emisión", "10 de septiembre de 2026"),
            ("Estado", "Borrador para validación con la docente y el cliente piloto"),
            ("Clasificación", "Uso académico e interno del equipo de proyecto"),
            (
                "Marco de referencia",
                "ISO/IEC 25010 (calidad de producto); ISO/IEC 25012 (calidad de datos); "
                "ISO/IEC/IEEE 29148:2018",
            ),
            ("Elaborado por", "Equipo SAAD — Camila A., Felipe R., Cristopher R., Catalina"),
            ("Documento complementario", "Especificación de Requisitos Funcionales (ERF-SAAD-2026)"),
            ("Repositorio", "docs/06-requisitos/"),
        ],
    )


def build_control(doc) -> None:
    h1(doc, "Control del documento")

    h2(doc, "Historial de versiones")
    add_table(
        doc,
        ["Versión", "Fecha", "Autor(es)", "Descripción del cambio"],
        [
            ("0.1", "27-08-2026", "Equipo SAAD",
             "Listado preliminar de atributos de calidad derivados de la definición de alcance."),
            ("1.0", "10-09-2026", "Equipo SAAD",
             "Primera versión formal completa: catálogo de requisitos no funcionales cuantificados "
             "sobre ISO/IEC 25010, calidad de datos ISO/IEC 25012, restricciones, matriz de "
             "verificación y trazabilidad."),
        ],
        [Cm(2.0), Cm(2.6), Cm(3.6), Cm(9.5)],
        font_size=9.5,
    )

    h2(doc, "Revisión y aprobación")
    add_table(
        doc,
        ["Rol", "Nombre", "Estado", "Fecha"],
        [
            ("Jefa de Proyecto", "Camila A.", "Emitido para revisión", "10-09-2026"),
            ("Docente guía", "Por confirmar", "Pendiente de validación", "—"),
            ("Cliente piloto", "Por confirmar", "Pendiente de validación", "—"),
        ],
        [Cm(4.5), Cm(5.0), Cm(5.2), Cm(3.0)],
        font_size=9.5,
    )

    note(
        doc,
        "**Nota sobre los valores objetivo.** Las metas cuantitativas de este documento se fijan "
        "sobre el entorno de referencia declarado en la sección 3 y corresponden a un producto "
        "mínimo viable de contexto académico. Cualquier medición realizada en un entorno distinto "
        "debe declarar esa diferencia antes de comparar resultados.",
    )
    doc.add_page_break()


def build_introduction(doc) -> None:
    h1(doc, "1. Introducción")

    h2(doc, "1.1 Propósito")
    p(
        doc,
        "Este documento especifica **cómo debe comportarse** el sistema SAAD: los atributos de "
        "calidad, restricciones y condiciones de operación que debe satisfacer, expresados de "
        "forma medible y verificable. Complementa a la **Especificación de Requisitos Funcionales "
        "(ERF-SAAD-2026)**, que define las capacidades del sistema.",
    )

    h2(doc, "1.2 Alcance")
    p(
        doc,
        "Los requisitos aquí definidos aplican a la totalidad del sistema: pipeline de datos, "
        "modelos analíticos, API, aplicación web, documentación y proceso de desarrollo. Cuando un "
        "requisito aplica solo a un componente, se indica explícitamente en su enunciado.",
    )

    h2(doc, "1.3 Audiencia")
    bullets(
        doc,
        [
            "**Equipo de desarrollo:** define los criterios técnicos que condicionan el diseño y la implementación.",
            "**Docente guía y comisión evaluadora:** entrega criterios objetivos de evaluación del producto.",
            "**Cliente piloto y usuarios validadores:** explicita el nivel de servicio, la accesibilidad y las salvaguardas de uso responsable comprometidas.",
        ],
    )

    h2(doc, "1.4 Referencias normativas y técnicas")
    add_table(
        doc,
        ["Código", "Referencia"],
        [
            ("NOR-01", "ISO/IEC 25010 — Modelo de calidad del producto software y de sistemas."),
            ("NOR-02", "ISO/IEC 25012 — Modelo de calidad de datos."),
            ("NOR-03", "ISO/IEC/IEEE 29148:2018 — Ingeniería de requisitos y métodos de verificación."),
            ("NOR-04", "W3C WCAG 2.1, nivel AA — Pautas de accesibilidad para el contenido web."),
            ("NOR-05", "OWASP Top 10 — Riesgos de seguridad más críticos en aplicaciones web."),
            ("NOR-06", "OpenAPI Specification 3.x — Descripción de interfaces REST."),
            ("NOR-07", "RFC 7946 (GeoJSON) y estándar OGC GeoPackage — Intercambio geoespacial."),
            ("NOR-08", "PEP 8 y PEP 484 — Estilo y anotaciones de tipo en Python."),
            ("NOR-09", "Semantic Versioning 2.0.0 y Conventional Commits — Versionado y trazabilidad de cambios."),
            ("DEF-01", "Documento de Definición de Proyecto APT — Fase 1."),
            ("BIT-01", "Bitácora de reunión de definición y alcance, 23-08-2026."),
            ("REG-01", "Reglas de rigor analítico, portabilidad y seguridad del proyecto (.agent/rules.md)."),
            ("CTX-01", "Contexto técnico y decisiones vigentes (PROJECT_CONTEXT.md)."),
            ("ERF-01", "Especificación de Requisitos Funcionales SAAD (ERF-SAAD-2026)."),
        ],
        [Cm(2.4), Cm(15.3)],
        font_size=9.5,
        first_col_bold=True,
    )
    doc.add_page_break()


def build_conventions(doc) -> None:
    h1(doc, "2. Marco de calidad y convenciones")

    h2(doc, "2.1 Modelo de calidad adoptado")
    p(
        doc,
        "El catálogo se organiza según las características de calidad de producto de **ISO/IEC "
        "25010**, complementadas con las características de **calidad de datos de ISO/IEC 25012** "
        "—determinantes en un sistema cuyo valor depende íntegramente de la fidelidad de sus "
        "fuentes—. La tabla siguiente relaciona cada característica con su prefijo de "
        "identificación en este documento.",
    )
    add_table(
        doc,
        ["Característica de calidad", "Prefijo", "Foco en el contexto de SAAD"],
        [
            ("Adecuación funcional", "RNF-AFU", "Corrección y exactitud de las cifras y de las proyecciones publicadas."),
            ("Eficiencia de desempeño", "RNF-EDE", "Tiempos de respuesta de la plataforma y costo de procesar volúmenes de decenas de millones de registros."),
            ("Compatibilidad", "RNF-CMP", "Interoperabilidad mediante formatos abiertos y ejecución en navegadores y entornos estándar."),
            ("Capacidad de interacción (usabilidad)", "RNF-USA", "Comprensión correcta de indicadores por profesionales no técnicos y accesibilidad."),
            ("Fiabilidad", "RNF-FIA", "Disponibilidad durante la validación, degradación controlada y recuperación."),
            ("Seguridad", "RNF-SEG", "Protección del acceso administrativo, del canal y de la información publicada."),
            ("Mantenibilidad", "RNF-MAN", "Modularidad, pruebas, estándares de código y capacidad del equipo de sostener el sistema."),
            ("Flexibilidad (portabilidad y escalabilidad)", "RNF-FLE", "Extensión a otros períodos o territorios e instalación reproducible."),
            ("Seguridad de uso y uso responsable", "RNF-USO", "Salvaguardas frente a interpretación indebida, reidentificación y estigmatización territorial."),
            ("Calidad de datos (ISO/IEC 25012)", "RNF-DAT", "Exactitud, completitud, consistencia, actualidad y trazabilidad del dato publicado."),
        ],
        [Cm(5.6), Cm(2.3), Cm(9.8)],
        font_size=9.5,
    )

    h2(doc, "2.2 Identificación, prioridad y verificación")
    p(
        doc,
        "Cada requisito se identifica como **RNF-<PREFIJO>-<NNN>**. La prioridad utiliza la misma "
        "escala MoSCoW del ERF (**Debe**, **Debería**, **Podría**). El método de verificación sigue "
        "ISO/IEC/IEEE 29148:2018: **I** inspección, **A** análisis, **D** demostración, **P** "
        "prueba. Ambos valores se indican en la columna «Prio. / Ver.».",
    )

    h2(doc, "2.3 Criterios de redacción")
    bullets(
        doc,
        [
            "Todo requisito declara una **métrica**, una **unidad** y un **valor objetivo** verificable.",
            "Los percentiles se expresan como P95 salvo indicación distinta y se miden sobre al menos 30 observaciones.",
            "Cuando la meta depende del entorno, se refiere explícitamente al entorno de referencia de la sección 3.",
            "Los requisitos que no puedan medirse con los recursos del proyecto se declaran como *Podría* o se documentan como limitación, en lugar de enunciarse sin verificación.",
        ],
    )
    doc.add_page_break()


def build_environment(doc) -> None:
    h1(doc, "3. Entorno de referencia para las mediciones")
    p(
        doc,
        "Los valores objetivo de las secciones 4 y 5 se verifican sobre las siguientes condiciones. "
        "Toda medición debe registrar el entorno efectivo utilizado; si difiere del de referencia, "
        "la diferencia se documenta junto al resultado.",
    )
    add_table(
        doc,
        ["Dimensión", "Condición de referencia"],
        [
            ("Equipo de procesamiento",
             "Estación de trabajo del equipo con 4 núcleos físicos, 16 GB de memoria RAM y almacenamiento SSD "
             "(valores a confirmar y registrar por el equipo antes de la primera medición formal)."),
            ("Sistema operativo y entorno",
             "Windows 11 con WSL2 y Docker Compose; alternativamente, distribución Linux equivalente."),
            ("Base de datos", "PostgreSQL con extensión PostGIS, ejecutándose en contenedor local."),
            ("Cliente web de referencia",
             "Navegador Chrome en su última versión estable, resolución 1366×768 o superior."),
            ("Red de referencia",
             "Conexión de banda ancha de al menos 20 Mbps de bajada y latencia inferior a 50 ms hacia el servidor."),
            ("Conjunto de datos de referencia",
             "Atenciones de urgencia de la RM 2020–2026 (13.897.800 registros normalizados), maestro de "
             "1.172 establecimientos de la RM y cartografía comunal de las 52 comunas del Censo 2024."),
            ("Carga de usuarios de referencia",
             "Hasta 20 sesiones concurrentes, coherente con un uso institucional acotado durante la validación."),
        ],
        [Cm(4.4), Cm(13.3)],
        font_size=9.5,
        first_col_bold=True,
    )
    doc.add_page_break()


# --------------------------------------------------------------------------------------
# Catálogo de requisitos no funcionales
# --------------------------------------------------------------------------------------
RNF_AFU = [
    ("RNF-AFU-001", "Corrección de las cifras publicadas",
     "Discrepancia entre la cifra publicada y su recálculo desde el código versionado: **0** en el 100 % de los indicadores auditados.",
     "Debe / P"),
    ("RNF-AFU-002", "Utilidad mínima del modelo de pronóstico",
     "Error escalado del modelo respecto del baseline estacional ingenuo (MASE) **< 1,0** en el horizonte comprometido de 4 a 8 semanas; en caso contrario, no se publica el modelo y se declara el resultado.",
     "Debe / A"),
    ("RNF-AFU-003", "Cobertura de la incertidumbre",
     "Cobertura empírica del intervalo de predicción en backtesting dentro de ±10 puntos porcentuales del nivel nominal declarado.",
     "Debería / A"),
    ("RNF-AFU-004", "Completitud funcional del MVP",
     "Requisitos funcionales de prioridad *Debe* implementados y verificados: **100 %** al cierre de la Fase 2.3.",
     "Debe / I"),
    ("RNF-AFU-005", "Exactitud de la agregación territorial",
     "Diferencia entre la suma de las unidades y el valor del agregado mostrado (comuna, servicio de salud, tipo de establecimiento): **0**.",
     "Debe / P"),
]

RNF_EDE = [
    ("RNF-EDE-001", "Tiempo de carga del panel principal",
     "Tiempo hasta contenido útil del panel inicial: **≤ 3 s** (P95) en el entorno de referencia.",
     "Debe / P"),
    ("RNF-EDE-002", "Latencia de los servicios de consulta",
     "Tiempo de respuesta de los endpoints de series e indicadores: **≤ 800 ms** (P95) para consultas de hasta 5 años y 52 comunas.",
     "Debe / P"),
    ("RNF-EDE-003", "Respuesta a la aplicación de filtros",
     "Actualización de las vistas activas tras cambiar un filtro: **≤ 1,5 s** (P95).",
     "Debe / P"),
    ("RNF-EDE-004", "Desempeño del mapa interactivo",
     "Render inicial del mapa con las 52 comunas y la capa de establecimientos: **≤ 3 s** (P95); las operaciones de zoom y desplazamiento no bloquean la interfaz por más de **200 ms**.",
     "Debe / P"),
    ("RNF-EDE-005", "Costo de procesamiento del pipeline",
     "Ejecución completa del pipeline de urgencias de la RM (7 años): **≤ 90 min**, con uso máximo de memoria **≤ 8 GB** mediante procesamiento por particiones.",
     "Debe / P"),
    ("RNF-EDE-006", "Costo de entrenamiento y evaluación",
     "Entrenamiento del modelo más backtesting completo: **≤ 60 min** por ciclo.",
     "Debería / P"),
    ("RNF-EDE-007", "Precálculo geoespacial",
     "Cálculo de isócronas para todos los establecimientos de urgencia: **≤ 4 h**, ejecutado fuera de línea; las isócronas nunca se calculan durante una petición del usuario.",
     "Debe / A"),
    ("RNF-EDE-008", "Exportación de datos",
     "Generación y descarga de un archivo tabular de hasta 100.000 filas: **≤ 10 s**.",
     "Debería / P"),
    ("RNF-EDE-009", "Capacidad concurrente",
     "Con 20 sesiones concurrentes, la degradación del P95 de los endpoints de consulta no supera el **50 %** respecto de la medición con un usuario.",
     "Debería / P"),
    ("RNF-EDE-010", "Peso de la interfaz",
     "Peso total transferido en la carga inicial de la aplicación: **≤ 2,5 MB** comprimido.",
     "Debería / P"),
]

RNF_CMP = [
    ("RNF-CMP-001", "Navegadores soportados",
     "Funcionamiento sin defectos bloqueantes en las **dos últimas versiones estables** de Chrome, Edge y Firefox, y en Safari 16 o superior.",
     "Debe / P"),
    ("RNF-CMP-002", "Formatos abiertos de intercambio",
     "Las exportaciones se entregan en **CSV (UTF-8)**, **Parquet**, **GeoJSON (RFC 7946)** y, opcionalmente, **GeoPackage**; ningún resultado se publica exclusivamente en un formato propietario.",
     "Debe / I"),
    ("RNF-CMP-003", "Interfaz de programación estándar",
     "La API se describe conforme a **OpenAPI 3.x**, con especificación válida según un validador estándar y consumible desde clientes de terceros.",
     "Debe / P"),
    ("RNF-CMP-004", "Sistemas de referencia de coordenadas",
     "Las capas se almacenan preservando el CRS de origen (`EPSG:4674`), se publican en `EPSG:4326` y los cálculos métricos se realizan en un CRS proyectado métrico apropiado; cada capa declara su CRS.",
     "Debe / I"),
    ("RNF-CMP-005", "Compatibilidad de las descargas con herramientas ofimáticas",
     "Los archivos tabulares se abren correctamente en hojas de cálculo de uso común, sin pérdida de acentos ni desalineación de columnas.",
     "Debería / D"),
    ("RNF-CMP-006", "Coexistencia en el entorno de despliegue",
     "El sistema se ejecuta en contenedores con puertos y credenciales configurables, sin exigir versiones específicas del sistema operativo anfitrión.",
     "Debe / D"),
]

RNF_USA = [
    ("RNF-USA-001", "Eficacia en la tarea principal",
     "Al menos el **80 %** de los participantes de la prueba de usabilidad (mínimo 5 usuarios del perfil objetivo) completa sin asistencia la tarea «obtener la proyección de demanda de una comuna» en **≤ 3 min** en su primera sesión.",
     "Debe / D"),
    ("RNF-USA-002", "Satisfacción percibida",
     "Puntaje **SUS ≥ 70** promedio en la sesión de validación con usuarios.",
     "Debería / D"),
    ("RNF-USA-003", "Comprensión correcta de los indicadores",
     "Al menos el **80 %** de los participantes interpreta correctamente el significado de la banda de intervalo de predicción tras consultar la ayuda contextual.",
     "Debe / D"),
    ("RNF-USA-004", "Accesibilidad WCAG 2.1 AA",
     "Contraste mínimo **4,5:1** en texto normal y **3:1** en texto grande y elementos gráficos; navegación completa por teclado con foco visible; etiquetas asociadas en todos los controles; **0** incidencias críticas en auditoría automatizada de accesibilidad.",
     "Debe / P"),
    ("RNF-USA-005", "Independencia del color",
     "Ninguna información se transmite únicamente mediante color; las paletas de series y mapas son distinguibles bajo simulación de deuteranopía y protanopía.",
     "Debe / I"),
    ("RNF-USA-006", "Convenciones locales",
     "Interfaz íntegramente en español de Chile, con formato de fecha `dd-mm-aaaa`, separador de miles y decimal según convención local, y toda semana epidemiológica acompañada de su rango de fechas.",
     "Debe / I"),
    ("RNF-USA-007", "Diseño adaptable",
     "La aplicación es plenamente funcional desde **1366 × 768 px** y utilizable en tablet a partir de **768 px** de ancho, sin desbordes horizontales ni pérdida de controles.",
     "Debe / D"),
    ("RNF-USA-008", "Estados de la interfaz",
     "El **100 %** de las vistas que consumen datos implementa estado de carga, estado vacío y estado de error con mensaje accionable en lenguaje no técnico.",
     "Debe / D"),
    ("RNF-USA-009", "Protección frente al error del usuario",
     "Los filtros que producen un resultado vacío o no comparable informan la causa y ofrecen una acción de recuperación, en lugar de mostrar una vista en blanco.",
     "Debe / D"),
    ("RNF-USA-010", "Ayuda contextual accesible",
     "La definición y el método de cálculo de cualquier indicador visible se alcanzan en **un clic** desde la vista en que aparece.",
     "Debe / D"),
]

RNF_FIA = [
    ("RNF-FIA-001", "Disponibilidad durante la validación",
     "Disponibilidad del servicio **≥ 95 %** en horario hábil (lunes a viernes, 08:00–20:00) durante el período de validación con usuarios, medida mediante verificación automática cada 5 minutos.",
     "Debe / P"),
    ("RNF-FIA-002", "Degradación controlada",
     "Ante indisponibilidad del servicio de pronóstico, la plataforma continúa mostrando los datos observados con aviso explícito; **no** se presenta una pantalla de error total.",
     "Debe / D"),
    ("RNF-FIA-003", "Publicación atómica de datos",
     "Una actualización fallida deja el sistema en la versión anterior íntegra: **0** publicaciones parciales visibles para el usuario.",
     "Debe / P"),
    ("RNF-FIA-004", "Respaldo y recuperación",
     "Respaldo diario de la base de datos y de los artefactos de modelo, con **RPO ≤ 24 h** y **RTO ≤ 4 h**; la restauración se prueba con éxito al menos una vez antes de la entrega final.",
     "Debe / D"),
    ("RNF-FIA-005", "Tolerancia a fallos de las fuentes externas",
     "Las descargas de fuentes externas reintentan al menos **3 veces** con espera incremental y registran el fallo definitivo sin interrumpir el resto del pipeline.",
     "Debería / P"),
    ("RNF-FIA-006", "Madurez del producto entregado",
     "**0** defectos de severidad crítica o alta abiertos al momento de la entrega final; los defectos menores conocidos se documentan en el backlog.",
     "Debe / I"),
]

RNF_SEG = [
    ("RNF-SEG-001", "Ausencia de datos personales",
     "El sistema no almacena, procesa ni expone datos personales identificables: **0** atributos identificables en el modelo de datos, verificado por inspección del esquema y del diccionario de datos.",
     "Debe / I"),
    ("RNF-SEG-002", "Cifrado del canal",
     "Todo el tráfico externo se sirve sobre **TLS 1.2 o superior**, con redirección automática desde HTTP; **0** endpoints accesibles sin cifrado.",
     "Debe / P"),
    ("RNF-SEG-003", "Almacenamiento de credenciales",
     "Las contraseñas se almacenan con función de derivación de clave resistente (Argon2id o bcrypt con factor de costo suficiente); **0** credenciales en texto plano o con hash simple.",
     "Debe / I"),
    ("RNF-SEG-004", "Gestión de secretos",
     "Ningún secreto, credencial o cadena de conexión se versiona en el repositorio; la configuración sensible se inyecta por variables de entorno. Verificación mediante análisis de secretos sobre el historial: **0** hallazgos.",
     "Debe / P"),
    ("RNF-SEG-005", "Autorización verificada en el servidor",
     "Toda operación privilegiada valida el perfil del usuario en el servidor, con independencia de lo que muestre la interfaz: **0** operaciones administrativas ejecutables sin autorización efectiva.",
     "Debe / P"),
    ("RNF-SEG-006", "Protección frente a riesgos OWASP Top 10",
     "Uso de consultas parametrizadas, validación de entrada, cabeceras de seguridad y CORS restringido a orígenes autorizados; **0** vulnerabilidades críticas o altas detectadas en la revisión de seguridad previa a la entrega.",
     "Debe / P"),
    ("RNF-SEG-007", "Registro de auditoría",
     "Las acciones administrativas registran actor, acción, marca de tiempo y resultado, con retención **≥ 90 días**; los registros no contienen credenciales ni datos sensibles.",
     "Debe / I"),
    ("RNF-SEG-008", "Gestión de dependencias",
     "**0** dependencias con vulnerabilidades conocidas de severidad crítica o alta al momento de la entrega; las versiones se fijan explícitamente.",
     "Debe / P"),
    ("RNF-SEG-009", "Límite de uso de la API pública",
     "Los endpoints públicos aplican un límite de tasa configurable por origen, que impide el agotamiento del servicio por consumo masivo.",
     "Debería / P"),
    ("RNF-SEG-010", "Mensajes de error seguros",
     "Los errores devueltos al cliente no exponen trazas de ejecución, rutas absolutas del sistema ni detalles internos de la base de datos.",
     "Debe / P"),
]

RNF_MAN = [
    ("RNF-MAN-001", "Modularidad de la arquitectura",
     "El código respeta la separación por responsabilidad (`src/data/`, `src/models/`, `src/geo/`, `src/hospitalization/`, `src/api/`); **0** dependencias cruzadas no justificadas entre módulos analíticos.",
     "Debe / I"),
    ("RNF-MAN-002", "Estándar de código",
     "Conformidad con **PEP 8** y anotaciones de tipo en todas las funciones públicas; **0** errores del analizador estático en la rama principal.",
     "Debe / P"),
    ("RNF-MAN-003", "Cobertura de pruebas",
     "Cobertura de pruebas unitarias **≥ 80 %** en los módulos de transformación y cálculo analítico, y **≥ 60 %** en el total del proyecto.",
     "Debe / P"),
    ("RNF-MAN-004", "Prueba obligatoria por función analítica",
     "El **100 %** de las funciones de transformación o cálculo analítico cuenta con al menos una prueba unitaria en `tests/`.",
     "Debe / I"),
    ("RNF-MAN-005", "Integración continua",
     "Cada solicitud de incorporación ejecuta automáticamente el analizador estático y la suite de pruebas; no se integra código con la verificación en rojo.",
     "Debe / D"),
    ("RNF-MAN-006", "Revisión por pares",
     "El **100 %** de las incorporaciones a la rama principal es revisado por un integrante distinto del autor.",
     "Debe / I"),
    ("RNF-MAN-007", "Trazabilidad de cambios",
     "Uso de mensajes de commit convencionales y versionado semántico de las publicaciones; cada versión entregable declara su alcance de cambios.",
     "Debería / I"),
    ("RNF-MAN-008", "Complejidad acotada",
     "Complejidad ciclomática **≤ 10** por función; las excepciones se justifican explícitamente en el código.",
     "Debería / A"),
    ("RNF-MAN-009", "Documentación técnica viva",
     "Módulos y funciones públicas documentados; guía de instalación y despliegue actualizada en cada entrega: **0** instrucciones obsoletas verificadas en una instalación limpia.",
     "Debe / D"),
    ("RNF-MAN-010", "Capacidad de corrección",
     "Tiempo medio de corrección de un defecto de severidad menor **≤ 3 días hábiles** durante la fase de construcción.",
     "Podría / A"),
]

RNF_FLE = [
    ("RNF-FLE-001", "Parametrización sin cambios de código",
     "Territorio, período, horizonte de pronóstico y umbrales de isócrona se definen por configuración; **0** de estos parámetros embebidos en el código fuente.",
     "Debe / I"),
    ("RNF-FLE-002", "Extensibilidad territorial",
     "Incorporar una región adicional requiere únicamente cargar sus fuentes y ajustar la configuración, sin modificar la arquitectura ni el modelo de datos.",
     "Debería / A"),
    ("RNF-FLE-003", "Escalabilidad del volumen",
     "Un aumento del **30 %** en el volumen de datos (equivalente a dos años adicionales) no incrementa el tiempo del pipeline en más de **1,5 veces** ni obliga a rediseñar la arquitectura.",
     "Debería / A"),
    ("RNF-FLE-004", "Instalación reproducible",
     "En una máquina limpia con Docker, el entorno completo se levanta con un único comando en **≤ 20 min**, siguiendo la guía documentada.",
     "Debe / D"),
    ("RNF-FLE-005", "Reemplazabilidad del modelo",
     "El algoritmo de pronóstico puede sustituirse implementando la interfaz común definida, sin modificar la API ni la aplicación web.",
     "Debería / I"),
    ("RNF-FLE-006", "Portabilidad del entorno de ejecución",
     "El sistema se ejecuta indistintamente sobre Linux y sobre Windows con WSL2, con el mismo resultado funcional.",
     "Debe / D"),
    ("RNF-FLE-007", "Portabilidad de las rutas y referencias",
     "El **100 %** de las rutas en código, informes y documentación es relativo al proyecto; **0** rutas absolutas, nombres de usuario o unidades de disco del entorno local.",
     "Debe / I"),
]

RNF_USO = [
    ("RNF-USO-001", "Incertidumbre siempre visible",
     "**0** proyecciones presentadas como valor puntual sin su intervalo de predicción, tanto en la interfaz como en las exportaciones.",
     "Debe / I"),
    ("RNF-USO-002", "Advertencia de uso no clínico",
     "El **100 %** de las vistas analíticas y de los archivos exportados incluye la declaración de que el sistema apoya la planificación y no sustituye el juicio clínico ni orienta decisiones sobre personas.",
     "Debe / I"),
    ("RNF-USO-003", "Prevención de reidentificación",
     "Ninguna vista, exportación o endpoint publica celdas cuyo recuento sea inferior al umbral mínimo definido para desagregaciones cruzadas; los valores suprimidos se señalan como tales.",
     "Debe / P"),
    ("RNF-USO-004", "Prevención de interpretación indebida",
     "Toda comparación territorial advierte que el volumen de atenciones refleja utilización y oferta instalada, no prevalencia poblacional.",
     "Debe / I"),
    ("RNF-USO-005", "Lenguaje no estigmatizante",
     "Las glosas y etiquetas visibles usan denominaciones diagnósticas oficiales y evitan formulaciones que atribuyan condiciones a territorios o comunidades.",
     "Debe / I"),
    ("RNF-USO-006", "Distinción de la naturaleza de la afirmación",
     "Hechos observados, inferencias e hipótesis se identifican de forma diferenciada en informes y vistas: **0** inferencias presentadas con el mismo estatus que un valor observado.",
     "Debe / I"),
    ("RNF-USO-007", "Declaración de datos no determinables",
     "Cuando una relación o causa no puede establecerse con los datos disponibles, el sistema lo declara explícitamente en lugar de estimarla o insinuarla.",
     "Debe / I"),
]

RNF_DAT = [
    ("RNF-DAT-001", "Exactitud",
     "**0** discrepancias en los controles de consistencia definidos (total contra suma de tramos etarios; macro-agregador de salud mental contra suma de subcausas).",
     "Debe / P"),
    ("RNF-DAT-002", "Completitud",
     "**0 %** de nulos en identificadores clave; la cobertura geoespacial efectiva se reporta explícitamente en cada publicación, junto al detalle de los registros excluidos.",
     "Debe / P"),
    ("RNF-DAT-003", "Consistencia",
     "Esquema homogéneo en el **100 %** de los períodos de una misma fuente: mismos nombres de columna, tipos y dominios.",
     "Debe / P"),
    ("RNF-DAT-004", "Unicidad",
     "**0** duplicados exactos y **0** duplicados en la llave natural de cada tabla publicada.",
     "Debe / P"),
    ("RNF-DAT-005", "Credibilidad",
     "El **100 %** de las fuentes utilizadas es oficial y declara origen, fecha de descarga y versión; no se incorporan fuentes sin procedencia verificable.",
     "Debe / I"),
    ("RNF-DAT-006", "Actualidad",
     "La publicación vigente se actualiza dentro de los **30 días** posteriores a la publicación oficial de una nueva fuente; la fecha de corte es visible en toda vista.",
     "Debería / I"),
    ("RNF-DAT-007", "Trazabilidad",
     "El **100 %** de las tablas publicadas registra su linaje: fuente de origen, transformaciones aplicadas y ejecución que la generó.",
     "Debe / I"),
    ("RNF-DAT-008", "Comprensibilidad",
     "El **100 %** de las variables expuestas está definido en el diccionario de datos, con tipo, dominio, granularidad y unidad.",
     "Debe / I"),
    ("RNF-DAT-009", "Conformidad de la interpretación temporal",
     "Todo período con captura incompleta o discontinuidad documentada se marca como no comparable en los datos publicados y en las vistas que lo incluyan.",
     "Debe / I"),
]

CATEGORIES = [
    ("4.1 Adecuación funcional (RNF-AFU)",
     "Grado en que el sistema entrega resultados correctos y adecuados al propósito declarado. En "
     "SAAD, la corrección de la cifra y la utilidad demostrada del pronóstico son condiciones de "
     "aceptación, no aspiraciones.", RNF_AFU),
    ("4.2 Eficiencia de desempeño (RNF-EDE)",
     "Comportamiento temporal y uso de recursos, tanto en la interacción del usuario como en el "
     "procesamiento por lotes de volúmenes del orden de decenas de millones de registros.", RNF_EDE),
    ("4.3 Compatibilidad e interoperabilidad (RNF-CMP)",
     "Capacidad del sistema de intercambiar información mediante formatos abiertos y de operar en "
     "los entornos de sus usuarios sin requisitos especiales.", RNF_CMP),
    ("4.4 Capacidad de interacción y accesibilidad (RNF-USA)",
     "Condiciones para que profesionales no técnicos comprendan e interpreten correctamente los "
     "indicadores, incluidas las exigencias de accesibilidad.", RNF_USA),
    ("4.5 Fiabilidad (RNF-FIA)",
     "Capacidad de mantener el servicio durante el período de validación, degradarse de forma "
     "controlada y recuperarse ante fallos.", RNF_FIA),
    ("4.6 Seguridad (RNF-SEG)",
     "Protección del acceso administrativo, del canal de comunicación y de la información "
     "publicada. El sistema no trata datos personales, lo que reduce —pero no elimina— la "
     "superficie de riesgo.", RNF_SEG),
    ("4.7 Mantenibilidad (RNF-MAN)",
     "Condiciones que permiten al equipo modificar el sistema con bajo riesgo y sostener su "
     "calidad a lo largo del proyecto y después de él.", RNF_MAN),
    ("4.8 Flexibilidad, portabilidad y escalabilidad (RNF-FLE)",
     "Capacidad de adaptar el sistema a nuevos períodos, territorios o volúmenes y de instalarlo de "
     "forma reproducible en entornos distintos.", RNF_FLE),
    ("4.9 Seguridad de uso y uso responsable (RNF-USO)",
     "Salvaguardas específicas del dominio sanitario: evitar la interpretación indebida, la "
     "reidentificación y la estigmatización territorial. Son requisitos verificables, no "
     "recomendaciones editoriales.", RNF_USO),
]


def build_requirements(doc) -> None:
    h1(doc, "4. Requisitos no funcionales de producto")
    p(
        doc,
        "El catálogo se organiza según las características de ISO/IEC 25010. La columna "
        "«Prio. / Ver.» indica prioridad MoSCoW y método de verificación, conforme a la sección 2.2.",
    )
    for title, description, rows in CATEGORIES:
        h2(doc, title)
        p(doc, description)
        add_table(doc, RNF_HEADERS, rows, RNF_WIDTHS, font_size=8.5, first_col_bold=True)
    doc.add_page_break()


def build_data_quality(doc) -> None:
    h1(doc, "5. Requisitos de calidad de datos (ISO/IEC 25012)")
    p(
        doc,
        "El valor de SAAD depende íntegramente de la fidelidad de sus fuentes. Estos requisitos "
        "condicionan la publicación: un dato que no los satisface no se expone a los usuarios, "
        "aunque la funcionalidad opere correctamente.",
    )
    add_table(doc, RNF_HEADERS, RNF_DAT, RNF_WIDTHS, font_size=8.5, first_col_bold=True)
    note(
        doc,
        "**Relación con el ERF.** Estos requisitos se materializan en los controles funcionales del "
        "módulo `RF-CAL` y en las reglas de negocio RN-01 a RN-11 del ERF. Este documento fija el "
        "umbral de aceptación; el ERF define el mecanismo que lo hace cumplir.",
    )
    doc.add_page_break()


def build_constraints(doc) -> None:
    h1(doc, "6. Restricciones")

    h2(doc, "6.1 Restricciones tecnológicas")
    add_table(
        doc,
        ["ID", "Restricción", "Justificación"],
        [
            ("RT-01", "Backend en Python con FastAPI; procesamiento de datos con el ecosistema Pandas/PyArrow.", "Competencias del equipo y ecosistema analítico del proyecto."),
            ("RT-02", "Frontend en React con Tailwind CSS, Recharts para gráficos y React Leaflet para cartografía.", "Stack definido en la definición de proyecto (DEF-01)."),
            ("RT-03", "Persistencia en PostgreSQL con extensión PostGIS.", "Requerido por el análisis espacial de polígonos y rutas."),
            ("RT-04", "Empaquetado y despliegue mediante Docker Compose; entorno de desarrollo con WSL2.", "Reproducibilidad del entorno entre integrantes."),
            ("RT-05", "Control de versiones en Git/GitHub, con gestión ágil en GitHub Projects (Kanban).", "Metodología declarada del proyecto."),
            ("RT-06", "Pruebas automatizadas con pytest.", "Regla técnica del proyecto (REG-01)."),
            ("RT-07", "Las dependencias se registran de forma explícita y justificada en `requirements.txt`.", "Trazabilidad y reproducibilidad del entorno."),
        ],
        [Cm(1.7), Cm(9.5), Cm(6.5)],
        font_size=9,
        first_col_bold=True,
    )

    h2(doc, "6.2 Restricciones normativas y de licenciamiento")
    p(
        doc,
        "El proyecto no trata datos personales, por lo que su exposición normativa es acotada. No "
        "obstante, el marco aplicable se declara explícitamente y **debe validarse en cuanto a "
        "vigencia y aplicabilidad** antes de la entrega final, dado que parte de la normativa "
        "chilena de protección de datos se encuentra en proceso de transición.",
    )
    add_table(
        doc,
        ["ID", "Marco o condición", "Implicancia para SAAD", "Estado"],
        [
            ("RN-L1", "Normativa chilena de protección de datos personales.", "El sistema opera exclusivamente con datos agregados o disociados, sin tratamiento de datos personales; esta condición debe mantenerse en toda extensión futura.", "Aplicable; verificar vigencia específica"),
            ("RN-L2", "Normativa sobre derechos y deberes de los pacientes e información clínica.", "Se excluye del alcance el acceso a fichas clínicas y a información clínica identificable (FA-03 del ERF).", "Aplicable"),
            ("RN-L3", "Condiciones de uso de los datos abiertos publicados por el DEIS/MINSAL.", "Se cita la fuente, la fecha de descarga y la versión de cada base utilizada.", "Aplicable"),
            ("RN-L4", "Licencia de OpenStreetMap (ODbL).", "La atribución a OpenStreetMap y sus condiciones de licenciamiento deben respetarse en toda visualización y en las capas derivadas que se distribuyan.", "Aplicable"),
            ("RN-L5", "Condiciones de uso del feed GTFS del transporte público.", "Deben revisarse antes de redistribuir capas derivadas del feed.", "Por validar"),
            ("RN-L6", "Buenas prácticas de ciberseguridad para servicios públicos digitales.", "Se adoptan como referencia para los requisitos RNF-SEG, sin que el proyecto constituya un servicio institucional.", "Referencial"),
        ],
        [Cm(1.7), Cm(4.8), Cm(8.0), Cm(3.2)],
        font_size=9,
        first_col_bold=True,
    )

    h2(doc, "6.3 Restricciones organizacionales y académicas")
    add_table(
        doc,
        ["ID", "Restricción", "Efecto sobre los requisitos"],
        [
            ("RO-01", "Equipo de cuatro integrantes con dedicación parcial y calendario académico fijo (entrega final: 26-11-2026).", "Los requisitos de prioridad *Podría* se ejecutan solo si el avance lo permite."),
            ("RO-02", "Sin presupuesto de infraestructura asignado.", "Las metas de disponibilidad y capacidad se fijan en niveles compatibles con alojamiento gratuito o local."),
            ("RO-03", "Cliente piloto en proceso de confirmación.", "Las metas de usabilidad se verifican con el número de participantes efectivamente disponible, documentando la limitación."),
            ("RO-04", "Metodología ágil con Kanban y construcción incremental.", "Los requisitos se verifican por incremento y no exclusivamente al cierre."),
        ],
        [Cm(1.7), Cm(8.5), Cm(7.5)],
        font_size=9,
        first_col_bold=True,
    )

    h2(doc, "6.4 Restricciones derivadas de los datos")
    add_table(
        doc,
        ["ID", "Restricción", "Efecto sobre los requisitos"],
        [
            ("RD-01", "La serie de urgencias de salud mental es comparable a partir de 2021.", "Limita la evaluación de estacionalidad y condiciona RNF-AFU-002 y RNF-DAT-009."),
            ("RD-02", "Los datos de urgencia son recuentos agregados por causa y establecimiento, no registros individuales.", "Impide requisitos de precisión a nivel de persona; sustenta FA-07 del ERF."),
            ("RD-03", "La base de egresos no incluye el código de establecimiento individual.", "Obliga al cruce ecológico y condiciona RNF-USO-004."),
            ("RD-04", "Un porcentaje de los establecimientos del maestro carece de coordenadas válidas.", "Condiciona la cobertura declarable en RNF-DAT-002 y las advertencias de RF-CAL-007."),
            ("RD-05", "Los tramos etarios publicados son amplios.", "Impide comprometer desagregación infanto-adolescente precisa."),
        ],
        [Cm(1.7), Cm(8.5), Cm(7.5)],
        font_size=9,
        first_col_bold=True,
    )
    doc.add_page_break()


def build_verification(doc) -> None:
    h1(doc, "7. Estrategia de verificación y evidencias")
    p(
        doc,
        "Cada método de verificación se materializa en un tipo de evidencia específico, que se "
        "adjunta al plan y a los resultados de pruebas del proyecto.",
    )
    add_table(
        doc,
        ["Método", "Requisitos que lo utilizan", "Instrumento", "Evidencia esperada"],
        [
            ("Prueba (P)", "Desempeño, seguridad, calidad de datos, compatibilidad y cobertura de código.", "Suite automatizada de pruebas, pruebas de carga, auditoría automatizada de accesibilidad y análisis de dependencias.", "Reporte de ejecución con métricas, fecha y entorno."),
            ("Análisis (A)", "Utilidad del modelo, escalabilidad, complejidad y cobertura de intervalos.", "Backtesting, cálculo de métricas y análisis estático.", "Informe con el cálculo reproducible y su interpretación."),
            ("Demostración (D)", "Usabilidad, degradación controlada, instalación e integración continua.", "Sesión de validación con usuarios y ejecución guiada ante evaluador.", "Bitácora de sesión, resultados y grabación o registro fotográfico."),
            ("Inspección (I)", "Rigor documental, uso responsable, trazabilidad y estándares de código.", "Revisión de código, de esquema y de documentación por un integrante distinto del autor.", "Registro de revisión con hallazgos y su resolución."),
        ],
        [Cm(2.6), Cm(5.0), Cm(5.0), Cm(5.1)],
        font_size=9,
        first_col_bold=True,
    )
    p(
        doc,
        "**Criterio de aceptación no funcional:** el sistema se considera conforme cuando el "
        "**100 %** de los requisitos de prioridad *Debe* de este documento ha sido verificado por "
        "su método declarado y sus resultados están registrados en el plan de pruebas del proyecto. "
        "Los requisitos *Debería* y *Podría* no verificados se documentan como limitación explícita "
        "en el informe final, con su justificación.",
    )


def build_traceability(doc) -> None:
    h1(doc, "8. Trazabilidad")
    p(
        doc,
        "La tabla relaciona cada grupo de requisitos no funcionales con los requisitos funcionales "
        "y los objetivos del proyecto sobre los que actúa.",
    )
    add_table(
        doc,
        ["Grupo RNF", "Requisitos funcionales relacionados (ERF)", "Objetivo específico"],
        [
            ("RNF-AFU", "RF-DEM-004 a RF-DEM-007; RF-CAL-003; RF-TRZ-002", "OE2, OE3"),
            ("RNF-EDE", "RF-ETL-005, RF-ETL-009; RF-ACC-002, RF-ACC-003; RF-VIS-002 a RF-VIS-004; RF-API-002, RF-API-007", "OE2, OE4, OE6"),
            ("RNF-CMP", "RF-API-001, RF-API-003; RF-ACC-007; RF-VIS-011", "OE6"),
            ("RNF-USA", "RF-VIS-001 a RF-VIS-011", "OE6"),
            ("RNF-FIA", "RF-ADM-003 a RF-ADM-007; RF-CAL-006", "OE2, OE6"),
            ("RNF-SEG", "RF-ADM-001, RF-ADM-002; RF-API-006; RF-HOS-005", "OE6"),
            ("RNF-MAN", "RF-ETL-009; RF-TRZ-001, RF-TRZ-005", "OE2, OE6"),
            ("RNF-FLE", "RF-ETL-014; RF-DEM-010; RF-ACC-008", "OE2, OE3, OE4"),
            ("RNF-USO", "RF-DEM-003, RF-DEM-008; RF-HOS-005, RF-HOS-006; RF-VIS-008; RF-TRZ-003, RF-TRZ-004", "OE5, OE6"),
            ("RNF-DAT", "RF-CAL-001 a RF-CAL-010; RF-ETL-002 a RF-ETL-004; RF-TRZ-001", "OE2"),
        ],
        [Cm(2.6), Cm(11.5), Cm(3.6)],
        font_size=9,
        first_col_bold=True,
    )


def build_risks(doc) -> None:
    h1(doc, "9. Riesgos asociados a los atributos de calidad")
    add_table(
        doc,
        ["ID", "Riesgo", "Requisito afectado", "Mitigación comprometida"],
        [
            ("RQ-01", "El costo de cómputo geoespacial excede el tiempo disponible.", "RNF-EDE-007", "Precálculo fuera de línea, reducción de umbrales y simplificación documentada de geometrías."),
            ("RQ-02", "El feed GTFS no está disponible o es inconsistente.", "RNF-EDE-007, RNF-CMP-004", "Degradar el análisis a red vial, declarando explícitamente la sustitución (SD-03 del ERF)."),
            ("RQ-03", "La longitud de la serie impide superar al baseline.", "RNF-AFU-002", "No publicar el modelo, informar el resultado y mantener la vista descriptiva con datos observados."),
            ("RQ-04", "No se dispone de suficientes participantes para la prueba de usabilidad.", "RNF-USA-001, RNF-USA-002", "Reducir la muestra documentando la limitación y complementar con revisión heurística."),
            ("RQ-05", "Las metas de disponibilidad no son alcanzables sin infraestructura pagada.", "RNF-FIA-001", "Acordar con la docente una ventana de disponibilidad acotada al período de validación."),
            ("RQ-06", "Sobreinterpretación de los resultados por parte de un usuario.", "RNF-USO-001 a RNF-USO-007", "Advertencias obligatorias en vistas y exportaciones, más validación de comprensión en la sesión con usuarios."),
            ("RQ-07", "Cambio de esquema en una publicación oficial futura.", "RNF-DAT-003", "Validación de esquema bloqueante y auditoría de la nueva fuente antes de publicar."),
        ],
        [Cm(1.5), Cm(5.2), Cm(3.4), Cm(7.6)],
        font_size=9,
        first_col_bold=True,
    )


def build_change_control(doc) -> None:
    h1(doc, "10. Control de cambios")
    numbered(
        doc,
        [
            "Toda modificación de un requisito no funcional aprobado se solicita mediante una incidencia etiquetada `requisito`, indicando identificador, valor objetivo actual, valor propuesto y evidencia que justifica el cambio.",
            "Un valor objetivo solo se relaja cuando existe evidencia de medición que demuestre su inviabilidad en el entorno de referencia; la evidencia se adjunta a la solicitud.",
            "Los cambios aprobados se registran en el historial de versiones de este documento y se comunican a la docente en la instancia de avance más próxima.",
            "Los requisitos no verificados al cierre se documentan como limitación explícita en el informe final, junto con su causa y su impacto.",
        ],
    )
    note(
        doc,
        "**Documento complementario:** Especificación de Requisitos Funcionales (ERF-SAAD-2026), "
        "disponible en `docs/06-requisitos/Especificacion_Requisitos_Funcionales.docx`.",
    )


def main() -> None:
    doc = new_document("ERNF-SAAD-2026 · Especificación de Requisitos No Funcionales · v1.0")
    build_cover(doc)
    build_control(doc)
    add_toc(doc)
    build_introduction(doc)
    build_conventions(doc)
    build_environment(doc)
    build_requirements(doc)
    build_data_quality(doc)
    build_constraints(doc)
    build_verification(doc)
    build_traceability(doc)
    build_risks(doc)
    build_change_control(doc)
    path = save(doc, OUTPUT_PATH)
    print(f"Documento generado: {path.as_posix()}")


if __name__ == "__main__":
    main()
