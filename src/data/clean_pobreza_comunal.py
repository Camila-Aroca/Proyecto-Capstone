"""Construye la dimensión de vulnerabilidad socioeconómica comunal RM (Casen 2022).

Fuente: Ministerio de Desarrollo Social y Familia (MDS), Observatorio Social,
"Estimaciones de Tasa de Pobreza por Ingresos por Comuna. Aplicación de
Metodología de Estimación para Áreas Pequeñas (SAE)", Encuesta Casen 2022. El
cuadro publica, para las 346 comunas del país, el porcentaje de personas en
situación de pobreza por ingresos 2022 junto con su intervalo de confianza,
la presencia de la comuna en la muestra Casen y el tipo de estimación SAE
(directa+sintética Fay-Herriot cuando la comuna tiene presencia muestral
suficiente, o sintética pura en caso contrario).

`indicador_vulnerabilidad` es la tasa de pobreza por ingresos 2022 (proporción
0-1): a mayor valor, mayor vulnerabilidad socioeconómica comunal. Es un
indicador estático de un único período (Casen 2022, no una serie histórica) y
no debe replicarse artificialmente por semana/mes ni combinarse aquí con
demanda de Urgencias para construir un puntaje de riesgo compuesto.

Las 52 comunas RM verificadas en el cuadro fuente tienen presencia en la
muestra Casen (columna "Presencia muestra" = "Sí"); esa condición se valida
explícitamente antes de publicar y su ausencia detiene la publicación. Dentro
de esas 52, el tipo de estimación SAE observado no es uniforme: 47 comunas
usan la estimación combinada directa+sintética (Fay-Herriot) y 5 (Calera de
Tango, Curacaví, San Pedro, El Monte e Isla de Maipo) usan estimación
sintética pura por no cumplir los criterios de inclusión muestral, pese a
tener presencia en Casen. `tipo_estimacion_sae` se conserva sin filtrar para
que esa menor precisión relativa quede documentada, no oculta.
"""

from __future__ import annotations

import argparse
import logging
from pathlib import Path
import tempfile
from typing import Any

from openpyxl import load_workbook
import pyarrow as pa
import pyarrow.parquet as pq

from src.data.clean_censo_comunas import EXPECTED_RM_CUTS


logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
logger = logging.getLogger(__name__)

RAW_PATH = Path("data/raw/pobreza/mds_casen_2022_tasa_pobreza_ingresos_comunal.xlsx")
GLOSA_CATALOG_PATH = Path("data/processed/censo/dim_poblacion_comuna_censo2024.parquet")
OUTPUT_PATH = Path("data/processed/pobreza/dim_vulnerabilidad_comuna.parquet")

SHEET_NAME = "Estimaciones"
HEADER_ROW = 3
FIRST_DATA_ROW = 4
REGION_METROPOLITANA_CODIGO = 13
ANO_REFERENCIA = 2022
TASA_MIN = 0.0
TASA_MAX = 1.0
NOMBRE_INDICADOR = "Tasa de pobreza por ingresos (estimación comunal SAE, Casen 2022)"
DIRECCION_INDICADOR = "mayor_valor_mayor_vulnerabilidad"
FUENTE = (
    "Ministerio de Desarrollo Social y Familia (MDS), Observatorio Social; "
    "Estimaciones de Tasa de Pobreza por Ingresos por Comuna, Encuesta Casen "
    "2022, Metodologia de Estimacion para Areas Pequenas (SAE, Fay-Herriot)."
)
PRESENCIA_MUESTRA_ESPERADA = "Sí"

SCHEMA = pa.schema([
    ("comuna_codigo", pa.string()),
    ("comuna_glosa", pa.string()),
    ("ano_referencia", pa.int32()),
    ("indicador_vulnerabilidad", pa.float64()),
    ("nombre_indicador", pa.string()),
    ("fuente", pa.string()),
    ("direccion_indicador", pa.string()),
    ("intervalo_confianza_inferior", pa.float64()),
    ("intervalo_confianza_superior", pa.float64()),
    ("tipo_estimacion_sae", pa.string()),
])


def _load_comuna_glosas(catalog_path: Path) -> dict[str, str]:
    """Lee comuna_codigo->comuna_glosa desde el catálogo comunal ya validado."""
    if not catalog_path.exists():
        raise FileNotFoundError(f"No existe el catálogo de glosas comunales requerido: {catalog_path.as_posix()}")
    table = pq.read_table(catalog_path, columns=["comuna_codigo", "comuna_glosa"])
    glosas = dict(zip(table["comuna_codigo"].to_pylist(), table["comuna_glosa"].to_pylist()))
    missing = {str(cut) for cut in EXPECTED_RM_CUTS} - glosas.keys()
    if missing:
        raise ValueError(f"Catálogo de glosas comunales incompleto, faltan CUT: {sorted(missing)}")
    return glosas


def _parse_rm_rows(raw_path: Path) -> list[dict[str, Any]]:
    """Extrae filas RM (región 13) del cuadro comunal, deteniéndose en las notas al pie."""
    workbook = load_workbook(raw_path, read_only=True, data_only=True)
    try:
        sheet = workbook[SHEET_NAME]
        rows: list[dict[str, Any]] = []
        for row_idx in range(FIRST_DATA_ROW, sheet.max_row + 1):
            codigo = sheet.cell(row_idx, 1).value
            if not isinstance(codigo, (int, float)):
                break
            codigo = int(codigo)
            if not (13000 <= codigo < 14000):
                continue
            tasa = sheet.cell(row_idx, 6).value
            limite_inferior = sheet.cell(row_idx, 7).value
            limite_superior = sheet.cell(row_idx, 8).value
            presencia_muestra = sheet.cell(row_idx, 9).value
            tipo_estimacion = sheet.cell(row_idx, 10).value
            if tasa is None or limite_inferior is None or limite_superior is None:
                raise ValueError(f"Valor nulo en la fuente MDS/Casen para comuna {codigo} (fila {row_idx}).")
            rows.append({
                "comuna_codigo": str(codigo),
                "indicador_vulnerabilidad": float(tasa),
                "intervalo_confianza_inferior": float(limite_inferior),
                "intervalo_confianza_superior": float(limite_superior),
                "presencia_muestra": presencia_muestra,
                "tipo_estimacion_sae": tipo_estimacion,
            })
        return rows
    finally:
        workbook.close()


def _validate_rows(rows: list[dict[str, Any]]) -> None:
    expected_codes = {str(cut) for cut in EXPECTED_RM_CUTS}
    observed_codes = [row["comuna_codigo"] for row in rows]
    if len(observed_codes) != len(set(observed_codes)):
        raise ValueError("Comunas RM duplicadas en la fuente de pobreza comunal.")
    observed_set = set(observed_codes)
    if observed_set != expected_codes:
        missing = expected_codes - observed_set
        extra = observed_set - expected_codes
        raise ValueError(
            f"Comunas RM inconsistentes con el CUT oficial. Faltantes: {sorted(missing)}, "
            f"inesperadas: {sorted(extra)}."
        )
    for row in rows:
        tasa = row["indicador_vulnerabilidad"]
        if not (TASA_MIN <= tasa <= TASA_MAX):
            raise ValueError(
                f"Tasa de pobreza fuera de dominio [0,1] para comuna {row['comuna_codigo']}: {tasa}."
            )
        lower, upper = row["intervalo_confianza_inferior"], row["intervalo_confianza_superior"]
        if not (lower <= tasa <= upper):
            raise ValueError(
                f"Intervalo de confianza inconsistente con la tasa para comuna {row['comuna_codigo']}: "
                f"[{lower}, {tasa}, {upper}]."
            )
        if row["presencia_muestra"] != PRESENCIA_MUESTRA_ESPERADA:
            raise ValueError(
                f"Comuna RM {row['comuna_codigo']} sin presencia en la muestra Casen "
                f"({row['presencia_muestra']!r}); requiere revisión antes de publicar."
            )


def build_dim_vulnerabilidad_comuna(
    raw_path: Path = RAW_PATH,
    glosa_catalog_path: Path = GLOSA_CATALOG_PATH,
    output_path: Path = OUTPUT_PATH,
) -> dict[str, Any]:
    """Genera `dim_vulnerabilidad_comuna.parquet` de forma validada y atómica."""
    if not raw_path.exists():
        raise FileNotFoundError(f"No existe el RAW requerido: {raw_path.as_posix()}")

    glosas = _load_comuna_glosas(glosa_catalog_path)
    rows = _parse_rm_rows(raw_path)
    _validate_rows(rows)

    records = [
        {
            "comuna_codigo": row["comuna_codigo"],
            "comuna_glosa": glosas[row["comuna_codigo"]],
            "ano_referencia": ANO_REFERENCIA,
            "indicador_vulnerabilidad": row["indicador_vulnerabilidad"],
            "nombre_indicador": NOMBRE_INDICADOR,
            "fuente": FUENTE,
            "direccion_indicador": DIRECCION_INDICADOR,
            "intervalo_confianza_inferior": row["intervalo_confianza_inferior"],
            "intervalo_confianza_superior": row["intervalo_confianza_superior"],
            "tipo_estimacion_sae": row["tipo_estimacion_sae"],
        }
        for row in sorted(rows, key=lambda item: item["comuna_codigo"])
    ]

    output_path.parent.mkdir(parents=True, exist_ok=True)
    table = pa.Table.from_pylist(records, schema=SCHEMA)
    with tempfile.NamedTemporaryFile(
        delete=False, dir=output_path.parent, prefix=f"{output_path.name}.", suffix=".tmp"
    ) as temp:
        temp_path = Path(temp.name)
    try:
        pq.write_table(table, temp_path, compression="snappy")
        temp_path.replace(output_path)
    except Exception:
        temp_path.unlink(missing_ok=True)
        raise

    verify_table = pq.read_table(output_path)
    return {
        "archivo_generado": output_path.as_posix(),
        "filas": verify_table.num_rows,
        "ano_referencia": ANO_REFERENCIA,
        "comunas_unicas": len(set(verify_table["comuna_codigo"].to_pylist())),
    }


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Construye la dimensión de vulnerabilidad socioeconómica comunal RM (MDS/Casen 2022).",
    )
    parser.add_argument("--force", action="store_true", help="Aceptado para interoperar con el orquestador.")
    parser.parse_args()
    result = build_dim_vulnerabilidad_comuna()
    logger.info("Dimensión de vulnerabilidad comunal generada: %s", result)


if __name__ == "__main__":
    main()
