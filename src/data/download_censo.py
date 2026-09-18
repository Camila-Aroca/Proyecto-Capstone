"""Ingesta segura de la cartografía RAW del Censo 2024 (comunal y subcomunal).

El ZIP fuente es un contenedor de transporte único que agrupa capas nacionales
a distintos grados de detalle territorial (comunal, distrital, zonal,
entidades, manzanas) más su diccionario oficial de variables geográficas.
Este módulo descarga ese ZIP una sola vez a `.cache/downloads/censo/`, valida
cada miembro requerido de forma independiente y publica únicamente los
artefactos RAW que falten o estén inválidos, sin tocar los que ya sean
snapshots válidos.
"""

from __future__ import annotations

import argparse
import json
import logging
import os
from pathlib import Path
import shutil
import ssl
from tempfile import TemporaryDirectory
import urllib.request
import zipfile

from openpyxl import load_workbook
import pyarrow.parquet as pq
from shapely import wkb

from src.data.raw_provenance import file_sha256, record_raw_snapshot


logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
logger = logging.getLogger(__name__)

URL_CENSO = "https://storage.googleapis.com/bktdescargascenso2024/Cartografia/GEOPARQUET/Cartografia_censo2024_Pais.zip"
DEST_DIR = Path("data/raw/censo")
CACHE_DIR = Path(".cache/downloads/censo")

# Nombres reales observados dentro del ZIP (confirmados por inspección directa,
# no asumidos): una capa geoparquet por nivel territorial más el diccionario
# oficial de variables geográficas. El ZIP también trae Aldeas, Limite_Urbano,
# Localidades, Provincial y Regional, que no son necesarios para este alcance
# y por lo tanto no se publican como RAW.
TARGET_FILE_NAME = "Cartografia_censo2024_Pais_Comunal.parquet"  # compatibilidad histórica
LAYER_FILE_NAMES = {
    "Comunal": "Cartografia_censo2024_Pais_Comunal.parquet",
    "Distrital": "Cartografia_censo2024_Pais_Distrital.parquet",
    "Zonal": "Cartografia_censo2024_Pais_Zonal.parquet",
    "Entidades": "Cartografia_censo2024_Pais_Entidades.parquet",
    "Manzanas": "Cartografia_censo2024_Pais_Manzanas.parquet",
}
DICTIONARY_FILE_NAME = "Diccionario_variables_geograficas_CPV24.xlsx"
REQUIRED_DICTIONARY_SHEETS = {
    "Comunal_CPV24", "Distrital_CPV24", "Zonal_CPV24", "Entidades_CPV24", "Manzanas_CPV24",
}

RAW_PATHS = {name: DEST_DIR / fname for name, fname in LAYER_FILE_NAMES.items()}
RAW_DICTIONARY_PATH = DEST_DIR / DICTIONARY_FILE_NAME
FINAL_OUTPUT_PATH = RAW_PATHS["Comunal"]  # compatibilidad histórica

MANIFEST_PATH = Path("data/raw/provenance_manifest.json")
REQUIRED_COLUMNS = {"CUT", "COMUNA", "COD_REGION", "SHAPE"}

_SOURCE_IDS = {
    "Comunal": "censo_2024_cartografia_comunal",
    "Distrital": "censo_2024_cartografia_distrital",
    "Zonal": "censo_2024_cartografia_zonal",
    "Entidades": "censo_2024_cartografia_entidades",
    "Manzanas": "censo_2024_cartografia_manzanas",
    "Diccionario": "censo_2024_diccionario_variables_geograficas",
}


def _download_zip(url: str, cache_path: Path) -> None:
    """Descarga el transporte a cache; nunca a ``data/raw``."""
    cache_path.parent.mkdir(parents=True, exist_ok=True)
    partial_path = cache_path.with_suffix(f"{cache_path.suffix}.part")
    request = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    context = ssl.create_default_context()
    context.check_hostname = False
    context.verify_mode = ssl.CERT_NONE
    with urllib.request.urlopen(request, context=context) as response, partial_path.open("wb") as output:
        shutil.copyfileobj(response, output, length=1024 * 1024)
    if partial_path.stat().st_size == 0:
        raise ValueError("ZIP del Censo descargado vacío.")
    os.replace(partial_path, cache_path)


def _find_member(archive: zipfile.ZipFile, filename: str) -> zipfile.ZipInfo:
    matches = [member for member in archive.infolist()
               if not member.is_dir() and Path(member.filename).name == filename]
    if len(matches) != 1:
        raise ValueError(f"Se esperaba exactamente un miembro '{filename}' en el ZIP del Censo.")
    return matches[0]


def _validate_transport(zip_path: Path) -> zipfile.ZipFile:
    if not zip_path.is_file() or zip_path.stat().st_size == 0 or not zipfile.is_zipfile(zip_path):
        raise ValueError("Transporte ZIP del Censo inexistente, vacío o inválido.")
    archive = zipfile.ZipFile(zip_path)
    corrupt_member = archive.testzip()
    if corrupt_member:
        archive.close()
        raise zipfile.BadZipFile(f"Miembro corrupto en ZIP Censo: {corrupt_member}")
    return archive


def validate_censo_geoparquet(path: Path) -> None:
    """Valida legibilidad, esquema mínimo y una geometría de la capa dada."""
    if not path.is_file() or path.stat().st_size == 0:
        raise ValueError(f"GeoParquet Censo inexistente o vacío: {path.name}")
    parquet = pq.ParquetFile(path)
    schema = parquet.schema_arrow
    missing = REQUIRED_COLUMNS.difference(schema.names)
    if missing:
        raise ValueError(f"GeoParquet Censo '{path.name}' sin columnas requeridas: {sorted(missing)}")
    if parquet.metadata is None or parquet.metadata.num_rows <= 0:
        raise ValueError(f"GeoParquet Censo '{path.name}' sin filas.")
    geo_metadata = (schema.metadata or {}).get(b"geo")
    if not geo_metadata:
        raise ValueError(f"GeoParquet Censo '{path.name}' sin metadata geoespacial.")
    try:
        geo_columns = json.loads(geo_metadata.decode("utf-8")).get("columns", {})
    except (UnicodeDecodeError, json.JSONDecodeError) as error:
        raise ValueError(f"Metadata geoespacial inválida en GeoParquet Censo '{path.name}'.") from error
    if "SHAPE" not in geo_columns:
        raise ValueError(f"Metadata geoespacial sin columna SHAPE en '{path.name}'.")
    geometry_batch = next(parquet.iter_batches(batch_size=1, columns=["SHAPE"]), None)
    if geometry_batch is None or geometry_batch.num_rows != 1:
        raise ValueError(f"GeoParquet Censo '{path.name}' sin geometría legible.")
    geometry = geometry_batch.column(0)[0].as_py()
    if not geometry:
        raise ValueError(f"GeoParquet Censo '{path.name}' con geometría vacía.")
    try:
        wkb.loads(geometry)
    except Exception as error:
        raise ValueError(f"GeoParquet Censo '{path.name}' con geometría WKB inválida.") from error


def validate_censo_dictionary(path: Path) -> None:
    """Valida que el diccionario oficial sea un XLSX legible con las hojas esperadas."""
    if not path.is_file() or path.stat().st_size == 0:
        raise ValueError("Diccionario de variables geográficas del Censo inexistente o vacío.")
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        sheets = set(workbook.sheetnames)
    finally:
        workbook.close()
    missing = REQUIRED_DICTIONARY_SHEETS.difference(sheets)
    if missing:
        raise ValueError(f"Diccionario del Censo sin hojas requeridas: {sorted(missing)}")


def _is_target_valid(name: str, raw_path: Path) -> bool:
    if not raw_path.is_file() or raw_path.stat().st_size == 0:
        return False
    try:
        if name == "Diccionario":
            validate_censo_dictionary(raw_path)
        else:
            validate_censo_geoparquet(raw_path)
        return True
    except Exception:
        return False


def _publish_candidate(candidate: Path, raw_path: Path, *, source_id: str, source_url: str,
                       zip_path: Path, archive_member: str, manifest_path: Path) -> None:
    """Publica candidato y provenance como una unidad con rollback del RAW previo."""
    raw_path.parent.mkdir(parents=True, exist_ok=True)
    backup = candidate.parent / f"previous_{raw_path.name}"
    had_previous = raw_path.exists()
    try:
        if had_previous:
            os.replace(raw_path, backup)
        os.replace(candidate, raw_path)
        record_raw_snapshot(
            source_id, source_url, raw_path,
            transport={"filename": zip_path.name, "size": zip_path.stat().st_size,
                       "sha256": file_sha256(zip_path), "archive_member": archive_member},
            source_metadata={"dataset": "censo", "year": 2024}, manifest_path=manifest_path,
        )
    except Exception:
        raw_path.unlink(missing_ok=True)
        if had_previous and backup.exists():
            os.replace(backup, raw_path)
        raise
    finally:
        backup.unlink(missing_ok=True)


def ingest_censo(*, force: bool = False, url: str = URL_CENSO,
                 raw_paths: dict[str, Path] | None = None,
                 dictionary_path: Path = RAW_DICTIONARY_PATH,
                 cache_dir: Path = CACHE_DIR,
                 manifest_path: Path = MANIFEST_PATH) -> dict[str, bool]:
    """Descarga, valida y publica las capas de cartografía Censo 2024 y su diccionario.

    Devuelve un dict por objetivo (5 capas + diccionario) indicando si se publicó
    un snapshot nuevo (True) o si se mantuvo el existente por SKIP (False).
    Solo descarga el ZIP si al menos un objetivo requiere publicación.
    """
    raw_paths = dict(raw_paths) if raw_paths is not None else dict(RAW_PATHS)
    targets: dict[str, tuple[str, Path]] = {name: (LAYER_FILE_NAMES[name], path) for name, path in raw_paths.items()}
    targets["Diccionario"] = (DICTIONARY_FILE_NAME, dictionary_path)

    published = {name: False for name in targets}
    pending = {
        name: (member_name, path) for name, (member_name, path) in targets.items()
        if force or not _is_target_valid(name, path)
    }

    if not pending:
        logger.info("[SKIP] RAW Censo 2024 (cartografía + diccionario) consolidado y válido.")
        return published

    cache_dir.mkdir(parents=True, exist_ok=True)
    zip_path = cache_dir / "Cartografia_censo2024_Pais.zip"
    all_succeeded = False
    try:
        _download_zip(url, zip_path)
        archive = _validate_transport(zip_path)
        try:
            with TemporaryDirectory(prefix="censo_2024_", dir=cache_dir) as temporary:
                staging = Path(temporary)
                for name, (member_name, raw_path) in pending.items():
                    member = _find_member(archive, member_name)
                    candidate = staging / member_name
                    with archive.open(member) as source, candidate.open("wb") as output:
                        shutil.copyfileobj(source, output, length=1024 * 1024)
                    if name == "Diccionario":
                        validate_censo_dictionary(candidate)
                    else:
                        validate_censo_geoparquet(candidate)
                    _publish_candidate(
                        candidate, raw_path, source_id=_SOURCE_IDS[name], source_url=url,
                        zip_path=zip_path, archive_member=member.filename, manifest_path=manifest_path,
                    )
                    published[name] = True
                    logger.info("RAW Censo publicado: %s", raw_path.as_posix())
        finally:
            archive.close()
        all_succeeded = True
        return published
    finally:
        # Un transporte fallido se conserva fuera de RAW para diagnóstico.
        if all_succeeded:
            zip_path.unlink(missing_ok=True)


def main() -> None:
    parser = argparse.ArgumentParser(description="Descarga segura de cartografía Censo 2024 (comunal y subcomunal).")
    parser.add_argument("--force", action="store_true", help="Solicita explícitamente reemplazar los snapshots consolidados.")
    args = parser.parse_args()
    try:
        ingest_censo(force=args.force)
    except Exception as error:
        logger.error("Fallo de ingesta Censo: %s", error)
        raise SystemExit(1) from error


if __name__ == "__main__":
    main()
