"""Ingesta segura de la dotación de camas hospitalarias (DEIS, DIGERA).

Fuente oficial: Departamento de Estadísticas e Información de Salud (DEIS) del
Ministerio de Salud, con la División de Gestión de la Red Asistencial (DIGERA).
Se publican dos libros XLSX independientes:

- Establecimientos pertenecientes al SNSS (públicos): una sola hoja
  ``2010-2025`` por Servicio de Salud, establecimiento y área funcional, con una
  columna por año.
- Establecimientos no pertenecientes al SNSS (privados): una hoja por año
  (``2019``-``2025``) por SEREMI, establecimiento y área. La propia fuente los
  rotula como "dato referencial" y cita a la Superintendencia de Salud.

Es dotación anual (stock de camas declarado por año), no ocupación ni
disponibilidad en tiempo real. Ambos libros cubren hasta 2025, año cerrado. Si el
DEIS publica una versión con otro rango de años (se infiere que con otro nombre
de archivo, por el patrón de las URLs actuales; no está comprobado), se incorpora
cambiando esta configuración, no sobrescribiendo en silencio.
"""

from __future__ import annotations

import argparse
from dataclasses import dataclass
import json
import logging
import os
from pathlib import Path
import shutil
import ssl
import tempfile
from typing import Any, Callable
import urllib.request

from openpyxl import load_workbook

from src.data.raw_provenance import file_sha256, record_raw_snapshot


logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
logger = logging.getLogger(__name__)

DEST_DIR = Path("data/raw/camas")
CACHE_DIR = Path(".cache/downloads/camas")
MANIFEST_PATH = Path("data/raw/provenance_manifest.json")

_BASE_URL = "https://repositoriodeis.minsal.cl/ContenidoSitioWeb2020/Estadisticas/DOTACION_CAMAS/"

HEADER_ROW = 3  # filas 1-2 son título y fuente; la 3 es el encabezado
SNSS_SHEET = "2010-2025"
SNSS_YEARS = tuple(range(2010, 2026))
SNSS_COLUMNS = (
    "Nombre Servicio de Salud", "Código Establecimiento", "Tipo de Establecimiento",
    "Nombre Establecimiento", "Cod. Área Funcional", "Área Funcional",
)
SNSS_MIN_ROWS = 1000
NO_SNSS_YEARS = tuple(range(2019, 2026))
NO_SNSS_COLUMNS = (
    "Nombre SEREMI", "Código establecimiento", "Nombre establecimiento",
    "Código área", "Area",
)
NO_SNSS_MIN_ROWS = 500


def _normalize(value: Any) -> str:
    """Normaliza un encabezado para compararlo: sin espacios extremos ni asteriscos de nota."""
    return str(value).strip().rstrip("*").strip() if value is not None else ""


def _header_of(sheet: Any, sheet_name: str) -> list[str]:
    rows = sheet.iter_rows(min_row=HEADER_ROW, max_row=HEADER_ROW, values_only=True)
    header = next(rows, None)
    if header is None:
        raise ValueError(f"Hoja {sheet_name!r} sin fila de encabezado.")
    return [_normalize(cell) for cell in header]


def _count_rows(sheet: Any) -> int:
    return sum(1 for row in sheet.iter_rows(values_only=True) if any(cell is not None for cell in row))


def validate_camas_snss_workbook(path: Path) -> None:
    """Valida legibilidad, hoja, encabezado y años del libro de establecimientos SNSS."""
    if not path.is_file() or path.stat().st_size == 0:
        raise ValueError("XLSX de dotación de camas SNSS inexistente o vacío.")
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if SNSS_SHEET not in workbook.sheetnames:
            raise ValueError(f"XLSX de camas SNSS sin la hoja requerida: {SNSS_SHEET!r}")
        sheet = workbook[SNSS_SHEET]
        header = _header_of(sheet, SNSS_SHEET)
        if tuple(header[:len(SNSS_COLUMNS)]) != SNSS_COLUMNS:
            raise ValueError(f"Encabezado SNSS inesperado: {header[:len(SNSS_COLUMNS)]}")
        missing = [str(year) for year in SNSS_YEARS if str(year) not in header]
        if missing:
            raise ValueError(f"Años ausentes en el libro de camas SNSS: {missing}")
        if _count_rows(sheet) < SNSS_MIN_ROWS:
            raise ValueError("Hoja de camas SNSS con menos filas de las esperadas.")
    finally:
        workbook.close()


def validate_camas_no_snss_workbook(path: Path) -> None:
    """Valida legibilidad y que cada hoja anual 2019-2025 traiga su encabezado y total."""
    if not path.is_file() or path.stat().st_size == 0:
        raise ValueError("XLSX de dotación de camas no SNSS inexistente o vacío.")
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        missing = [str(year) for year in NO_SNSS_YEARS if str(year) not in workbook.sheetnames]
        if missing:
            raise ValueError(f"Hojas anuales ausentes en el libro de camas no SNSS: {missing}")
        for year in NO_SNSS_YEARS:
            sheet = workbook[str(year)]
            header = _header_of(sheet, str(year))
            expected = (*NO_SNSS_COLUMNS, f"Total {year}")
            if tuple(header[:len(expected)]) != expected:
                raise ValueError(f"Encabezado no SNSS inesperado en hoja {year}: {header[:len(expected)]}")
            if _count_rows(sheet) < NO_SNSS_MIN_ROWS:
                raise ValueError(f"Hoja {year} de camas no SNSS con menos filas de las esperadas.")
    finally:
        workbook.close()


@dataclass(frozen=True)
class CamasSource:
    """Una publicación de dotación de camas y cómo validarla y registrarla."""

    key: str
    source_id: str
    url: str
    filename: str
    validator: Callable[[Path], None]
    source_metadata: dict[str, str]


CAMAS_SOURCES: tuple[CamasSource, ...] = (
    CamasSource(
        key="snss",
        source_id="deis_dotacion_camas_snss",
        url=_BASE_URL + "Dotaci%C3%B3n%20de%20camas%202010-2025%20Establecimientos%20"
                        "Pertenecientes%20al%20SNSS_01-12-2025.xlsx",
        filename="dotacion_camas_snss_2010_2025.xlsx",
        validator=validate_camas_snss_workbook,
        source_metadata={
            "dataset": "deis_dotacion_camas_establecimientos_snss",
            "institucion": "Ministerio de Salud, DEIS y DIGERA",
            "alcance": "Establecimientos pertenecientes al SNSS (públicos)",
            "periodo": "2010-2025",
            # Se conserva literal el sufijo del nombre de archivo de origen; no se interpreta.
            "sufijo_nombre_archivo_origen": "01-12-2025",
        },
    ),
    CamasSource(
        key="no_snss",
        source_id="deis_dotacion_camas_no_snss",
        url=_BASE_URL + "Dotaci%C3%B3n%20de%20camas%202019-2025%20Establecimientos%20"
                        "No%20pertenecientes%20al%20SNSS.xlsx",
        filename="dotacion_camas_no_snss_2019_2025.xlsx",
        validator=validate_camas_no_snss_workbook,
        source_metadata={
            "dataset": "deis_dotacion_camas_establecimientos_no_snss",
            "institucion": "Ministerio de Salud, DEIS y DIGERA, con Superintendencia de Salud",
            "alcance": "Establecimientos no pertenecientes al SNSS (privados); dato referencial según la fuente",
            "periodo": "2019-2025",
        },
    ),
)


def _download_xlsx(url: str, cache_path: Path) -> dict[str, str]:
    """Descarga el XLSX a cache (nunca a ``data/raw``) y devuelve cabeceras HTTP observadas.

    Usa verificación TLS estándar: el endpoint DEIS de camas la supera.
    """
    cache_path.parent.mkdir(parents=True, exist_ok=True)
    partial_path = cache_path.with_suffix(f"{cache_path.suffix}.part")
    request = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    observed: dict[str, str] = {}
    try:
        with urllib.request.urlopen(request, context=ssl.create_default_context()) as response, \
                partial_path.open("wb") as output:
            headers = getattr(response, "headers", None)
            if headers is not None:
                for field, key in (("Last-Modified", "http_last_modified"), ("ETag", "http_etag")):
                    value = headers.get(field)
                    if value:
                        observed[key] = value
            shutil.copyfileobj(response, output, length=1024 * 1024)
        if partial_path.stat().st_size == 0:
            raise ValueError("XLSX de dotación de camas descargado vacío.")
        os.replace(partial_path, cache_path)
    finally:
        partial_path.unlink(missing_ok=True)
    return observed


def _snapshot_registered(source_id: str, sha256: str, manifest_path: Path) -> bool:
    if not manifest_path.exists():
        return False
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    return any(
        snapshot.get("raw_sha256") == sha256
        for entry in manifest if entry.get("source_id") == source_id
        for snapshot in entry.get("snapshots", [])
    )


def _publish_candidate(
    candidate: Path, raw_path: Path, source: CamasSource, *, http_headers: dict[str, str],
    manifest_path: Path,
) -> None:
    """Publica el candidato validado y su provenance como una unidad, con rollback del RAW previo."""
    raw_path.parent.mkdir(parents=True, exist_ok=True)
    backup = candidate.parent / f"previous_{raw_path.name}"
    had_previous = raw_path.exists()
    try:
        if had_previous:
            os.replace(raw_path, backup)
        os.replace(candidate, raw_path)
        record_raw_snapshot(
            source.source_id, source.url, raw_path,
            transport={"filename": raw_path.name, "size": raw_path.stat().st_size,
                       "sha256": file_sha256(raw_path), **http_headers},
            source_metadata=source.source_metadata,
            manifest_path=manifest_path,
        )
    except Exception:
        raw_path.unlink(missing_ok=True)
        if had_previous and backup.exists():
            os.replace(backup, raw_path)
        raise
    finally:
        backup.unlink(missing_ok=True)


def ingest_camas_source(
    source: CamasSource, *, force: bool = False, dest_dir: Path = DEST_DIR,
    cache_dir: Path = CACHE_DIR, manifest_path: Path = MANIFEST_PATH,
) -> bool:
    """Descarga, valida y publica una fuente; devuelve si publicó un snapshot nuevo."""
    raw_path = dest_dir / source.filename
    if raw_path.is_file() and raw_path.stat().st_size > 0 and not force:
        source.validator(raw_path)
        logger.info("[SKIP] RAW de camas (%s) consolidado y válido: %s", source.key, raw_path.as_posix())
        return False

    cache_dir.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(delete=False, dir=cache_dir, prefix=f"camas_{source.key}_",
                                     suffix=".xlsx") as temp:
        candidate = Path(temp.name)
    try:
        http_headers = _download_xlsx(source.url, candidate)
        source.validator(candidate)
        candidate_hash = file_sha256(candidate)
        if raw_path.is_file() and file_sha256(raw_path) == candidate_hash:
            # Misma versión que el RAW vigente: no se reemplaza ni se duplica el historial.
            if not _snapshot_registered(source.source_id, candidate_hash, manifest_path):
                record_raw_snapshot(
                    source.source_id, source.url, raw_path,
                    transport={"filename": raw_path.name, "size": raw_path.stat().st_size,
                               "sha256": candidate_hash, **http_headers},
                    source_metadata=source.source_metadata, manifest_path=manifest_path,
                )
            logger.info("RAW de camas (%s) sin cambios respecto de la fuente.", source.key)
            return False
        _publish_candidate(candidate, raw_path, source, http_headers=http_headers,
                           manifest_path=manifest_path)
        logger.info("RAW de camas (%s) publicado: %s", source.key, raw_path.as_posix())
        return True
    finally:
        candidate.unlink(missing_ok=True)


def ingest_all(
    *, force: bool = False, sources: tuple[CamasSource, ...] = CAMAS_SOURCES,
    dest_dir: Path = DEST_DIR, cache_dir: Path = CACHE_DIR, manifest_path: Path = MANIFEST_PATH,
) -> dict[str, str]:
    """Procesa cada fuente de forma independiente; devuelve el error por fuente fallida."""
    errors: dict[str, str] = {}
    for source in sources:
        try:
            ingest_camas_source(source, force=force, dest_dir=dest_dir, cache_dir=cache_dir,
                                manifest_path=manifest_path)
        except Exception as error:  # una fuente fallida no impide publicar la otra
            logger.error("Fallo de ingesta de camas (%s): %s", source.key, error)
            errors[source.key] = str(error)
    return errors


def main() -> None:
    parser = argparse.ArgumentParser(description="Descarga segura de la dotación de camas hospitalarias (DEIS).")
    parser.add_argument("--force", action="store_true",
                        help="Solicita explícitamente un nuevo snapshot de ambas fuentes.")
    args = parser.parse_args()
    if ingest_all(force=args.force):
        raise SystemExit(1)


if __name__ == "__main__":
    main()
