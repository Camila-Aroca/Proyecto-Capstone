"""Construye la dimensión de población censada 2024 por comuna de la RM.

Extrae el cuadro comunal oficial (hoja "2") del tabulado INE descargado por
``download_censo_poblacion`` y lo contrasta contra el total regional
independiente publicado en la hoja "1" del mismo workbook. Esta dimensión es
un insumo territorial estático (2024): no debe usarse como proxy poblacional
para otros años ni combinarse aquí con las series históricas de Urgencias.
"""

from __future__ import annotations

import argparse
import logging
from pathlib import Path
import tempfile
from typing import Any

from openpyxl import load_workbook
import pyarrow as pa
import pyarrow.parquet as pq

from src.data.clean_censo_comunas import EXPECTED_RM_CUTS


logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
logger = logging.getLogger(__name__)

RAW_PATH = Path("data/raw/censo/D1_Poblacion-censada-por-sexo-y-edad-en-grupos-quinquenales.xlsx")
OUTPUT_PATH = Path("data/processed/censo/dim_poblacion_comuna_censo2024.parquet")

REGION_METROPOLITANA_CODIGO = 13
ANO_REFERENCIA = 2024
FUENTE = (
    "INE, Censo de Poblacion y Vivienda 2024, cuadro 'Poblacion censada por sexo y "
    "razon hombre-mujer, segun comuna' "
    "(D1_Poblacion-censada-por-sexo-y-edad-en-grupos-quinquenales.xlsx, hoja 2)"
)

SCHEMA = pa.schema([
    ("comuna_codigo", pa.string()),
    ("comuna_glosa", pa.string()),
    ("ano_referencia", pa.int32()),
    ("poblacion_censada", pa.int64()),
    ("fuente", pa.string()),
])


def _read_comunal_rm_rows(raw_path: Path) -> list[dict[str, Any]]:
    """Lee la hoja 2 (comuna) y conserva únicamente las filas de la RM."""
    workbook = load_workbook(raw_path, read_only=True, data_only=True)
    try:
        sheet = workbook["2"]
        records: list[dict[str, Any]] = []
        for row in sheet.iter_rows(min_row=5, max_col=7, values_only=True):
            cod_region, _region, _cod_provincia, _provincia, cod_comuna, comuna, poblacion = row
            if cod_region != REGION_METROPOLITANA_CODIGO:
                continue
            records.append({
                "comuna_codigo": str(int(cod_comuna)),
                "comuna_glosa": comuna,
                "ano_referencia": ANO_REFERENCIA,
                "poblacion_censada": int(poblacion) if poblacion is not None else None,
                "fuente": FUENTE,
            })
        return records
    finally:
        workbook.close()


def _read_region_total(raw_path: Path) -> int:
    """Lee el total regional independiente de la hoja 1 para validación cruzada no circular."""
    workbook = load_workbook(raw_path, read_only=True, data_only=True)
    try:
        sheet = workbook["1"]
        for row in sheet.iter_rows(min_row=5, max_col=3, values_only=True):
            cod_region, _region, poblacion = row
            if cod_region == REGION_METROPOLITANA_CODIGO:
                return int(poblacion)
        raise ValueError("No se encontró el total regional de la RM en la hoja 1 del tabulado.")
    finally:
        workbook.close()


def _validate_records(records: list[dict[str, Any]], region_total: int) -> None:
    if len(records) != 52:
        raise ValueError(f"Se esperaban 52 comunas RM en el tabulado comunal, se obtuvieron {len(records)}.")

    codes = [record["comuna_codigo"] for record in records]
    if len(set(codes)) != len(codes):
        raise ValueError("Códigos de comuna duplicados en el tabulado comunal de población.")

    observed_cuts = {int(code) for code in codes}
    if observed_cuts != EXPECTED_RM_CUTS:
        missing = EXPECTED_RM_CUTS - observed_cuts
        extra = observed_cuts - EXPECTED_RM_CUTS
        raise ValueError(f"Comunas RM inconsistentes con el CUT oficial. Faltantes: {sorted(missing)}, "
                          f"inesperadas: {sorted(extra)}.")

    for record in records:
        poblacion = record["poblacion_censada"]
        if poblacion is None or poblacion <= 0:
            raise ValueError(f"Población censada nula o no positiva para {record['comuna_codigo']}.")
        if not record["comuna_glosa"] or not str(record["comuna_glosa"]).strip():
            raise ValueError(f"Glosa comunal vacía para {record['comuna_codigo']}.")
        if record["ano_referencia"] != ANO_REFERENCIA:
            raise ValueError("ano_referencia debe ser 2024; no usar 2024 como proxy de otros años.")

    comuna_total = sum(record["poblacion_censada"] for record in records)
    if comuna_total != region_total:
        raise ValueError(
            f"Suma comunal RM ({comuna_total}) no coincide con el total regional publicado "
            f"independientemente en la hoja 1 ({region_total})."
        )


def build_dim_poblacion_comuna_censo2024(
    raw_path: Path = RAW_PATH, output_path: Path = OUTPUT_PATH,
) -> dict[str, Any]:
    """Genera `dim_poblacion_comuna_censo2024.parquet` de forma validada y atómica."""
    if not raw_path.exists():
        raise FileNotFoundError(f"No existe el RAW requerido: {raw_path.as_posix()}")

    records = _read_comunal_rm_rows(raw_path)
    region_total = _read_region_total(raw_path)
    _validate_records(records, region_total)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    table = pa.Table.from_pylist(records, schema=SCHEMA)
    with tempfile.NamedTemporaryFile(
        delete=False, dir=output_path.parent, prefix=f"{output_path.name}.", suffix=".tmp"
    ) as temp:
        temp_path = Path(temp.name)
    try:
        pq.write_table(table, temp_path, compression="snappy")
        temp_path.replace(output_path)
    except Exception:
        temp_path.unlink(missing_ok=True)
        raise

    verify_table = pq.read_table(output_path)
    return {
        "archivo_generado": output_path.as_posix(),
        "filas": verify_table.num_rows,
        "poblacion_total_rm": region_total,
        "comunas_unicas": len(set(verify_table["comuna_codigo"].to_pylist())),
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Construye la dimensión de población comunal Censo 2024 (RM).")
    parser.add_argument("--force", action="store_true", help="Aceptado para interoperar con el orquestador.")
    parser.parse_args()
    result = build_dim_poblacion_comuna_censo2024()
    logger.info("Dimensión de población comunal generada: %s", result)


if __name__ == "__main__":
    main()
