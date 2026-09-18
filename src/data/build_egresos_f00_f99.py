"""Construye el dataset analítico canónico de Egresos Hospitalarios F00-F99
(capítulo CIE-10 "Trastornos mentales y del comportamiento"), 2020-2025.

Grano: 1 fila = 1 registro de egreso publicado por DEIS (dataset disociado y
anonimizado; no existe identificador de paciente). No se deduplica por
coincidencia de atributos: dos registros con todos los campos iguales son
dos egresos distintos que coinciden por azar en una población grande, no un
error del pipeline (ver `reports/eda/eda_egresos_encoding_detalle.md`,
sección 9).

Alcance geográfico: **nacional**, con la columna derivada `residente_rm`
(booleana, nula si `region_residencia` es nulo en origen) en lugar de un
filtro RM. Se prioriza esta opción sobre filtrar directamente a RM porque:
(a) `region_residencia`/`comuna_residencia` son residencia del paciente, no
ubicación del hospital (un residente de RM puede egresar de un hospital
fuera de RM y viceversa); (b) el capítulo F00-F99 es apenas el 2,3% de los
egresos nacionales 2020-2025, y limitar además a residencia RM reduciría la
muestra a un tercio adicional sin beneficio para la reproducibilidad; y (c)
mantener el universo nacional permite comparar RM contra el resto del país
sin tener que reconstruir el dataset. `residente_rm` NUNCA debe leerse como
"hospitalizado en la Región Metropolitana".

Filtro F00-F99: se aplica exclusivamente sobre `diag1` (diagnóstico
principal), contra el catálogo oficial CIE-10 extraído en tiempo de
ejecución de `data/raw/egresos/Diccionario BD egresos hospitalario.xlsx`
(hoja "codigo CIE-10", columna `CAPITULO` = "F00-F99"; 407 subcategorías).
Se verificó por código que este filtro por catálogo es idéntico, sin una
sola discrepancia, a `diag1` que comienza con "F" sobre las 9.379.787 filas
nacionales 2020-2025. `diag2` (causa externa) NO se usa para el filtro y
está vacío ("") en el 100% de los registros F00-F99 observados.

No se homologan silenciosamente heterogeneidades conocidas entre años (ver
`SEXO_DOMAIN`/`GRUPO_EDAD_DOMAIN` y el diccionario/EDA de este dataset):
- `sexo`: 2021 publica códigos numéricos como texto ("1","2","3"); el resto
  de los años publica texto ("HOMBRE","MUJER","INTERSEX (INDETERMINDADO)")
  más el marcador "*". El diccionario oficial no incluye una tabla de
  códigos para `SEXO`, por lo que no existe evidencia documental para
  mapear 2021 a las categorías de texto; se conserva tal cual.
- `grupo_edad`: 2021 usa un esquema quinquenal/fino distinto (mayúsculas,
  con tramos de días/meses para menores de 1 año) del esquema decenal en
  minúsculas que usan 2020 y 2022-2025. El tramo "85 A MAS" de 2021 no
  puede homologarse sin ambigüedad al esquema decenal (se solaparía entre
  "80 a 89" y "90 y más"), por lo que no se intenta ningún cruce.
- El mojibake U+FFFD en `grupo_edad`/`glosa_comuna_residencia`/
  `glosa_region_residencia` de 2024-2025 ya está documentado en
  `reports/eda/eda_egresos_encoding_detalle.md`; se conserva tal cual.

Todos los dominios categóricos observados (`sexo`, `pertenencia_establecimiento_salud`,
`grupo_edad`, `glosa_prevision`, `prevision`, `condicion_egreso`) se validan
contra listas explícitas construidas desde la evidencia 2020-2025: un valor
nuevo fuera de esas listas hace fallar la etapa en vez de pasar
silenciosamente, forzando su documentación explícita.
"""

from __future__ import annotations

from pathlib import Path
from typing import Final

import openpyxl
import pandas as pd
import pyarrow as pa
import pyarrow.compute as pc
import pyarrow.dataset as ds
import pyarrow.parquet as pq

INPUT_DIR: Final[Path] = Path("data/processed/egresos")
YEARS: Final[tuple[int, ...]] = tuple(range(2020, 2026))
DICCIONARIO_PATH: Final[Path] = Path("data/raw/egresos/Diccionario BD egresos hospitalario.xlsx")
OUTPUT_PATH: Final[Path] = INPUT_DIR / "egresos_f00_f99_nacional_2020_2025.parquet"
CATALOG_OUTPUT_PATH: Final[Path] = INPUT_DIR / "catalogo_cie10_f00_f99.csv"
RM_REGION_CODE: Final[int] = 13
EXPECTED_F00_F99_SUBCATEGORIAS: Final[int] = 407

COLUMNS: Final[tuple[str, ...]] = (
    "ano_egreso", "sexo", "grupo_edad",
    "region_residencia", "glosa_region_residencia",
    "comuna_residencia", "glosa_comuna_residencia",
    "prevision", "glosa_prevision",
    "pertenencia_establecimiento_salud",
    "diag1", "diag2",
    "dias_estada", "condicion_egreso",
    "interv_q", "proced", "error",
)

# Dominios observados y documentados sobre el subconjunto F00-F99 2020-2025
# (ver reports/hospitalization/eda_egresos_f00_f99.md). Un valor fuera de
# estos conjuntos debe forzar una revisión explícita del contrato, no pasar
# silenciosamente.
SEXO_DOMAIN: Final[set[str]] = {
    "*", "1", "2", "3", "9", "HOMBRE", "MUJER", "INTERSEX (INDETERMINDADO)",
}
PERTENENCIA_DOMAIN: Final[set[str]] = {
    "*",
    "Pertenecientes al Sistema Nacional de Servicios de Salud, SNSS",
    "No Pertenecientes al Sistema Nacional de Servicios de Salud, SNSS",
}
PREVISION_DOMAIN: Final[set[int]] = {1, 2, 3, 4, 5, 96, 99}
GLOSA_PREVISION_DOMAIN: Final[set[str]] = {
    "*", "FONASA", "ISAPRE", "CAPREDENA", "DIPRECA", "SISA", "NINGUNA", "DESCONOCIDO",
}
CONDICION_EGRESO_DOMAIN: Final[set[int]] = {1, 2}
# Unión de los dos esquemas de grupo_edad observados en F00-F99 2020-2025
# (decenal 2020/2022-2025 y quinquenal/fino 2021), incluidas las variantes
# con mojibake U+FFFD de 2024-2025 documentadas en
# reports/eda/eda_egresos_encoding_detalle.md.
GRUPO_EDAD_DOMAIN: Final[set[str]] = {
    "*",
    "1 A 4 AÑOS", "1 a 9", "10 A 14 AÑOS", "10 a 19", "15 A 19 AÑOS",
    "2 MESES A MENOS DE 1 AÑO", "20 A 24 AÑOS", "20 a 29", "25 A 29 AÑOS",
    "28 DIAS A 2 MES", "30 A 34 AÑOS", "30 a 39", "35 A 39 AÑOS",
    "40 A 44 AÑOS", "40 a 49", "45 A 49 AÑOS", "5 A 9 AÑOS", "50 A 54 AÑOS",
    "50 a 59", "55 A 59 AÑOS", "60 A 64 AÑOS", "60 a 69", "65 A 69 AÑOS",
    "7 A 27 DIAS", "70 A 74 AÑOS", "70 a 79", "75 A 79 AÑOS", "80 A 84 AÑOS",
    "80 a 89", "85 A MAS", "90 y más", "90 y m�s", "menor a 7 días",
    "menor de un año", "menor de un a�o",
}


def load_catalog_f00_f99(path: Path = DICCIONARIO_PATH) -> pd.DataFrame:
    """Extrae el subconjunto CIE-10 F00-F99 del diccionario oficial DEIS de Egresos.

    Falla si la hoja "codigo CIE-10" no tiene la estructura esperada o si el
    número de subcategorías F00-F99 cambia respecto del contrato vigente
    (407), en vez de propagar silenciosamente un catálogo distinto.
    """
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if "codigo CIE-10" not in workbook.sheetnames:
        raise ValueError(f"Hoja 'codigo CIE-10' no encontrada en {path.as_posix()}.")
    rows = list(workbook["codigo CIE-10"].iter_rows(values_only=True))
    header_idx = next(
        (i for i, row in enumerate(rows) if row and row[0] == "CODIGO SUBCATEGORIA"), None
    )
    if header_idx is None:
        raise ValueError("No se encontró la fila de encabezado 'CODIGO SUBCATEGORIA'.")
    header = [str(cell).strip() for cell in rows[header_idx]]
    catalog = pd.DataFrame(rows[header_idx + 1 :], columns=header)
    catalog = catalog[catalog["CAPITULO"] == "F00-F99"].copy()
    catalog["CODIGO SUBCATEGORIA"] = catalog["CODIGO SUBCATEGORIA"].str.strip().str.upper()
    if catalog["CODIGO SUBCATEGORIA"].duplicated().any():
        raise ValueError("Catálogo CIE-10 F00-F99 con códigos de subcategoría duplicados.")
    if len(catalog) != EXPECTED_F00_F99_SUBCATEGORIAS:
        raise ValueError(
            f"El catálogo F00-F99 extraído tiene {len(catalog)} filas; se esperaban "
            f"{EXPECTED_F00_F99_SUBCATEGORIAS} según el diccionario DEIS vigente."
        )
    return catalog.reset_index(drop=True)


def load_source_year(year: int, input_dir: Path = INPUT_DIR) -> pa.Table:
    """Lee un año de Egresos ya normalizados por `clean_egresos`, sin filtrar."""
    path = input_dir / f"egresos_{year}.parquet"
    if not path.is_file():
        raise FileNotFoundError(f"Parquet de Egresos no disponible: {path.as_posix()}")
    return pq.read_table(path, columns=list(COLUMNS))


def filter_f00_f99(table: pa.Table, catalog_codes: set[str]) -> pa.Table:
    """Filtra por diagnóstico principal (`diag1`) perteneciente al capítulo F00-F99."""
    codes_array = pa.array(sorted(catalog_codes), type=pa.string())
    mask = pc.is_in(table["diag1"], value_set=codes_array)
    return table.filter(mask)


def add_residente_rm_flag(table: pa.Table) -> pa.Table:
    """Agrega `residente_rm` (bool, nulo si `region_residencia` es nulo en origen).

    No representa la ubicación del establecimiento hospitalario: Egresos no
    tiene identificador de establecimiento, solo residencia del paciente.
    """
    region = table["region_residencia"]
    is_rm = pc.equal(region, pa.scalar(RM_REGION_CODE, type=region.type))
    return table.append_column("residente_rm", is_rm)


def independent_total_count(catalog_codes: set[str], years: tuple[int, ...] = YEARS,
                             input_dir: Path = INPUT_DIR) -> int:
    """Recuento F00-F99 vía un camino de código independiente (dataset multi-archivo).

    Sirve como control cruzado frente al recuento año-por-año usado para
    construir el dataset final, para detectar duplicación/pérdida accidental
    introducida por el pipeline.
    """
    files = [input_dir / f"egresos_{year}.parquet" for year in years]
    dataset = ds.dataset(files, format="parquet")
    table = dataset.to_table(columns=["diag1"])
    codes_array = pa.array(sorted(catalog_codes), type=pa.string())
    mask = pc.is_in(table["diag1"], value_set=codes_array)
    return int(pc.sum(mask.cast(pa.int64())).as_py())


def validate_diag1_domain(table: pa.Table, catalog_codes: set[str]) -> None:
    """Verifica que todo `diag1` esté en el catálogo F00-F99 (nada fuera de contrato)."""
    observed = set(table["diag1"].to_pylist())
    fuera = observed - catalog_codes
    if fuera:
        raise ValueError(f"diag1 fuera del catálogo F00-F99: {sorted(fuera)[:10]}")


def validate_years(table: pa.Table, years: tuple[int, ...] = YEARS) -> None:
    """Verifica que `ano_egreso` cubra exactamente 2020-2025, sin años extra ni faltantes."""
    observed = set(table["ano_egreso"].to_pylist())
    if observed != set(years):
        raise ValueError(f"Años observados {sorted(observed)} != contrato {sorted(years)}.")


def validate_no_pipeline_duplication(
    per_year_counts: dict[int, int], table: pa.Table, expected_total: int
) -> None:
    """Verifica que concatenar por año no haya introducido/perdido filas."""
    if table.num_rows != sum(per_year_counts.values()):
        raise ValueError(
            f"Filas finales ({table.num_rows}) != suma de filtrados por año "
            f"({sum(per_year_counts.values())})."
        )
    if table.num_rows != expected_total:
        raise ValueError(
            f"Filas finales ({table.num_rows}) != recuento independiente "
            f"({expected_total}); posible duplicación/pérdida introducida por el pipeline."
        )
    observed_counts = pd.Series(table["ano_egreso"].to_pylist()).value_counts().to_dict()
    if observed_counts != per_year_counts:
        raise ValueError("Conteo por año no coincide entre el filtrado por año y el dataset final.")


def validate_dias_estada(table: pa.Table) -> None:
    """Verifica que `dias_estada` no tenga nulos ni valores no positivos (sin modificarla)."""
    dias = table["dias_estada"]
    if pc.any(pc.is_null(dias)).as_py():
        raise ValueError("dias_estada tiene nulos en el dataset F00-F99 (no observado en la fuente).")
    if pc.any(pc.less_equal(dias, 0)).as_py():
        raise ValueError("dias_estada <= 0 detectado (no observado en la fuente).")


def validate_whitelisted_domains(table: pa.Table) -> None:
    """Verifica que los campos categóricos no publiquen valores fuera de lo ya documentado."""
    text_checks = [
        ("sexo", SEXO_DOMAIN),
        ("pertenencia_establecimiento_salud", PERTENENCIA_DOMAIN),
        ("grupo_edad", GRUPO_EDAD_DOMAIN),
        ("glosa_prevision", GLOSA_PREVISION_DOMAIN),
    ]
    for column, domain in text_checks:
        observed = set(table[column].to_pylist())
        fuera = observed - domain
        if fuera:
            raise ValueError(f"Valores no documentados en {column}: {sorted(fuera)}")

    prevision_observed = {v for v in table["prevision"].to_pylist() if v is not None}
    if not prevision_observed <= PREVISION_DOMAIN:
        raise ValueError(
            f"Valores no documentados en prevision: {sorted(prevision_observed - PREVISION_DOMAIN)}"
        )
    condicion_observed = {v for v in table["condicion_egreso"].to_pylist() if v is not None}
    if not condicion_observed <= CONDICION_EGRESO_DOMAIN:
        raise ValueError(
            f"Valores no documentados en condicion_egreso: {sorted(condicion_observed - CONDICION_EGRESO_DOMAIN)}"
        )


def validate_grain(table: pa.Table) -> None:
    """Verifica que el dataset no esté vacío.

    El grano es 1 fila = 1 registro publicado; el dataset es disociado y no
    tiene llave natural de paciente, por lo que no se exige unicidad de
    atributos (ver duplicados naturales documentados en el EDA).
    """
    if table.num_rows == 0:
        raise ValueError("Dataset F00-F99 vacío.")


def atomic_write_table(table: pa.Table, output_path: Path) -> None:
    """Escribe un Parquet temporal y lo publica sólo si su esquema es legible."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary = output_path.with_suffix(f"{output_path.suffix}.tmp")
    temporary.unlink(missing_ok=True)
    pq.write_table(table, temporary, compression="snappy")
    if not pq.read_schema(temporary).names:
        raise ValueError(f"Parquet temporal sin columnas: {temporary.as_posix()}")
    temporary.replace(output_path)


def build_dataset() -> tuple[pa.Table, dict[int, int], pd.DataFrame]:
    """Filtra y consolida Egresos F00-F99 2020-2025 (nacional), año por año."""
    catalog = load_catalog_f00_f99()
    catalog_codes = set(catalog["CODIGO SUBCATEGORIA"])

    per_year_counts: dict[int, int] = {}
    filtered_tables: list[pa.Table] = []
    for year in YEARS:
        year_table = load_source_year(year)
        year_filtered = filter_f00_f99(year_table, catalog_codes)
        per_year_counts[year] = year_filtered.num_rows
        filtered_tables.append(year_filtered)

    combined = pa.concat_tables(filtered_tables)
    combined = add_residente_rm_flag(combined)
    return combined, per_year_counts, catalog


def main() -> None:
    """Genera, valida y publica el dataset canónico F00-F99 y su catálogo CIE-10."""
    combined, per_year_counts, catalog = build_dataset()
    catalog_codes = set(catalog["CODIGO SUBCATEGORIA"])
    expected_total = independent_total_count(catalog_codes)

    validate_diag1_domain(combined, catalog_codes)
    validate_years(combined)
    validate_no_pipeline_duplication(per_year_counts, combined, expected_total)
    validate_dias_estada(combined)
    validate_whitelisted_domains(combined)
    validate_grain(combined)

    atomic_write_table(combined, OUTPUT_PATH)

    CATALOG_OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    catalog.to_csv(CATALOG_OUTPUT_PATH, index=False, encoding="utf-8")

    print(
        "Dataset F00-F99 de Egresos generado: "
        f"{combined.num_rows:,} filas nacionales 2020-2025; "
        f"catálogo CIE-10 F00-F99: {len(catalog)} subcategorías."
    )


if __name__ == "__main__":
    main()
