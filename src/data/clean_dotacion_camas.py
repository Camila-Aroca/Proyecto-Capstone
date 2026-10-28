"""Limpieza de la dotación de camas hospitalarias (DEIS, DIGERA): RAW XLSX -> Parquet.

Cada libro de ``data/raw/camas/`` produce su propio Parquet canónico, con las
mismas columnas que el libro. Es dotación anual declarada, no ocupación ni
disponibilidad en tiempo real.

Salidas en ``data/processed/camas/`` (una por libro):

- ``dotacion_camas_snss.parquet``: establecimientos del SNSS (públicos), 2010-2025.
  Columnas del libro: Servicio de Salud, código, tipo y nombre del establecimiento,
  código y nombre del área funcional y una columna ``camas_AAAA`` por año. Una fila
  por fila del libro, incluida la fila ``Total`` de cada hospital (área 100).
- ``dotacion_camas_no_snss.parquet``: establecimientos no pertenecientes al SNSS
  (privados), 2019-2025; "dato referencial" según la fuente. El libro trae una hoja
  por año, y el Parquet las apila con una columna ``ano``. La fila de total del
  hospital no trae área en el libro, y queda con el área nula.

Solo datos de dotación. Quedan fuera del Parquet, y siguen intactos en el RAW: las
notas al pie, las marcas que remiten a ellas (asteriscos en nombres y encabezados,
``*******`` en celdas de valor, ``ok`` en columnas adicionales) y los totales
nacionales publicados. Estos últimos se usan para validar: la suma de los totales
por hospital debe coincidir con el total nacional del libro; una discrepancia se
informa, nunca se corrige.

Principios (``.agents/rules.md``): no se imputa ni se corrige nada. Un blanco queda
nulo y no es un cero; el cero explícito del origen se conserva como 0. Una celda con
marca de nota en lugar de una cifra (``*******``: "no declara") queda nula. Los
nombres se normalizan solo en espacios y asteriscos finales. Un hospital repetido
en una hoja no se elimina. Toda fila del origen debe responder a un patrón conocido;
si no, o si hay datos fuera de las columnas contempladas, o si una fila repite la
llave lógica de otra, la etapa falla en vez de descartar o elegir en silencio.
"""

from __future__ import annotations

import argparse
from dataclasses import dataclass, field
import logging
import os
from pathlib import Path
import re
import tempfile
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
import pyarrow as pa
import pyarrow.parquet as pq

from src.data import download_dotacion_camas as download
from src.data.raw_provenance import RAW_SHA256_METADATA_KEY, file_sha256


logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
logger = logging.getLogger(__name__)

_SOURCES = {source.key: source for source in download.CAMAS_SOURCES}
RAW_SNSS_PATH = download.DEST_DIR / _SOURCES["snss"].filename
RAW_NO_SNSS_PATH = download.DEST_DIR / _SOURCES["no_snss"].filename
OUTPUT_DIR = Path("data/processed/camas")
SNSS_FILENAME = "dotacion_camas_snss.parquet"
NO_SNSS_FILENAME = "dotacion_camas_no_snss.parquet"

FIRST_DATA_ROW = download.HEADER_ROW + 1
SNSS_TOTAL_AREA_CODE = 100  # el libro SNSS rotula "Total" como el área 100
SNSS_TOTAL_LABEL = "Total"
ZSTD_LEVEL = 19  # el formato largo repite etiquetas; con zstd pesa menos que el libro de origen

_YEAR_HEADER = re.compile(r"^(\d{4})\*{0,}$")
_NAME_MARKER = re.compile(r"^(.*?)\s*\*+$")
_MARKER_ONLY = re.compile(r"^\*+$")

# Cada libro tiene su esquema: las columnas del libro, con nombres de columna válidos.
SNSS_SCHEMA = pa.schema([
    ("servicio_salud", pa.string()),
    ("establecimiento_codigo", pa.int64()),
    ("tipo_establecimiento", pa.string()),
    ("establecimiento_nombre", pa.string()),
    ("area_funcional_codigo", pa.int32()),
    ("area_funcional", pa.string()),
    *[(f"camas_{year}", pa.int32()) for year in download.SNSS_YEARS],
])
NO_SNSS_SCHEMA = pa.schema([
    ("ano", pa.int32()),
    ("seremi", pa.string()),
    ("establecimiento_codigo", pa.int64()),
    ("establecimiento_nombre", pa.string()),
    ("area_funcional_codigo", pa.int32()),
    ("area_funcional", pa.string()),
    ("camas", pa.int32()),
])
SNSS_KEY = ("servicio_salud", "establecimiento_codigo", "area_funcional_codigo")
NO_SNSS_KEY = ("ano", "seremi", "establecimiento_codigo", "area_funcional_codigo")


@dataclass
class _Source:
    """Lo leído de un libro: filas del Parquet y lo necesario para validarlas."""

    rows: list[dict[str, Any]] = field(default_factory=list)
    national: dict[int, int | None] = field(default_factory=dict)
    hospital_totals: list[dict[str, Any]] = field(default_factory=list)


# --------------------------------------------------------------------------- #
# Utilidades de lectura
# --------------------------------------------------------------------------- #

def _clean(value: Any) -> str | None:
    """Colapsa espacios (incluido el espacio duro) y recorta; vacío -> None."""
    if value is None:
        return None
    text = re.sub(r"\s+", " ", str(value).replace("\xa0", " ")).strip()
    return text or None


def _name(value: Any) -> str | None:
    """Nombre sin espacios sobrantes ni asteriscos finales de nota."""
    text = _clean(value)
    if text is None:
        return None
    match = _NAME_MARKER.match(text)
    return (match.group(1).strip() or None) if match else text


def _has_content(value: Any) -> bool:
    return value is not None and not (isinstance(value, str) and not value.strip())


def _ref(column_index: int, row_number: int) -> str:
    return f"{get_column_letter(column_index + 1)}{row_number}"


def _read_rows(path: Path, sheet_name: str) -> list[tuple[Any, ...]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = [tuple(row) for row in workbook[sheet_name].iter_rows(min_row=1, values_only=True)]
    finally:
        workbook.close()
    width = max((len(row) for row in rows), default=0)
    return [row + (None,) * (width - len(row)) for row in rows]


def _parse_value(value: Any, ref: str) -> int | None:
    """Camas de una celda: blanco o marca de nota (``*****``) dan None; otra cosa que un entero >= 0, error."""
    if not _has_content(value):
        return None
    if isinstance(value, bool):
        raise ValueError(f"Celda {ref}: valor de camas inválido ({value!r}).")
    if isinstance(value, int):
        number = value
    elif isinstance(value, float) and value.is_integer():
        number = int(value)
    elif isinstance(value, str) and _MARKER_ONLY.match(value.strip()):
        return None
    else:
        raise ValueError(f"Celda {ref}: valor de camas inválido ({value!r}).")
    if number < 0:
        raise ValueError(f"Celda {ref}: valor de camas inválido ({value!r}).")
    return number


def _as_code(value: Any, ref: str) -> int:
    if isinstance(value, int) and not isinstance(value, bool):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    raise ValueError(f"Celda {ref}: código numérico esperado, se encontró {value!r}.")


def _unrecognized(sheet: str, row_number: int, row: tuple[Any, ...]) -> ValueError:
    return ValueError(
        f"Fila {row_number} de la hoja {sheet!r} con estructura no reconocida: {list(row[:8])!r}. "
        "Se detiene para no descartar datos en silencio."
    )


def _last_data_row(rows: list[tuple[Any, ...]]) -> int:
    """Número de la última fila con contenido en alguna columna distinta de A."""
    last = 0
    for row_number, row in enumerate(rows, start=1):
        if any(_has_content(cell) for cell in row[1:]):
            last = row_number
    return last


def _is_footnote(row: tuple[Any, ...], row_number: int, last_data_row: int) -> bool:
    """Texto solo en la columna A después de la última fila de datos: nota al pie."""
    return (isinstance(row[0], str) and _has_content(row[0])
            and not any(_has_content(cell) for cell in row[1:]) and row_number > last_data_row)


# --------------------------------------------------------------------------- #
# Libro SNSS
# --------------------------------------------------------------------------- #

def _snss_year_columns(rows: list[tuple[Any, ...]]) -> list[tuple[int, int]]:
    """(índice de columna, año) desde el encabezado; ``2014*`` se lee como 2014."""
    header = rows[download.HEADER_ROW - 1]
    first = len(download.SNSS_COLUMNS)
    columns: list[tuple[int, int]] = []
    for index in range(first, first + len(download.SNSS_YEARS)):
        cell = header[index]
        match = _YEAR_HEADER.match(str(cell).strip()) if _has_content(cell) else None
        if match is None:
            raise ValueError(f"Encabezado de año no reconocido en {_ref(index, download.HEADER_ROW)}: {cell!r}")
        columns.append((index, int(match.group(1))))
    if [year for _, year in columns] != list(download.SNSS_YEARS):
        raise ValueError("Los años del encabezado SNSS no coinciden con los esperados.")
    return columns


def _parse_snss(path: Path) -> _Source:
    sheet = download.SNSS_SHEET
    rows = _read_rows(path, sheet)
    year_columns = _snss_year_columns(rows)
    last_column = year_columns[-1][0]
    last_data_row = _last_data_row(rows)
    source = _Source()
    block: dict[str, Any] | None = None

    for row_number, row in enumerate(rows[FIRST_DATA_ROW - 1:], start=FIRST_DATA_ROW):
        if any(_has_content(cell) for cell in row[last_column + 1:]):
            raise ValueError(f"Fila {row_number} de la hoja {sheet!r} con datos fuera de las columnas contempladas.")
        if not any(_has_content(cell) for cell in row):
            continue
        a, b, c, d, e, f = row[:6]

        if not any(_has_content(cell) for cell in (a, b, c, d, e)) and _has_content(f):
            if (_clean(f) or "").casefold() != "total país":
                raise _unrecognized(sheet, row_number, row)
            for index, year in year_columns:
                source.national[year] = _parse_value(row[index], _ref(index, row_number))
            continue
        if _is_footnote(row, row_number, last_data_row):
            continue  # nota al pie: no es dato de dotación

        is_total = (all(_has_content(cell) for cell in (a, b, c, d, e, f))
                    and isinstance(e, int) and e == SNSS_TOTAL_AREA_CODE
                    and (_clean(f) or "").casefold() == SNSS_TOTAL_LABEL.casefold())
        is_area = (_has_content(a) and not _has_content(b) and not _has_content(c)
                   and _has_content(d) and _has_content(e) and _has_content(f))
        if is_total:
            block = {"id": row_number, "servicio": _clean(a), "codigo": _as_code(b, _ref(1, row_number)),
                     "tipo": _clean(c), "nombre": _name(d)}
            area_codigo, area = SNSS_TOTAL_AREA_CODE, SNSS_TOTAL_LABEL
        elif is_area:
            if block is None:
                raise ValueError(f"Fila {row_number} de la hoja {sheet!r}: área sin establecimiento previo.")
            if _clean(a) != block["servicio"]:
                raise ValueError(
                    f"Fila {row_number} de la hoja {sheet!r}: el servicio {_clean(a)!r} difiere del de "
                    f"su establecimiento {block['servicio']!r}.")
            area_codigo, area = _as_code(e, _ref(4, row_number)), _clean(f)
        else:
            raise _unrecognized(sheet, row_number, row)

        values = {year: _parse_value(row[index], _ref(index, row_number)) for index, year in year_columns}
        source.rows.append({
            "servicio_salud": block["servicio"], "establecimiento_codigo": block["codigo"],
            "tipo_establecimiento": block["tipo"], "establecimiento_nombre": block["nombre"],
            "area_funcional_codigo": area_codigo, "area_funcional": area,
            **{f"camas_{year}": value for year, value in values.items()},
        })
        if is_total:
            source.hospital_totals.extend(
                {"ano": year, "establecimiento_codigo": block["codigo"], "unidad": block["servicio"],
                 "bloque": block["id"], "valor": value}
                for year, value in values.items() if value is not None)
    return source


# --------------------------------------------------------------------------- #
# Libro no SNSS
# --------------------------------------------------------------------------- #

def _parse_no_snss(path: Path) -> _Source:
    source = _Source()
    for year in download.NO_SNSS_YEARS:
        sheet = str(year)
        rows = _read_rows(path, sheet)
        last_data_row = _last_data_row(rows)
        block: dict[str, Any] | None = None
        for row_number, row in enumerate(rows[FIRST_DATA_ROW - 1:], start=FIRST_DATA_ROW):
            if not any(_has_content(cell) for cell in row):
                continue
            a, b, c, d, e, f = row[:6]
            extras = [cell for cell in row[6:] if _has_content(cell)]

            if (_clean(a) or "").casefold() == "total" and not any(_has_content(cell) for cell in (b, c, d, e)):
                if extras:
                    raise ValueError(f"Fila {row_number} de la hoja {sheet!r} con datos fuera de las columnas contempladas.")
                source.national[year] = _parse_value(f, _ref(5, row_number))
                continue
            if _is_footnote(row, row_number, last_data_row):
                continue  # nota al pie: no es dato de dotación

            is_establishment = (_has_content(a) and _has_content(b) and _has_content(c)
                                and not _has_content(d) and not _has_content(e))
            is_area = (not any(_has_content(cell) for cell in (a, b, c))
                       and _has_content(d) and _has_content(e))
            if not (is_establishment or is_area):
                raise _unrecognized(sheet, row_number, row)
            # Una columna adicional solo puede traer una anotación de texto (``ok``), nunca datos.
            if len(extras) > 1 or any(not isinstance(cell, str) or re.search(r"\d", cell) for cell in extras):
                raise ValueError(f"Fila {row_number} de la hoja {sheet!r} con datos fuera de las columnas contempladas.")

            value = _parse_value(f, _ref(5, row_number))
            if is_establishment:
                block = {"id": row_number, "seremi": _clean(a), "codigo": _as_code(b, _ref(1, row_number)),
                         "nombre": _name(c)}
                area_codigo, area = None, None
                if value is not None:
                    source.hospital_totals.append(
                        {"ano": year, "establecimiento_codigo": block["codigo"], "unidad": block["seremi"],
                         "bloque": row_number, "valor": value})
            else:
                if block is None:
                    raise ValueError(f"Fila {row_number} de la hoja {sheet!r}: área sin establecimiento previo.")
                area_codigo, area = _as_code(d, _ref(3, row_number)), _clean(e)
            source.rows.append({
                "ano": year, "seremi": block["seremi"], "establecimiento_codigo": block["codigo"],
                "establecimiento_nombre": block["nombre"], "area_funcional_codigo": area_codigo,
                "area_funcional": area, "camas": value,
            })
    return source


# --------------------------------------------------------------------------- #
# Validación
# --------------------------------------------------------------------------- #

def _assert_unique(rows: list[dict[str, Any]], key: tuple[str, ...], label: str) -> None:
    """La llave lógica identifica una fila; si se repite no se puede elegir cuál vale."""
    if not rows:
        raise ValueError(f"Libro {label}: no produjo ninguna fila.")
    seen: set[tuple[Any, ...]] = set()
    for row in rows:
        identity = tuple(row[column] for column in key)
        if identity in seen:
            raise ValueError(f"Libro {label}: fila duplicada para {dict(zip(key, identity))}.")
        seen.add(identity)


def _reconcile(source: _Source, years: tuple[int, ...], label: str) -> list[tuple[int, int, int]]:
    """Contrasta el total nacional publicado con la suma de los totales por hospital."""
    missing = sorted(set(years) - set(source.national))
    if missing:
        raise ValueError(f"Libro {label}: sin total nacional publicado para los años {missing}.")
    sums: dict[int, int] = {}
    for total in source.hospital_totals:
        sums[total["ano"]] = sums.get(total["ano"], 0) + total["valor"]
    return [(year, source.national[year], sums.get(year, 0)) for year in years
            if source.national[year] is not None and source.national[year] != sums.get(year, 0)]


def _repeated_establishments(source: _Source) -> list[dict[str, Any]]:
    """Hospitales con más de un total no nulo en el mismo año (la propia fuente los cuenta dos veces)."""
    groups: dict[tuple[int, int], dict[int, str]] = {}
    for total in source.hospital_totals:
        groups.setdefault((total["ano"], total["establecimiento_codigo"]), {})[total["bloque"]] = total["unidad"]
    return [{"ano": year, "establecimiento_codigo": code, "unidades": sorted(set(blocks.values()))}
            for (year, code), blocks in sorted(groups.items()) if len(blocks) > 1]


# --------------------------------------------------------------------------- #
# Escritura atómica
# --------------------------------------------------------------------------- #

def _write_all(tables: dict[Path, pa.Table]) -> None:
    """Escribe todas las salidas en temporales y las publica juntas, con rollback."""
    directory = next(iter(tables)).parent
    directory.mkdir(parents=True, exist_ok=True)
    staged: dict[Path, Path] = {}
    backups: dict[Path, Path] = {}
    published: list[Path] = []
    try:
        for target, table in tables.items():
            with tempfile.NamedTemporaryFile(delete=False, dir=directory, prefix=f".{target.name}.",
                                             suffix=".tmp") as temp:
                temporary = Path(temp.name)
            staged[target] = temporary
            pq.write_table(table, temporary, compression="zstd", compression_level=ZSTD_LEVEL)
        for target, temporary in staged.items():
            if target.exists():
                backup = directory / f".{target.name}.previous"
                os.replace(target, backup)
                backups[target] = backup
            os.replace(temporary, target)
            published.append(target)
    except Exception:
        for target in published:
            target.unlink(missing_ok=True)
        for target, backup in backups.items():
            if backup.exists():
                os.replace(backup, target)
        raise
    finally:
        for temporary in staged.values():
            temporary.unlink(missing_ok=True)
        for backup in backups.values():
            backup.unlink(missing_ok=True)


# --------------------------------------------------------------------------- #
# Orquestación
# --------------------------------------------------------------------------- #

def _table(source: _Source, schema: pa.Schema, raw_path: Path) -> pa.Table:
    """Tabla del libro; el footer registra solo la procedencia (SHA256 del RAW de origen)."""
    table = pa.Table.from_pylist(source.rows, schema=schema)
    return table.replace_schema_metadata({RAW_SHA256_METADATA_KEY: file_sha256(raw_path).encode("ascii")})


def build_dotacion_camas(
    raw_snss_path: Path = RAW_SNSS_PATH,
    raw_no_snss_path: Path = RAW_NO_SNSS_PATH,
    output_dir: Path = OUTPUT_DIR,
) -> dict[str, Any]:
    """Genera los dos Parquet de dotación de camas (uno por libro) de forma validada y atómica."""
    for path in (raw_snss_path, raw_no_snss_path):
        if not path.exists():
            raise FileNotFoundError(f"No existe el RAW requerido: {path.as_posix()}")
    download.validate_camas_snss_workbook(raw_snss_path)
    download.validate_camas_no_snss_workbook(raw_no_snss_path)

    snss, no_snss = _parse_snss(raw_snss_path), _parse_no_snss(raw_no_snss_path)
    mismatches: list[tuple[str, int, int, int]] = []
    repeated = 0
    for source, key, years, label in ((snss, SNSS_KEY, download.SNSS_YEARS, "SNSS"),
                                      (no_snss, NO_SNSS_KEY, download.NO_SNSS_YEARS, "no SNSS")):
        _assert_unique(source.rows, key, label)
        mismatches += [(label, *item) for item in _reconcile(source, years, label)]
        for item in _repeated_establishments(source):
            repeated += 1
            logger.warning("Libro %s: el hospital %s figura más de una vez en %s (%s); se conservan todas las filas.",
                           label, item["establecimiento_codigo"], item["ano"], ", ".join(item["unidades"]))
    for label, year, published, total in mismatches:
        logger.warning("Libro %s: el total nacional de %s (%s) no coincide con la suma de los hospitales (%s).",
                       label, year, published, total)

    tables = {
        output_dir / SNSS_FILENAME: _table(snss, SNSS_SCHEMA, raw_snss_path),
        output_dir / NO_SNSS_FILENAME: _table(no_snss, NO_SNSS_SCHEMA, raw_no_snss_path),
    }
    _write_all(tables)
    return {
        "archivos_generados": [path.as_posix() for path in tables],
        "filas_snss": len(snss.rows),
        "filas_no_snss": len(no_snss.rows),
        "totales_que_no_coinciden": len(mismatches),
        "establecimientos_repetidos": repeated,
    }


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Limpia la dotación de camas hospitalarias (DEIS) a Parquet.")
    parser.add_argument("--force", action="store_true", help="Aceptado para interoperar con el orquestador.")
    parser.parse_args()
    logger.info("Dotación de camas generada: %s", build_dotacion_camas())


if __name__ == "__main__":
    main()
