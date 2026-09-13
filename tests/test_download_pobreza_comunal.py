"""Pruebas de ingesta RAW segura para el cuadro comunal MDS/Casen de pobreza por ingresos."""

import io
from pathlib import Path
from unittest.mock import patch

import openpyxl
import pytest

from src.data.download_pobreza_comunal import (
    ingest_pobreza_comunal,
    validate_pobreza_comunal_workbook,
)
import src.data.download_pobreza_comunal as module


def _workbook_bytes(*, rows: int = 341, cols: int = 10) -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = module.REQUIRED_SHEET
    header = [
        "Código", "Región", "Nombre comuna", "Personas proyectadas",
        "Personas en pobreza", "Porcentaje pobreza", "Límite inferior",
        "Límite superior", "Presencia muestra", "Tipo estimación SAE",
    ]
    sheet.append(header)
    for _ in range(rows - 1):
        sheet.append([None] * cols)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_transport_outside_raw_and_cleaned_after_success(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setattr(module, "MIN_ROWS", 10)
    raw_path = tmp_path / "raw" / "pobreza_comunal.xlsx"
    cache_dir = tmp_path / ".cache" / "downloads"
    with patch(
        "src.data.download_pobreza_comunal.urllib.request.urlopen",
        return_value=io.BytesIO(_workbook_bytes(rows=15)),
    ):
        assert ingest_pobreza_comunal(
            force=True, url="https://example.test/pobreza_comunal.xlsx", raw_path=raw_path,
            cache_dir=cache_dir, manifest_path=tmp_path / "manifest.json",
        )

    assert raw_path.is_file()
    assert not list(cache_dir.glob("*.xlsx"))
    validate_pobreza_comunal_workbook(raw_path)


def test_invalid_download_preserves_previous_raw(tmp_path: Path) -> None:
    raw_path = tmp_path / "raw" / "pobreza_comunal.xlsx"
    raw_path.parent.mkdir()
    raw_path.write_bytes(b"previous raw")
    cache_dir = tmp_path / ".cache" / "downloads"
    with patch(
        "src.data.download_pobreza_comunal.urllib.request.urlopen",
        return_value=io.BytesIO(b"not an xlsx"),
    ):
        with pytest.raises(Exception):
            ingest_pobreza_comunal(
                force=True, url="https://example.test/pobreza_comunal.xlsx", raw_path=raw_path,
                cache_dir=cache_dir, manifest_path=tmp_path / "manifest.json",
            )

    assert raw_path.read_bytes() == b"previous raw"


def test_skip_when_raw_already_valid(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setattr(module, "MIN_ROWS", 10)
    raw_path = tmp_path / "raw" / "pobreza_comunal.xlsx"
    raw_path.parent.mkdir()
    raw_path.write_bytes(_workbook_bytes(rows=15))
    cache_dir = tmp_path / ".cache" / "downloads"
    with patch("src.data.download_pobreza_comunal.urllib.request.urlopen") as mock_urlopen:
        published = ingest_pobreza_comunal(
            force=False, raw_path=raw_path, cache_dir=cache_dir, manifest_path=tmp_path / "manifest.json",
        )
    assert published is False
    mock_urlopen.assert_not_called()


def test_validate_rejects_missing_required_sheet(tmp_path: Path) -> None:
    workbook = openpyxl.Workbook()
    workbook.active.title = "otra_hoja"
    path = tmp_path / "sin_hoja.xlsx"
    workbook.save(path)
    with pytest.raises(ValueError):
        validate_pobreza_comunal_workbook(path)


def test_validate_rejects_undersized_sheet(tmp_path: Path) -> None:
    workbook = openpyxl.Workbook()
    workbook.active.title = module.REQUIRED_SHEET
    workbook.active.append(["a", "b"])
    path = tmp_path / "muy_chico.xlsx"
    workbook.save(path)
    with pytest.raises(ValueError):
        validate_pobreza_comunal_workbook(path)


def test_official_raw_exists() -> None:
    """Verifica que el snapshot RAW oficial esté presente para reproducibilidad local."""
    assert module.FINAL_OUTPUT_PATH.exists(), f"Falta archivo raw: {module.FINAL_OUTPUT_PATH}"
