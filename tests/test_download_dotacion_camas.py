"""Pruebas de ingesta RAW segura de la dotación de camas hospitalarias (DEIS)."""

import io
import json
from pathlib import Path
from unittest.mock import patch

import openpyxl
import pytest

import scripts.run_pipeline as pipeline
import src.data.download_dotacion_camas as module
from src.data.download_dotacion_camas import (
    CAMAS_SOURCES,
    ingest_all,
    ingest_camas_source,
    validate_camas_no_snss_workbook,
    validate_camas_snss_workbook,
)

SNSS, NO_SNSS = CAMAS_SOURCES


@pytest.fixture(autouse=True)
def _small_workbooks(monkeypatch) -> None:
    monkeypatch.setattr(module, "SNSS_MIN_ROWS", 5)
    monkeypatch.setattr(module, "NO_SNSS_MIN_ROWS", 5)


def _to_bytes(workbook: openpyxl.Workbook) -> bytes:
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _snss_bytes(*, years=module.SNSS_YEARS, sheet="2010-2025", rows=8, marker=0) -> bytes:
    workbook = openpyxl.Workbook()
    ws = workbook.active
    ws.title = sheet
    ws.append(["DOTACIÓN DE CAMAS HOSPITALARIAS (título)"])
    ws.append(["Fuente: DEIS, DIGERA"])
    ws.append([*module.SNSS_COLUMNS, *[f"{y}*" if y == 2014 else y for y in years]])
    for index in range(rows):
        ws.append(["Arica", 101100 + index, "Alta complejidad", "Hospital", 100, "Total",
                   *[marker + index] * len(years)])
    return _to_bytes(workbook)


def _no_snss_bytes(*, years=module.NO_SNSS_YEARS, rows=8, total_label=None) -> bytes:
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for year in years:
        ws = workbook.create_sheet(str(year))
        ws.append(["DATO REFERENCIAL DE DOTACIÓN DE CAMAS (título)"])
        ws.append(["Fuente: DEIS, DIGERA y Super Intendencia de Salud"])
        ws.append([*module.NO_SNSS_COLUMNS, total_label or f"Total {year}"])
        for index in range(rows):
            ws.append([" Arica y Parinacota", 101203 + index, "Clínica", None, None, 3])
    return _to_bytes(workbook)


class _Response(io.BytesIO):
    """Respuesta HTTP mínima con cabeceras, como la devuelta por urlopen."""

    def __init__(self, payload: bytes, headers: dict[str, str] | None = None) -> None:
        super().__init__(payload)
        self.headers = headers or {}


def _paths(tmp_path: Path) -> dict:
    return {"dest_dir": tmp_path / "raw" / "camas", "cache_dir": tmp_path / ".cache" / "downloads",
            "manifest_path": tmp_path / "manifest.json"}


def _manifest(tmp_path: Path) -> list[dict]:
    return json.loads((tmp_path / "manifest.json").read_text(encoding="utf-8"))


def test_validate_snss_accepts_expected_workbook(tmp_path: Path) -> None:
    path = tmp_path / "snss.xlsx"
    path.write_bytes(_snss_bytes())
    validate_camas_snss_workbook(path)


@pytest.mark.parametrize("payload, reason", [
    (lambda: _snss_bytes(sheet="otra"), "hoja requerida"),
    (lambda: _snss_bytes(years=tuple(range(2010, 2025))), "año 2025 ausente"),
    (lambda: _snss_bytes(rows=1), "pocas filas"),
])
def test_validate_snss_rejects_unexpected_workbook(tmp_path: Path, payload, reason) -> None:
    path = tmp_path / "snss.xlsx"
    path.write_bytes(payload())
    with pytest.raises(ValueError):
        validate_camas_snss_workbook(path)


def test_validate_snss_rejects_altered_header(tmp_path: Path) -> None:
    workbook = openpyxl.load_workbook(io.BytesIO(_snss_bytes()))
    workbook.active.cell(row=3, column=2).value = "Otro encabezado"
    path = tmp_path / "snss.xlsx"
    workbook.save(path)
    with pytest.raises(ValueError, match="Encabezado SNSS"):
        validate_camas_snss_workbook(path)


def test_validate_no_snss_accepts_expected_workbook(tmp_path: Path) -> None:
    path = tmp_path / "no_snss.xlsx"
    path.write_bytes(_no_snss_bytes())
    validate_camas_no_snss_workbook(path)


def test_validate_no_snss_rejects_missing_year_sheet(tmp_path: Path) -> None:
    path = tmp_path / "no_snss.xlsx"
    path.write_bytes(_no_snss_bytes(years=tuple(range(2019, 2025))))
    with pytest.raises(ValueError, match="Hojas anuales ausentes"):
        validate_camas_no_snss_workbook(path)


def test_validate_no_snss_rejects_undersized_sheets(tmp_path: Path) -> None:
    path = tmp_path / "no_snss.xlsx"
    path.write_bytes(_no_snss_bytes(rows=1))
    with pytest.raises(ValueError, match="menos filas"):
        validate_camas_no_snss_workbook(path)


def test_validate_no_snss_rejects_total_label_not_matching_sheet(tmp_path: Path) -> None:
    path = tmp_path / "no_snss.xlsx"
    path.write_bytes(_no_snss_bytes(total_label="Total 1999"))
    with pytest.raises(ValueError, match="Encabezado no SNSS"):
        validate_camas_no_snss_workbook(path)


def test_validators_reject_the_other_source(tmp_path: Path) -> None:
    snss, no_snss = tmp_path / "snss.xlsx", tmp_path / "no_snss.xlsx"
    snss.write_bytes(_snss_bytes())
    no_snss.write_bytes(_no_snss_bytes())
    with pytest.raises(ValueError):
        validate_camas_snss_workbook(no_snss)
    with pytest.raises(ValueError):
        validate_camas_no_snss_workbook(snss)


def test_ingest_publishes_with_provenance_and_cleans_cache(tmp_path: Path) -> None:
    headers = {"Last-Modified": "Mon, 20 Oct 2025 16:12:50 GMT", "ETag": '"abc"'}
    with patch("src.data.download_dotacion_camas.urllib.request.urlopen",
               return_value=_Response(_snss_bytes(), headers)):
        assert ingest_camas_source(SNSS, force=True, **_paths(tmp_path))

    raw_path = tmp_path / "raw" / "camas" / SNSS.filename
    assert raw_path.is_file()
    assert not list((tmp_path / ".cache" / "downloads").glob("*"))
    (entry,) = _manifest(tmp_path)
    assert entry["source_id"] == SNSS.source_id
    assert entry["periodo"] == "2010-2025"
    (snapshot,) = entry["snapshots"]
    assert snapshot["source_url"] == SNSS.url
    assert snapshot["raw_sha256"] == snapshot["transport"]["sha256"]
    assert snapshot["transport"]["http_last_modified"] == headers["Last-Modified"]
    assert snapshot["transport"]["http_etag"] == headers["ETag"]


def test_ingest_without_http_headers_still_publishes(tmp_path: Path) -> None:
    with patch("src.data.download_dotacion_camas.urllib.request.urlopen",
               return_value=io.BytesIO(_no_snss_bytes())):
        assert ingest_camas_source(NO_SNSS, force=True, **_paths(tmp_path))
    (snapshot,) = _manifest(tmp_path)[0]["snapshots"]
    assert "http_last_modified" not in snapshot["transport"]


def test_skip_when_raw_already_valid(tmp_path: Path) -> None:
    paths = _paths(tmp_path)
    paths["dest_dir"].mkdir(parents=True)
    (paths["dest_dir"] / SNSS.filename).write_bytes(_snss_bytes())
    with patch("src.data.download_dotacion_camas.urllib.request.urlopen") as mock_urlopen:
        assert ingest_camas_source(SNSS, force=False, **paths) is False
    mock_urlopen.assert_not_called()
    assert not paths["manifest_path"].exists()


def test_force_with_identical_content_does_not_duplicate_history(tmp_path: Path) -> None:
    payload = _snss_bytes()
    with patch("src.data.download_dotacion_camas.urllib.request.urlopen",
               side_effect=lambda *a, **k: _Response(payload)):
        assert ingest_camas_source(SNSS, force=True, **_paths(tmp_path))
        assert ingest_camas_source(SNSS, force=True, **_paths(tmp_path)) is False
    assert len(_manifest(tmp_path)[0]["snapshots"]) == 1


def test_force_with_changed_content_replaces_raw_and_appends_history(tmp_path: Path) -> None:
    first, second = _snss_bytes(marker=0), _snss_bytes(marker=10)
    with patch("src.data.download_dotacion_camas.urllib.request.urlopen",
               side_effect=[_Response(first), _Response(second)]):
        assert ingest_camas_source(SNSS, force=True, **_paths(tmp_path))
        assert ingest_camas_source(SNSS, force=True, **_paths(tmp_path))
    raw_path = tmp_path / "raw" / "camas" / SNSS.filename
    assert raw_path.read_bytes() == second
    snapshots = _manifest(tmp_path)[0]["snapshots"]
    assert len(snapshots) == 2
    assert snapshots[0]["raw_sha256"] != snapshots[1]["raw_sha256"]


def test_force_with_identical_raw_registers_missing_provenance_once(tmp_path: Path) -> None:
    """Un RAW sin entrada en el manifiesto se registra, sin inventar una fecha anterior."""
    paths = _paths(tmp_path)
    payload = _snss_bytes()
    paths["dest_dir"].mkdir(parents=True)
    (paths["dest_dir"] / SNSS.filename).write_bytes(payload)
    with patch("src.data.download_dotacion_camas.urllib.request.urlopen",
               side_effect=lambda *a, **k: _Response(payload)):
        assert ingest_camas_source(SNSS, force=True, **paths) is False
        assert ingest_camas_source(SNSS, force=True, **paths) is False
    assert len(_manifest(tmp_path)[0]["snapshots"]) == 1


def test_invalid_download_preserves_previous_raw(tmp_path: Path) -> None:
    paths = _paths(tmp_path)
    paths["dest_dir"].mkdir(parents=True)
    raw_path = paths["dest_dir"] / SNSS.filename
    raw_path.write_bytes(b"previous raw")
    with patch("src.data.download_dotacion_camas.urllib.request.urlopen",
               return_value=io.BytesIO(b"not an xlsx")):
        with pytest.raises(Exception):
            ingest_camas_source(SNSS, force=True, **paths)
    assert raw_path.read_bytes() == b"previous raw"
    assert not paths["manifest_path"].exists()
    assert not list(paths["cache_dir"].glob("*"))


def test_empty_download_is_rejected(tmp_path: Path) -> None:
    with patch("src.data.download_dotacion_camas.urllib.request.urlopen",
               return_value=io.BytesIO(b"")):
        with pytest.raises(ValueError, match="vacío"):
            ingest_camas_source(NO_SNSS, force=True, **_paths(tmp_path))
    assert not (tmp_path / "raw" / "camas" / NO_SNSS.filename).exists()


def test_provenance_failure_rolls_back_to_previous_raw(tmp_path: Path) -> None:
    paths = _paths(tmp_path)
    paths["dest_dir"].mkdir(parents=True)
    previous = _snss_bytes(marker=0)
    raw_path = paths["dest_dir"] / SNSS.filename
    raw_path.write_bytes(previous)
    with patch("src.data.download_dotacion_camas.urllib.request.urlopen",
               return_value=_Response(_snss_bytes(marker=10))), \
            patch("src.data.download_dotacion_camas.record_raw_snapshot",
                  side_effect=OSError("disco lleno")):
        with pytest.raises(OSError):
            ingest_camas_source(SNSS, force=True, **paths)
    assert raw_path.read_bytes() == previous


def test_sources_are_independent(tmp_path: Path) -> None:
    """Si una fuente falla, la otra igual se publica y el error queda informado."""
    def fake_urlopen(request, **_kwargs):
        if request.full_url == SNSS.url:
            return _Response(_snss_bytes())
        return io.BytesIO(b"not an xlsx")

    with patch("src.data.download_dotacion_camas.urllib.request.urlopen", side_effect=fake_urlopen):
        errors = ingest_all(force=True, **_paths(tmp_path))

    assert set(errors) == {"no_snss"}
    assert (tmp_path / "raw" / "camas" / SNSS.filename).is_file()
    assert not (tmp_path / "raw" / "camas" / NO_SNSS.filename).exists()


def test_sources_have_distinct_ids_files_and_https_urls() -> None:
    assert len({source.source_id for source in CAMAS_SOURCES}) == len(CAMAS_SOURCES)
    assert len({source.filename for source in CAMAS_SOURCES}) == len(CAMAS_SOURCES)
    assert all(source.url.startswith("https://repositoriodeis.minsal.cl/") for source in CAMAS_SOURCES)
    assert all(source.filename.endswith(".xlsx") for source in CAMAS_SOURCES)


def test_download_dotacion_camas_stage_is_registered() -> None:
    stage = pipeline.STAGES["download_dotacion_camas"]
    assert stage["module"] == "src.data.download_dotacion_camas"
    assert stage["depends_on"] == []
    assert stage["outputs"] == [f"data/raw/camas/{source.filename}" for source in CAMAS_SOURCES]
    assert "download_dotacion_camas" in pipeline.PIPELINE_ORDER


def test_output_contract_validates_camas_raw(tmp_path: Path) -> None:
    raw_dir = tmp_path / "raw" / "camas"
    raw_dir.mkdir(parents=True)
    snss, no_snss = raw_dir / SNSS.filename, raw_dir / NO_SNSS.filename
    snss.write_bytes(_snss_bytes())
    no_snss.write_bytes(_no_snss_bytes())
    assert pipeline.check_outputs_exist([snss.as_posix(), no_snss.as_posix()])

    # El contrato distingue la fuente: un libro cruzado no se acepta.
    no_snss.write_bytes(_snss_bytes())
    assert not pipeline.check_outputs_exist([snss.as_posix(), no_snss.as_posix()])


@pytest.mark.data_integration
def test_official_raw_exists_and_is_valid() -> None:
    """Verifica que los snapshots RAW oficiales estén presentes y sean válidos localmente."""
    for source in CAMAS_SOURCES:
        path = module.DEST_DIR / source.filename
        assert path.exists(), f"Falta archivo raw: {path}"
        source.validator(path)
