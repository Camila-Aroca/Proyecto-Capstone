"""Pruebas de ingesta RAW segura para el tabulado comunal de población Censo 2024."""

import io
from pathlib import Path
from unittest.mock import patch

import openpyxl
import pytest

from src.data.download_censo_poblacion import ingest_censo_poblacion, validate_censo_poblacion_workbook


def _workbook_bytes(*, rows_sheet2: int = 353) -> bytes:
    workbook = openpyxl.Workbook()
    sheet1 = workbook.active
    sheet1.title = "1"
    sheet1.append(["Código región", "Región", "Población censada", "Hombres", "Mujeres", "Razón"])
    sheet1.append([13, "Metropolitana de Santiago", 7400741, 3582833, 3817908, 93.8])

    sheet2 = workbook.create_sheet("2")
    for _ in range(rows_sheet2 - 1):
        sheet2.append([None] * 7)
    sheet2.append([13, "Metropolitana de Santiago", 131, "Santiago", 13101, "Santiago", 438856])

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_transport_outside_raw_and_cleaned_after_success(tmp_path: Path) -> None:
    raw_path = tmp_path / "raw" / "poblacion.xlsx"
    cache_dir = tmp_path / ".cache" / "downloads"
    with patch(
        "src.data.download_censo_poblacion.urllib.request.urlopen",
        return_value=io.BytesIO(_workbook_bytes()),
    ):
        assert ingest_censo_poblacion(
            force=True, url="https://example.test/censo_poblacion.xlsx", raw_path=raw_path,
            cache_dir=cache_dir, manifest_path=tmp_path / "manifest.json",
        )

    assert raw_path.is_file()
    assert not list(cache_dir.glob("*.xlsx"))
    validate_censo_poblacion_workbook(raw_path)


def test_invalid_download_preserves_previous_raw(tmp_path: Path) -> None:
    raw_path = tmp_path / "raw" / "poblacion.xlsx"
    raw_path.parent.mkdir()
    raw_path.write_bytes(b"previous raw")
    cache_dir = tmp_path / ".cache" / "downloads"
    with patch(
        "src.data.download_censo_poblacion.urllib.request.urlopen",
        return_value=io.BytesIO(b"not an xlsx"),
    ):
        with pytest.raises(Exception):
            ingest_censo_poblacion(
                force=True, url="https://example.test/censo_poblacion.xlsx", raw_path=raw_path,
                cache_dir=cache_dir, manifest_path=tmp_path / "manifest.json",
            )

    assert raw_path.read_bytes() == b"previous raw"


def test_skip_when_raw_already_valid(tmp_path: Path) -> None:
    raw_path = tmp_path / "raw" / "poblacion.xlsx"
    raw_path.parent.mkdir()
    raw_path.write_bytes(_workbook_bytes())
    cache_dir = tmp_path / ".cache" / "downloads"
    with patch("src.data.download_censo_poblacion.urllib.request.urlopen") as mock_urlopen:
        published = ingest_censo_poblacion(
            force=False, raw_path=raw_path, cache_dir=cache_dir, manifest_path=tmp_path / "manifest.json",
        )
    assert published is False
    mock_urlopen.assert_not_called()


def test_validate_rejects_missing_required_sheet(tmp_path: Path) -> None:
    workbook = openpyxl.Workbook()
    workbook.active.title = "1"
    path = tmp_path / "sin_hoja2.xlsx"
    workbook.save(path)
    with pytest.raises(ValueError):
        validate_censo_poblacion_workbook(path)
